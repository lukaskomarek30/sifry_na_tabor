APP_VERSION = "0.0.4"
APP_NAME = "Sifrator_Mraveniste"
"""
main_window.py – SifratorWindow, LiveLogDialog a všechny rozšiřující patche.

Obsahuje:
  - Monkey-patche Zednářské šifry pro SifratorSkinWidget
  - Monkey-patche Caesarova UI pro SifratorSkinWidget
  - Tisk (print patches)
  - SifratorWindow (QMainWindow obálka s ScrollArea)
  - LiveLogDialog (okno živého logu)

Závislosti:
    skin_widget     – SifratorSkinWidget
    cipher_registry – get_cipher_logic, get_cipher_widget_class, get_pirate_key_renderer
    ui_widgets      – Colors, CaesarDirectionCombo
    app_paths       – cesty
"""

import os
import hashlib
import json
import random
import shutil
import sys
import tempfile
import time
import unicodedata
import zipfile
from datetime import date, datetime, timedelta
from xml.etree import ElementTree as ET

from PySide6.QtCore import QEasingCurve, QPropertyAnimation, Qt, QRect, QRectF, QSize, QTimer, QUrl, QEvent, QDate
from PySide6.QtGui import QColor, QDesktopServices, QFont, QIcon, QImage, QPainter, QPixmap, QTextOption, QPen, QTextBlockFormat, QTextCursor
from PySide6.QtWidgets import (
    QApplication,
    QLabel,
    QLineEdit,
    QMainWindow,
    QPushButton,
    QScrollArea,
    QSizePolicy,
    QTabWidget,
    QTextEdit,
    QWidget,
    QGridLayout,
    QGraphicsOpacityEffect,
    QListWidget,
    QListWidgetItem,
    QMessageBox,
    QDialog,
    QVBoxLayout,
    QHBoxLayout,
    QComboBox,
    QCheckBox,
    QDateEdit,
    QFileDialog,
    QFrame,
    QSpinBox,
    QAbstractScrollArea,
    QAbstractSpinBox,
    QStackedWidget,
)

try:
    from PIL import Image
except Exception:
    Image = None

import update_manager

from app_paths import (
    get_app_dir,
    get_script_dir,
    get_pyinstaller_bundle_dir,
    get_icons_dir,
    get_user_data_dir,
    migrate_user_data_items,
)
from cipher_registry import get_cipher_logic, get_cipher_widget_class, get_pirate_key_renderer, list_cipher_names
from ui_widgets import Colors, CaesarDirectionCombo, CipherItem
from skin_widget import SifratorSkinWidget, BASE_W, BASE_H
from home_menu import PirateHomeWidget
from sports_day import SportsDayDialog
from diploma import DiplomaDialog
from groups import GroupsDialog
from fire_effects import PirateModuleDialog
from user_data_backup import (
    UserDataBackupError,
    default_backup_filename,
    export_user_data_zip,
    import_user_data_zip,
    inspect_user_data_backup,
)


_RANDOM_EASY_CIPHER_NAME = "Náhodná lehká šifra"
_RANDOM_EASY_CIPHER_ICON = "compass.png"


_RANDOM_EASY_ORIGINAL_BUILD_CIPHER_LIST = SifratorSkinWidget.build_cipher_list


def _random_easy_build_cipher_list(self):
    items = list(_RANDOM_EASY_ORIGINAL_BUILD_CIPHER_LIST(self))
    if not any(item.name == _RANDOM_EASY_CIPHER_NAME for item in items):
        items.append(CipherItem(_RANDOM_EASY_CIPHER_NAME, _RANDOM_EASY_CIPHER_ICON))

    def sort_key(item):
        text = unicodedata.normalize("NFKD", item.name.lower())
        return "".join(ch for ch in text if not unicodedata.combining(ch))

    return sorted(items, key=sort_key)


SifratorSkinWidget.build_cipher_list = _random_easy_build_cipher_list




# ============================================================
# ROZŠÍŘENÍ: Zednářská šifra
# Tato část je oddělená jako samostatná integrace, aby nebylo nutné
# neúměrně rozšiřovat původní metody hlavního widgetu.
# ============================================================

def _zednarska_sifra_is_selected(self):
    return self.selected_cipher == "Zednářská šifra"


def _zednarska_sifra_create_canvas(self):
    widget_class = get_cipher_widget_class("Zednářská šifra")

    if widget_class is not None:
        return widget_class()

    fallback = QLabel("Chybí kreslicí modul Zednářská šifra")
    fallback.setStyleSheet("background: transparent; color: #f3d79a;")
    return fallback


def _zednarska_sifra_resize_canvas_to_content(self, text: str | None = None):
    if not hasattr(self, "zednarska_sifra_scroll") or not hasattr(self, "zednarska_sifra_canvas"):
        return

    viewport_width = self.zednarska_sifra_scroll.viewport().width()
    if viewport_width <= 20:
        viewport_width = self.output_text.width()

    width = max(250, viewport_width)

    if hasattr(self.zednarska_sifra_canvas, "calculate_required_height"):
        height = max(
            self.zednarska_sifra_scroll.height(),
            self.zednarska_sifra_canvas.calculate_required_height(width),
        )
    else:
        height = max(self.zednarska_sifra_scroll.height(), 170)

    self.zednarska_sifra_canvas.setMinimumSize(width, height)
    self.zednarska_sifra_canvas.resize(width, height)

    if hasattr(self.zednarska_sifra_canvas, "update_content_size"):
        self.zednarska_sifra_canvas.update_content_size()


def _zednarska_sifra_hide_other_draw_outputs(self):
    """Schová všechny kreslené výstupy kromě Zednářské šifry."""
    for attr_name in dir(self):
        if not attr_name.endswith("_scroll"):
            continue

        if attr_name == "zednarska_sifra_scroll":
            continue

        widget = getattr(self, attr_name, None)
        if hasattr(widget, "hide"):
            widget.hide()


def _zednarska_sifra_scroll_style():
    return """
        QScrollArea {
            background: rgba(0, 0, 0, 0);
            border: none;
        }
        QScrollBar:vertical {
            background: rgba(20, 17, 12, 130);
            width: 11px;
            margin: 4px 2px 4px 2px;
            border-radius: 5px;
        }
        QScrollBar::handle:vertical {
            background: #b89b68;
            min-height: 42px;
            border-radius: 5px;
        }
        QScrollBar::add-line:vertical,
        QScrollBar::sub-line:vertical {
            height: 0px;
        }
    """


_ORIGINAL_CREATE_WIDGETS = SifratorSkinWidget.create_widgets
_ORIGINAL_UPDATE_LAYOUT_POSITIONS = SifratorSkinWidget.update_layout_positions
_ORIGINAL_APPLY_RESPONSIVE_FONTS = SifratorSkinWidget.apply_responsive_fonts
_ORIGINAL_SHOW_PLAIN_OUTPUT = SifratorSkinWidget.show_plain_output
_ORIGINAL_SET_RESULT_OUTPUT = SifratorSkinWidget.set_result_output
_ORIGINAL_UPDATE_OUTPUT_WIDGET_MODE = SifratorSkinWidget.update_output_widget_mode
_ORIGINAL_ENCRYPT_SELECTED_CIPHER = SifratorSkinWidget.encrypt_selected_cipher
_ORIGINAL_DECRYPT_SELECTED_CIPHER = SifratorSkinWidget.decrypt_selected_cipher


def _PATCHED_CREATE_WIDGETS(self):
    _ORIGINAL_CREATE_WIDGETS(self)

    self.zednarska_sifra_canvas = self.create_zednarska_sifra_canvas()

    self.zednarska_sifra_scroll = QScrollArea(self)
    self.zednarska_sifra_scroll.setWidgetResizable(False)
    self.zednarska_sifra_scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
    self.zednarska_sifra_scroll.setVerticalScrollBarPolicy(Qt.ScrollBarAsNeeded)
    self.zednarska_sifra_scroll.setFrameShape(QScrollArea.NoFrame)
    self.zednarska_sifra_scroll.setWidget(self.zednarska_sifra_canvas)
    self.zednarska_sifra_scroll.setStyleSheet(_zednarska_sifra_scroll_style())
    self.zednarska_sifra_scroll.hide()


def _PATCHED_UPDATE_LAYOUT_POSITIONS(self):
    _ORIGINAL_UPDATE_LAYOUT_POSITIONS(self)

    if hasattr(self, "zednarska_sifra_scroll"):
        self.zednarska_sifra_scroll.setGeometry(self.output_text.geometry())
        self.zednarska_sifra_scroll.raise_()
        self.resize_zednarska_sifra_canvas_to_content()


def _PATCHED_APPLY_RESPONSIVE_FONTS(self):
    _ORIGINAL_APPLY_RESPONSIVE_FONTS(self)

    if hasattr(self, "zednarska_sifra_canvas") and hasattr(self.zednarska_sifra_canvas, "set_scale"):
        self.zednarska_sifra_canvas.set_scale(self.sc())

    if hasattr(self, "zednarska_sifra_scroll"):
        self.resize_zednarska_sifra_canvas_to_content()


def _PATCHED_SHOW_PLAIN_OUTPUT(self, text: str):
    if hasattr(self, "zednarska_sifra_scroll"):
        self.zednarska_sifra_scroll.hide()

    return _ORIGINAL_SHOW_PLAIN_OUTPUT(self, text)


def _PATCHED_SET_RESULT_OUTPUT(self, result: str):
    if self.is_zednarska_sifra_selected() and self.result_mode == "encrypt":
        self.output_text.clear()
        self.output_text.hide()

        self.zednarska_sifra_hide_other_draw_outputs()

        self.zednarska_sifra_scroll.show()
        self.zednarska_sifra_scroll.raise_()
        self.zednarska_sifra_scroll.verticalScrollBar().setValue(0)

        self.resize_zednarska_sifra_canvas_to_content(result)

        if hasattr(self.zednarska_sifra_canvas, "set_cipher_text"):
            self.zednarska_sifra_canvas.set_cipher_text(result)
        elif hasattr(self.zednarska_sifra_canvas, "setText"):
            self.zednarska_sifra_canvas.setText(result)

        self.resize_zednarska_sifra_canvas_to_content(result)
        QTimer.singleShot(0, lambda: self.resize_zednarska_sifra_canvas_to_content(result))
        return

    if hasattr(self, "zednarska_sifra_scroll"):
        self.zednarska_sifra_scroll.hide()

    return _ORIGINAL_SET_RESULT_OUTPUT(self, result)


def _PATCHED_UPDATE_OUTPUT_WIDGET_MODE(self):
    _ORIGINAL_UPDATE_OUTPUT_WIDGET_MODE(self)

    if self.is_zednarska_sifra_selected() and self.result_mode == "encrypt" and hasattr(self, "zednarska_sifra_scroll"):
        self.output_text.hide()
        self.zednarska_sifra_hide_other_draw_outputs()
        self.zednarska_sifra_scroll.show()
        self.zednarska_sifra_scroll.raise_()
        self.resize_zednarska_sifra_canvas_to_content()
        return

    if hasattr(self, "zednarska_sifra_scroll"):
        self.zednarska_sifra_scroll.hide()


def _PATCHED_ENCRYPT_SELECTED_CIPHER(self, text: str) -> str:
    if self.is_zednarska_sifra_selected():
        zednarska_sifra_logic = get_cipher_logic("Zednářská šifra")

        if zednarska_sifra_logic is None:
            return (
                "Chybí soubor s logikou šifry Zednářská šifra:\n"
                "logika_sifer/Zednarska_sifra/zednarska_sifra.py"
            )

        return zednarska_sifra_logic.encrypt(text)

    return _ORIGINAL_ENCRYPT_SELECTED_CIPHER(self, text)


def _PATCHED_DECRYPT_SELECTED_CIPHER(self, text: str) -> str:
    if self.is_zednarska_sifra_selected():
        zednarska_sifra_logic = get_cipher_logic("Zednářská šifra")

        if zednarska_sifra_logic is None:
            return (
                "Chybí soubor s logikou šifry Zednářská šifra:\n"
                "logika_sifer/Zednarska_sifra/zednarska_sifra.py"
            )

        return zednarska_sifra_logic.decrypt(text)

    return _ORIGINAL_DECRYPT_SELECTED_CIPHER(self, text)


SifratorSkinWidget.is_zednarska_sifra_selected = _zednarska_sifra_is_selected
SifratorSkinWidget.create_zednarska_sifra_canvas = _zednarska_sifra_create_canvas
SifratorSkinWidget.resize_zednarska_sifra_canvas_to_content = _zednarska_sifra_resize_canvas_to_content
SifratorSkinWidget.zednarska_sifra_hide_other_draw_outputs = _zednarska_sifra_hide_other_draw_outputs

SifratorSkinWidget.create_widgets = _PATCHED_CREATE_WIDGETS
SifratorSkinWidget.update_layout_positions = _PATCHED_UPDATE_LAYOUT_POSITIONS
SifratorSkinWidget.apply_responsive_fonts = _PATCHED_APPLY_RESPONSIVE_FONTS
SifratorSkinWidget.show_plain_output = _PATCHED_SHOW_PLAIN_OUTPUT
SifratorSkinWidget.set_result_output = _PATCHED_SET_RESULT_OUTPUT
SifratorSkinWidget.update_output_widget_mode = _PATCHED_UPDATE_OUTPUT_WIDGET_MODE
SifratorSkinWidget.encrypt_selected_cipher = _PATCHED_ENCRYPT_SELECTED_CIPHER
SifratorSkinWidget.decrypt_selected_cipher = _PATCHED_DECRYPT_SELECTED_CIPHER



# ============================================================
# UI ROZŠÍŘENÍ PRO CAESAROVU ŠIFRU
#
# Rozšíření se aktivuje pouze pro Caesarovu šifru:
# - levá akční oblast slouží pro volbu směru posunu,
# - pravá akční oblast slouží pro nastavení velikosti posunu,
# - ostatní šifry zachovávají standardní akční tlačítka.
# ============================================================

_CAESAR_UI_ORIGINAL_CREATE_WIDGETS = SifratorSkinWidget.create_widgets
_CAESAR_UI_ORIGINAL_UPDATE_LAYOUT_POSITIONS = SifratorSkinWidget.update_layout_positions
_CAESAR_UI_ORIGINAL_APPLY_RESPONSIVE_FONTS = SifratorSkinWidget.apply_responsive_fonts
_CAESAR_UI_ORIGINAL_SELECT_CIPHER = SifratorSkinWidget.select_cipher
_CAESAR_UI_ORIGINAL_ENCRYPT_SELECTED_CIPHER = SifratorSkinWidget.encrypt_selected_cipher
_CAESAR_UI_ORIGINAL_DECRYPT_SELECTED_CIPHER = SifratorSkinWidget.decrypt_selected_cipher


class CaesarDirectionCombo(QComboBox):
    """Skutečný rozevírací seznam, ale vykreslený jako text uprostřed tlačítka."""

    def __init__(self, parent=None):
        super().__init__(parent)
        self._hovered = False
        self.setMouseTracking(True)

    def enterEvent(self, event):
        self._hovered = True
        self.update()
        super().enterEvent(event)

    def leaveEvent(self, event):
        self._hovered = False
        self.update()
        super().leaveEvent(event)

    def paintEvent(self, event):
        painter = QPainter(self)
        painter.setRenderHint(QPainter.Antialiasing, True)
        painter.setRenderHint(QPainter.TextAntialiasing, True)

        if self._hovered:
            painter.setPen(Qt.NoPen)
            painter.setBrush(QColor(255, 220, 120, 26))
            painter.drawRoundedRect(self.rect().adjusted(5, 5, -5, -5), 10, 10)

        painter.setFont(self.font())
        painter.setPen(QColor(Colors.GOLD_LIGHT if not self._hovered else "#fff0bd"))

        # Text je centrovaný přes celou plochu původního grafického tlačítka.
        # Vpravo zůstává pouze minimální prostor pro decentní indikaci rozbalení.
        text_rect = self.rect().adjusted(26, 0, -54, 0)
        painter.drawText(text_rect, Qt.AlignCenter, self.currentText())

        # Vlastní indikátor šipky zachovává vzhled skinu a současně naznačuje rozbalovací prvek.
        arrow_font = QFont(self.font())
        arrow_font.setPointSize(max(10, int(self.font().pointSize() * 0.58)))
        painter.setFont(arrow_font)
        arrow_rect = QRect(self.width() - 56, 0, 38, self.height())
        painter.drawText(arrow_rect, Qt.AlignCenter, "▼")


def _caesar_ui_is_selected(self):
    return self.selected_cipher == "Caesarova šifra"


def _caesar_ui_style():
    # Ovládací prvky jsou položeny přímo na grafický skin původních tlačítek.
    # Bez dodatečných rámečků, aby vizuálně navazovaly na původní tlačítka.
    return f"""
        QComboBox, QSpinBox {{
            color: {Colors.GOLD_LIGHT};
            background: rgba(0, 0, 0, 0);
            border: none;
            padding: 0px;
            margin: 0px;
            selection-background-color: {Colors.GOLD};
            selection-color: #111111;
        }}
        QComboBox:hover, QSpinBox:hover {{
            color: #fff0bd;
            background: rgba(255, 220, 120, 22);
            border: none;
            border-radius: 10px;
        }}
        QComboBox::drop-down {{
            border: none;
            background: transparent;
            width: 0px;
        }}
        QComboBox::down-arrow {{
            image: none;
            width: 0px;
            height: 0px;
        }}
        QSpinBox::up-button, QSpinBox::down-button {{
            width: 0px;
            height: 0px;
            border: none;
            background: transparent;
        }}
        QLineEdit {{
            color: {Colors.GOLD_LIGHT};
            background: transparent;
            border: none;
            padding: 0px;
            margin: 0px;
        }}
        QComboBox QAbstractItemView {{
            color: {Colors.GOLD_LIGHT};
            background: rgba(7, 18, 22, 248);
            border: 2px solid {Colors.GOLD};
            selection-background-color: rgba(200, 154, 76, 155);
            selection-color: #fff0bd;
            outline: 0px;
            padding: 8px;
        }}
    """


def _caesar_ui_create_widgets(self):
    _CAESAR_UI_ORIGINAL_CREATE_WIDGETS(self)

    self.caesar_direction_combo = CaesarDirectionCombo(self)
    self.caesar_direction_combo.addItems(["DOPŘEDU", "DOZADU"])
    self.caesar_direction_combo.setCursor(Qt.PointingHandCursor)
    self.caesar_direction_combo.setEditable(False)
    self.caesar_direction_combo.setMaxVisibleItems(2)
    self.caesar_direction_combo.setStyleSheet(_caesar_ui_style())
    self.caesar_direction_combo.view().setStyleSheet(_caesar_ui_style())
    self.caesar_direction_combo.currentIndexChanged.connect(self.auto_encrypt_action)
    self.caesar_direction_combo.hide()

    self.caesar_shift_input = QSpinBox(self)
    self.caesar_shift_input.setRange(0, 999)
    self.caesar_shift_input.setValue(3)
    self.caesar_shift_input.setPrefix("O KOLIK: ")
    self.caesar_shift_input.setAlignment(Qt.AlignCenter)
    self.caesar_shift_input.setKeyboardTracking(True)
    self.caesar_shift_input.setFrame(False)
    self.caesar_shift_input.setButtonSymbols(QAbstractSpinBox.NoButtons)
    self.caesar_shift_input.lineEdit().setAlignment(Qt.AlignCenter)
    self.caesar_shift_input.lineEdit().setStyleSheet(
        f"color: {Colors.GOLD_LIGHT}; background: transparent; border: none; padding: 0px; margin: 0px;"
    )
    self.caesar_shift_input.setStyleSheet(_caesar_ui_style())
    self.caesar_shift_input.valueChanged.connect(self.auto_encrypt_action)
    self.caesar_shift_input.hide()

    self.update_caesar_controls_visibility()


def _caesar_ui_update_layout_positions(self):
    _CAESAR_UI_ORIGINAL_UPDATE_LAYOUT_POSITIONS(self)

    # Využívá se plocha původních tlačítek s malou bezpečnostní rezervou,
    # aby text vizuálně nekolidoval s dekorativním okrajem skinu.
    inner_x = max(3, int(8 * self.sx()))
    inner_y = max(2, int(5 * self.sy()))

    if hasattr(self, "caesar_direction_combo"):
        rect = self.encrypt_button.geometry().adjusted(inner_x, inner_y, -inner_x, -inner_y)
        self.caesar_direction_combo.setGeometry(rect)
        self.caesar_direction_combo.raise_()

    if hasattr(self, "caesar_shift_input"):
        rect = self.decrypt_button.geometry().adjusted(inner_x, inner_y, -inner_x, -inner_y)
        self.caesar_shift_input.setGeometry(rect)
        self.caesar_shift_input.raise_()
        if self.caesar_shift_input.lineEdit() is not None:
            self.caesar_shift_input.lineEdit().setAlignment(Qt.AlignCenter)

    self.update_caesar_controls_visibility()


def _caesar_ui_apply_responsive_fonts(self):
    _CAESAR_UI_ORIGINAL_APPLY_RESPONSIVE_FONTS(self)

    # Velikost textu je opticky sjednocená s původními akčními tlačítky.
    font = QFont("Georgia", self.fs(20), QFont.Bold)
    if hasattr(self, "caesar_direction_combo"):
        self.caesar_direction_combo.setFont(font)
        self.caesar_direction_combo.update()
    if hasattr(self, "caesar_shift_input"):
        self.caesar_shift_input.setFont(font)
        if self.caesar_shift_input.lineEdit() is not None:
            self.caesar_shift_input.lineEdit().setFont(font)
            self.caesar_shift_input.lineEdit().setAlignment(Qt.AlignCenter)


def _caesar_ui_update_controls_visibility(self):
    is_caesar = self.is_caesar_selected()

    # U Caesarovy šifry se dvě akční oblasti přepnou na konfiguraci směru a posunu.
    # U ostatních šifer se zachová standardní režim tlačítek.
    if hasattr(self, "encrypt_button"):
        self.encrypt_button.setVisible(not is_caesar)
    if hasattr(self, "decrypt_button"):
        self.decrypt_button.setVisible(not is_caesar)

    if hasattr(self, "caesar_direction_combo"):
        self.caesar_direction_combo.setVisible(is_caesar)
        self.caesar_direction_combo.setEnabled(is_caesar)
        if is_caesar:
            self.caesar_direction_combo.raise_()
    if hasattr(self, "caesar_shift_input"):
        self.caesar_shift_input.setVisible(is_caesar)
        self.caesar_shift_input.setEnabled(is_caesar)
        if is_caesar:
            self.caesar_shift_input.raise_()


def _caesar_ui_select_cipher(self, name):
    # Původní select_cipher se nepoužívá přímo, protože při změně šifry
    # spouští auto_encrypt_action dříve, než se přepnou Caesarovy ovládací prvky.
    # Nejprve se nastaví viditelnost prvků a až poté se přepočítá výsledek.
    self.selected_cipher = name
    self.refresh_cipher_styles()
    self.update_selected_header()
    self.update_status()
    self.update_caesar_controls_visibility()
    self.update_output_widget_mode()
    self.auto_encrypt_action()
    self.update_caesar_controls_visibility()


def _caesar_ui_get_shift(self):
    if hasattr(self, "caesar_shift_input"):
        return int(self.caesar_shift_input.value())
    return 3


def _caesar_ui_get_direction(self):
    if hasattr(self, "caesar_direction_combo"):
        text = self.caesar_direction_combo.currentText().strip().lower()
        if "dozadu" in text:
            return "dozadu"
    return "dopredu"


def _caesar_ui_encrypt_selected_cipher(self, text: str) -> str:
    if self.is_caesar_selected():
        caesar_logic = get_cipher_logic("Caesarova šifra")

        if caesar_logic is None:
            return (
                "Chybí soubor s logikou šifry Caesarova šifra:\n"
                "logika_sifer/Caesarova_sifra/caesarova_sifra.py"
            )

        shift = self.get_caesar_shift()
        direction = self.get_caesar_direction()
        signed_shift = -shift if direction == "dozadu" else shift
        return caesar_logic.encrypt(text, signed_shift)

    return _CAESAR_UI_ORIGINAL_ENCRYPT_SELECTED_CIPHER(self, text)


def _caesar_ui_decrypt_selected_cipher(self, text: str) -> str:
    if self.is_caesar_selected():
        caesar_logic = get_cipher_logic("Caesarova šifra")

        if caesar_logic is None:
            return (
                "Chybí soubor s logikou šifry Caesarova šifra:\n"
                "logika_sifer/Caesarova_sifra/caesarova_sifra.py"
            )

        shift = self.get_caesar_shift()
        direction = self.get_caesar_direction()
        signed_shift = -shift if direction == "dozadu" else shift
        return caesar_logic.decrypt(text, signed_shift)

    return _CAESAR_UI_ORIGINAL_DECRYPT_SELECTED_CIPHER(self, text)


SifratorSkinWidget.is_caesar_selected = _caesar_ui_is_selected
SifratorSkinWidget.create_widgets = _caesar_ui_create_widgets
SifratorSkinWidget.update_layout_positions = _caesar_ui_update_layout_positions
SifratorSkinWidget.apply_responsive_fonts = _caesar_ui_apply_responsive_fonts
SifratorSkinWidget.select_cipher = _caesar_ui_select_cipher
SifratorSkinWidget.update_caesar_controls_visibility = _caesar_ui_update_controls_visibility
SifratorSkinWidget.get_caesar_shift = _caesar_ui_get_shift
SifratorSkinWidget.get_caesar_direction = _caesar_ui_get_direction
SifratorSkinWidget.encrypt_selected_cipher = _caesar_ui_encrypt_selected_cipher
SifratorSkinWidget.decrypt_selected_cipher = _caesar_ui_decrypt_selected_cipher


# ============================================================
# ROZŠÍŘENÍ TISKU
# Do rozhraní se doplňuje tlačítko TISK s ikonou icons/tiskarna.png.
# Tlačítko spustí tisk aktuálně zobrazeného výsledku.
# Podporován je textový výstup i grafické výstupy umístěné ve scrollovací oblasti.
# ============================================================

_PRINT_ORIGINAL_CREATE_WIDGETS = SifratorSkinWidget.create_widgets
_PRINT_ORIGINAL_UPDATE_LAYOUT_POSITIONS = SifratorSkinWidget.update_layout_positions
_PRINT_ORIGINAL_APPLY_RESPONSIVE_FONTS = SifratorSkinWidget.apply_responsive_fonts


def _print_create_widgets(self):
    _PRINT_ORIGINAL_CREATE_WIDGETS(self)

    self.print_button = QPushButton("TISK", self)
    self.print_button.setCursor(Qt.PointingHandCursor)
    self.print_button.setToolTip("Vytisknout aktuální výsledek")
    self.print_button.setFocusPolicy(Qt.NoFocus)

    printer_icon_path = self.icon_path("tiskarna.png")
    if printer_icon_path and os.path.exists(printer_icon_path):
        self.print_button.setIcon(QIcon(printer_icon_path))
        self.print_button.setIconSize(QSize(24, 24))

    self.print_button.clicked.connect(self.print_current_result)
    self.print_button.setStyleSheet(f"""
        QPushButton {{
            color: {Colors.GOLD_LIGHT};
            background: rgba(7, 18, 22, 145);
            border: 1px solid rgba(200, 154, 76, 170);
            border-radius: 8px;
            padding-left: 8px;
            padding-right: 8px;
            text-align: center;
        }}
        QPushButton:hover {{
            color: #fff0bd;
            background: rgba(20, 45, 50, 185);
            border: 2px solid {Colors.GOLD_LIGHT};
        }}
        QPushButton:pressed {{
            background: rgba(200, 154, 76, 75);
        }}
    """)
    self.print_button.hide()


def _print_update_layout_positions(self):
    _PRINT_ORIGINAL_UPDATE_LAYOUT_POSITIONS(self)

    if hasattr(self, "print_button"):
        # Tlačítko je umístěné nad výsledkem tak, aby neomezovalo čitelnost výstupu.
        # Pozice se odvozuje od output_text, proto zůstává responzivní při změně velikosti okna.
        out_rect = self.output_text.geometry()
        title_rect = self.result_title.geometry()
        btn_w = max(92, int(122 * self.sx()))
        btn_h = max(28, int(38 * self.sy()))
        x = out_rect.right() - btn_w + 1
        y = title_rect.y() - max(2, int(2 * self.sy()))
        self.print_button.setGeometry(x, y, btn_w, btn_h)
        self.print_button.raise_()
        self.print_button.show()


def _print_apply_responsive_fonts(self):
    _PRINT_ORIGINAL_APPLY_RESPONSIVE_FONTS(self)

    if hasattr(self, "print_button"):
        self.print_button.setFont(QFont("Georgia", self.fs(12), QFont.Bold))
        self.print_button.setIconSize(QSize(max(18, self.fs(24)), max(18, self.fs(24))))


def _print_find_visible_draw_widget(self):
    """Vrátí aktuálně viditelný grafický výstup umístěný ve scrollovací oblasti."""
    for name, widget in self.__dict__.items():
        if name == "scroll_area":
            continue
        if not name.endswith("_scroll"):
            continue
        if isinstance(widget, QScrollArea) and widget.isVisible():
            drawn_widget = widget.widget()
            if drawn_widget is not None:
                return drawn_widget
    return None


def _print_current_result(self):
    """Spustí tisk aktuálně zobrazeného textového nebo grafického výsledku."""
    try:
        from PySide6.QtPrintSupport import QPrinter, QPrintDialog
    except Exception as error:
        QMessageBox.warning(
            self,
            "Tisk není dostupný",
            f"Nepodařilo se načíst podporu tisku:\n{error}",
        )
        return

    has_plain_text = bool(self.output_text.toPlainText().strip())
    drawn_widget = self.print_find_visible_draw_widget()

    if not has_plain_text and drawn_widget is None:
        QMessageBox.information(self, "Není co tisknout", "Nejdřív vytvoř výsledek šifrování nebo dešifrování.")
        return

    printer = QPrinter(QPrinter.HighResolution)
    dialog = QPrintDialog(printer, self)
    dialog.setWindowTitle("Tisk výsledku")

    if dialog.exec() != 1:
        return

    # Textové šifry se tisknou přímo z obsahu QTextEdit.
    if self.output_text.isVisible() and has_plain_text:
        self.output_text.print_(printer)
        return

    # Grafické šifry se tisknou jako rastrový výstup celého vykreslovacího widgetu.
    if drawn_widget is not None:
        pixmap = drawn_widget.grab()
        if pixmap.isNull():
            QMessageBox.warning(self, "Tisk se nezdařil", "Nepodařilo se připravit kreslený výsledek pro tisk.")
            return

        painter = QPainter(printer)
        try:
            try:
                page_rect = printer.pageLayout().paintRectPixels(printer.resolution())
            except Exception:
                try:
                    page_rect = printer.pageRect(QPrinter.DevicePixel)
                except Exception:
                    page_rect = printer.pageRect()

            if hasattr(page_rect, "toRect"):
                page_rect = page_rect.toRect()

            scaled = pixmap.scaled(
                page_rect.size(),
                Qt.KeepAspectRatio,
                Qt.SmoothTransformation,
            )
            x = page_rect.x() + (page_rect.width() - scaled.width()) // 2
            y = page_rect.y()
            painter.drawPixmap(x, y, scaled)
        finally:
            painter.end()


SifratorSkinWidget.create_widgets = _print_create_widgets
SifratorSkinWidget.update_layout_positions = _print_update_layout_positions
SifratorSkinWidget.apply_responsive_fonts = _print_apply_responsive_fonts
SifratorSkinWidget.print_find_visible_draw_widget = _print_find_visible_draw_widget
SifratorSkinWidget.print_current_result = _print_current_result


# ============================================================
# TISKOVÝ DIALOG S VOLBAMI
# Po kliknutí na tlačítko TISK vyskočí okénko se zaškrtávátky:
#   - klíč šifry
#   - text k zašifrování
#   - zašifrovaný / výsledný text
# ============================================================

def _print_options_escape_html(value: str) -> str:
    import html
    return html.escape(value or "").replace("\n", "<br>")


def _print_options_get_key_image_path(self):
    """Najde obrázek klíče šifry, pokud existuje.

    Hledá například:
        icons/klic_caesarova_sifra.png
        icons/Cesarova šifra_klic.png
        icons/caesarova_sifra.png

    Když speciální klíč neexistuje, použije alespoň ikonu vybrané šifry.
    """
    icon_file = self.selected_icon_file() if hasattr(self, "selected_icon_file") else ""
    if not icon_file:
        return ""

    base, ext = os.path.splitext(icon_file)
    candidates = [
        f"klic_{icon_file}",
        f"klíč_{icon_file}",
        f"{base}_klic{ext}",
        f"{base}_klíč{ext}",
        f"{base}_key{ext}",
        icon_file,
    ]

    for file_name in candidates:
        image_path = self.icon_path(file_name)
        if image_path and os.path.exists(image_path):
            return image_path

    return ""




def _print_options_dialog(self, has_input: bool, has_output: bool):
    from PySide6.QtWidgets import QDialog, QVBoxLayout, QLabel, QCheckBox, QDialogButtonBox

    dialog = QDialog(self)
    dialog.setWindowTitle("Co chceš vytisknout?")
    dialog.setModal(True)
    dialog.setMinimumWidth(360)

    layout = QVBoxLayout(dialog)
    title = QLabel("Zaškrtni, co chceš vytisknout:", dialog)
    title.setStyleSheet("font-weight: bold; font-size: 14px;")
    layout.addWidget(title)

    key_check = QCheckBox("Klíč šifry", dialog)
    input_check = QCheckBox("Text k zašifrování", dialog)
    output_check = QCheckBox("Zašifrovaný text", dialog)

    key_check.setChecked(True)
    input_check.setChecked(has_input)
    output_check.setChecked(has_output)

    input_check.setEnabled(has_input)
    output_check.setEnabled(has_output)

    layout.addWidget(key_check)
    layout.addWidget(input_check)
    layout.addWidget(output_check)

    buttons = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel, dialog)
    buttons.button(QDialogButtonBox.Ok).setText("Pokračovat")
    buttons.button(QDialogButtonBox.Cancel).setText("Zrušit")
    buttons.accepted.connect(dialog.accept)
    buttons.rejected.connect(dialog.reject)
    layout.addWidget(buttons)

    if dialog.exec() != QDialog.Accepted:
        return None

    return {
        "key": key_check.isChecked(),
        "input": input_check.isChecked(),
        "output": output_check.isChecked(),
    }




def _print_page_size_points(paper_name: str, orientation_name: str):
    """Převede zvolený formát stránky na typografické body používané při tisku."""
    from PySide6.QtCore import QSizeF

    sizes = {
        "A5": QSizeF(420, 595),
        "A4": QSizeF(595, 842),
        "A3": QSizeF(842, 1191),
        "Letter": QSizeF(612, 792),
        "Legal": QSizeF(612, 1008),
    }
    size = sizes.get(paper_name or "A4", sizes["A4"])
    if orientation_name == "Na šířku":
        return QSizeF(size.height(), size.width())
    return size


def _print_page_setup_label(paper_name: str, orientation_name: str) -> str:
    """Popisek zvoleného papíru pro živý náhled tisku."""
    dimensions_mm = {
        "A5": (148, 210),
        "A4": (210, 297),
        "A3": (297, 420),
        "Letter": (216, 279),
        "Legal": (216, 356),
    }
    width, height = dimensions_mm.get(paper_name or "A4", dimensions_mm["A4"])
    if orientation_name == "Na šířku":
        width, height = height, width
    return f"Papír: {paper_name or 'A4'}, {orientation_name or 'Na výšku'}, {width} x {height} mm"


def _print_apply_page_setup_to_printer(printer, paper_name: str, orientation_name: str):
    """Nastaví QPrinter podle stejného papíru, který vidí uživatel v náhledu."""
    from PySide6.QtGui import QPageSize, QPageLayout

    page_size_map = {
        "A5": QPageSize.A5,
        "A4": QPageSize.A4,
        "A3": QPageSize.A3,
        "Letter": QPageSize.Letter,
        "Legal": QPageSize.Legal,
    }
    printer.setPageSize(QPageSize(page_size_map.get(paper_name or "A4", QPageSize.A4)))
    printer.setPageOrientation(QPageLayout.Landscape if orientation_name == "Na šířku" else QPageLayout.Portrait)
    try:
        printer.setFullPage(True)
    except Exception:
        pass


def _pdf_writer_for_path(path: str, paper_name: str = "A4", orientation_name: str = "Na výšku"):
    from PySide6.QtCore import QMarginsF
    from PySide6.QtGui import QPageLayout, QPageSize, QPdfWriter

    path = os.path.abspath(path)
    folder = os.path.dirname(path)
    if folder:
        os.makedirs(folder, exist_ok=True)

    page_size_map = {
        "A5": QPageSize.A5,
        "A4": QPageSize.A4,
        "A3": QPageSize.A3,
        "Letter": QPageSize.Letter,
        "Legal": QPageSize.Legal,
    }
    writer = QPdfWriter(path)
    writer.setPageSize(QPageSize(page_size_map.get(paper_name or "A4", QPageSize.A4)))
    writer.setPageOrientation(QPageLayout.Landscape if orientation_name == "Na šířku" else QPageLayout.Portrait)
    writer.setPageMargins(QMarginsF(12, 12, 12, 12), QPageLayout.Millimeter)
    return writer


class _PrintPaperPreviewWidget(QWidget):
    """Vykreslí QTextDocument jako skutečné stránky zvoleného papíru."""

    def __init__(self, parent=None):
        super().__init__(parent)
        self._document = None
        self._paper_name = "A4"
        self._orientation_name = "Na výšku"
        self._message = "Připravuji náhled..."
        self._scale = 1.0
        self._manual_scale_percent = None
        self._page_count = 1
        self._page_gap = 28
        self._outer_margin = 28
        self.setMinimumSize(260, 360)
        self.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Expanding)

    def set_message(self, message: str, paper_name: str = "A4", orientation_name: str = "Na výšku"):
        self._document = None
        self._message = message or ""
        self._paper_name = paper_name or "A4"
        self._orientation_name = orientation_name or "Na výšku"
        self._refresh_geometry()

    def set_preview_document(self, document, paper_name: str, orientation_name: str):
        self._document = document
        self._message = ""
        self._paper_name = paper_name or "A4"
        self._orientation_name = orientation_name or "Na výšku"
        self._refresh_geometry()

    def resizeEvent(self, event):
        super().resizeEvent(event)
        self._refresh_geometry()

    def _page_size(self):
        if self._document is not None and self._document.pageSize().isValid():
            return self._document.pageSize()
        return _print_page_size_points(self._paper_name, self._orientation_name)

    def _viewport_width(self) -> int:
        parent = self.parentWidget()
        if parent is not None:
            return max(260, parent.width())
        return max(260, self.width())

    def _calculate_scale(self, page_width: float) -> float:
        if self._manual_scale_percent is not None:
            return max(0.25, min(2.5, float(self._manual_scale_percent) / 100.0))
        available_width = max(220, self._viewport_width() - self._outer_margin * 2)
        base_a4_width = _print_page_size_points("A4", "Na výšku").width()
        return max(0.35, min(1.25, available_width / max(1.0, base_a4_width)))

    def set_zoom_percent(self, percent: int | None):
        if percent is None:
            self._manual_scale_percent = None
        else:
            self._manual_scale_percent = max(25, min(250, int(percent)))
        self._refresh_geometry()

    def zoom_percent(self) -> int:
        if self._manual_scale_percent is not None:
            return int(self._manual_scale_percent)
        return int(round(self._scale * 100))

    def reset_zoom(self):
        self.set_zoom_percent(None)

    def _calculate_page_count(self, page_height: float) -> int:
        if self._document is None:
            return 1

        try:
            count = int(self._document.documentLayout().pageCount())
            if count > 0:
                return count
        except Exception:
            pass

        try:
            doc_height = float(self._document.size().height())
            return max(1, int((doc_height / max(1.0, page_height)) + 0.999))
        except Exception:
            return 1

    def _refresh_geometry(self):
        page_size = self._page_size()
        page_width = max(1.0, page_size.width())
        page_height = max(1.0, page_size.height())
        self._scale = self._calculate_scale(page_width)
        self._page_count = self._calculate_page_count(page_height)

        page_px_w = int(page_width * self._scale)
        page_px_h = int(page_height * self._scale)
        total_w = max(self._viewport_width(), page_px_w + self._outer_margin * 2)
        total_h = self._outer_margin * 2 + self._page_count * page_px_h + max(0, self._page_count - 1) * self._page_gap
        self.setMinimumSize(total_w, total_h)
        self.resize(total_w, total_h)
        self.update()

    def paintEvent(self, event):
        painter = QPainter(self)
        painter.setRenderHint(QPainter.Antialiasing, True)
        painter.fillRect(self.rect(), QColor("#07111f"))

        page_size = self._page_size()
        page_width = max(1.0, page_size.width())
        page_height = max(1.0, page_size.height())
        page_px_w = int(page_width * self._scale)
        page_px_h = int(page_height * self._scale)
        x = max(self._outer_margin, (self.width() - page_px_w) // 2)

        for page_index in range(max(1, self._page_count)):
            y = self._outer_margin + page_index * (page_px_h + self._page_gap)
            shadow_rect = QRect(x + 7, y + 8, page_px_w, page_px_h)
            page_rect = QRect(x, y, page_px_w, page_px_h)

            painter.fillRect(shadow_rect, QColor(0, 0, 0, 95))
            painter.fillRect(page_rect, QColor("#ffffff"))
            painter.setPen(QPen(QColor("#d7d7d7"), 1))
            painter.drawRect(page_rect.adjusted(0, 0, -1, -1))

            if self._document is None:
                if page_index == 0 and self._message:
                    painter.setPen(QColor("#333333"))
                    painter.setFont(QFont("Georgia", 13))
                    painter.drawText(page_rect.adjusted(22, 22, -22, -22), Qt.AlignCenter | Qt.TextWordWrap, self._message)
                continue

            painter.save()
            painter.setClipRect(page_rect)
            painter.translate(page_rect.topLeft())
            painter.scale(self._scale, self._scale)
            painter.translate(0, -page_index * page_height)
            self._document.drawContents(painter, QRectF(0, page_index * page_height, page_width, page_height))
            painter.restore()

        painter.end()



def _print_settings_value(settings: dict, key: str, default):
    if not isinstance(settings, dict):
        return default
    return settings.get(key, default)



def _print_crop_image_to_content(image: QImage, margin: int = 18) -> QImage:
    """Ořízne prázdné okolí grafického výstupu, aby šla měnit jeho velikost.

    Kreslené šifry jsou často vložené ve velkém widgetu. Bez ořezu vypadá
    samotný symbol pořád malý, i když se obrázek zvětšuje. Tady najdeme pixely,
    které nejsou průhledné / bílé, a vytvoříme těsnější obrázek.
    """
    if image.isNull():
        return image

    width = image.width()
    height = image.height()
    if width <= 0 or height <= 0:
        return image

    left = width
    top = height
    right = -1
    bottom = -1

    # Pro rychlost přeskočíme úplně bílé/průhledné pixely.
    for y in range(height):
        for x in range(width):
            color = image.pixelColor(x, y)
            if color.alpha() <= 8:
                continue

            # Bílé nebo skoro bílé pozadí nepočítáme jako kresbu.
            if color.red() > 245 and color.green() > 245 and color.blue() > 245:
                continue

            # Velmi světle šedé pozadí také nepočítáme.
            if color.red() > 235 and color.green() > 235 and color.blue() > 235:
                continue

            left = min(left, x)
            top = min(top, y)
            right = max(right, x)
            bottom = max(bottom, y)

    if right < left or bottom < top:
        return image

    left = max(0, left - margin)
    top = max(0, top - margin)
    right = min(width - 1, right + margin)
    bottom = min(height - 1, bottom + margin)

    return image.copy(left, top, right - left + 1, bottom - top + 1)



def _print_section_start(html_parts, title: str, show_heading: bool, framed: bool, heading_size: int):
    if framed:
        html_parts.append("<div class='section'>")
    else:
        html_parts.append("<div class='section plain-section'>")
    if show_heading and title:
        html_parts.append(f"<h2 style='font-size:{int(heading_size)}pt;'>" + _print_options_escape_html(title) + "</h2>")



def _print_section_end(html_parts):
    html_parts.append("</div>")


def _print_current_key_image_width(self, default: int = 690) -> int:
    """Vrátí šířku obrázku klíče bezpečnou pro aktuální velikost papíru."""
    try:
        value = int(getattr(self, "_print_current_key_image_width", default))
    except Exception:
        value = default
    return max(80, min(int(default), value))




def _print_image_on_white(image: QImage) -> QImage:
    """Vrátí kopii obrázku se světlým pozadím místo průhlednosti.

    QTextDocument někdy při tisku / náhledu tisku zobrazí průhledné pixely černě.
    Proto kreslené výstupy před vložením do tisku podložíme bílým papírem.
    """
    if image.isNull():
        return image

    result = QImage(image.size(), QImage.Format_RGB32)
    result.fill(QColor("#fffdf8"))

    painter = QPainter(result)
    painter.drawImage(0, 0, image)
    painter.end()

    return result


def _print_image_black_on_white(image: QImage) -> QImage:
    """Tiskový převod grafické šifry na čistou černou na bílém papíru.

    Používá se jen pro tisk. UI zůstane barevné/průhledné.
    """
    if image.isNull():
        return image

    source = image.convertToFormat(QImage.Format_ARGB32)
    result = QImage(source.size(), QImage.Format_RGB32)
    result.fill(QColor("#ffffff"))

    for y in range(source.height()):
        for x in range(source.width()):
            color = source.pixelColor(x, y)
            alpha = color.alpha()
            if alpha <= 8:
                continue

            # Ignoruj bílé / skoro bílé pozadí, kdyby ho některý widget přece jen kreslil.
            if alpha > 245 and color.red() > 245 and color.green() > 245 and color.blue() > 245:
                continue

            result.setPixelColor(x, y, QColor(0, 0, 0))

    return result


def _print_options_build_document_base(self, options: dict, drawn_widget, paper_name="A4", orientation_name="Na výšku", settings=None):
    from PySide6.QtCore import QUrl
    from PySide6.QtGui import QTextDocument

    settings = settings or {}

    document = QTextDocument()
    document.setDefaultFont(QFont("Georgia", 11))
    page_size = _print_page_size_points(paper_name, orientation_name)
    document.setPageSize(page_size)

    content_width = max(120.0, page_size.width() - 72.0)

    show_headings = bool(_print_settings_value(settings, "show_headings", True))
    show_frames = bool(_print_settings_value(settings, "show_frames", True))
    heading_size = int(_print_settings_value(settings, "heading_size", 20))
    key_font_size = int(_print_settings_value(settings, "key_font_size", 12))
    key_scale = int(_print_settings_value(settings, "key_scale", 100))
    story_font_size = int(_print_settings_value(settings, "story_font_size", 12))
    input_font_size = int(_print_settings_value(settings, "input_font_size", 13))
    output_font_size = int(_print_settings_value(settings, "output_font_size", 13))
    cipher_scale = int(_print_settings_value(settings, "cipher_scale", 85))
    show_cipher_name = bool(_print_settings_value(settings, "show_cipher_name", False))

    key_title = str(_print_settings_value(settings, "key_title", "Klíč šifry"))
    story_title = str(_print_settings_value(settings, "story_title", "Vlastní text / příběh"))
    input_title = str(_print_settings_value(settings, "input_title", "Text k zašifrování"))
    output_title = str(_print_settings_value(settings, "output_title", "Zašifrovaný text"))
    story_text = str(_print_settings_value(settings, "story_text", ""))

    section_padding = 34.0 if show_frames else 0.0
    printable_width = max(80.0, content_width - section_padding)
    image_width = int(max(0, min(printable_width, printable_width * cipher_scale / 100.0)))
    key_image_width = int(max(80, min(printable_width, printable_width * key_scale / 100.0)))

    html_parts = [
        "<html><head><meta charset='utf-8'>",
        "<style>",
        "body { font-family: Georgia, serif; font-size: 12pt; color: #111; margin: 18px; }",
        "h2 { margin: 0 0 12px 0; color: #10223a; }",
        ".cipher-name { color: #555; font-size: 10pt; text-align: right; margin-bottom: 10px; }",
        ".section { border: 1px solid #777; border-radius: 8px; padding: 14px; margin: 16px 0; }",
        ".plain-section { border: none; padding: 0; margin: 14px 0; }",
        ".textbox { white-space: pre-wrap; border: 1px solid #999; padding: 10px; border-radius: 6px; }",
        ".textbox.no-border { border: none; padding: 0; }",
        ".key-print-box { background: #ffffff; color: #111111; padding: 10px; border-radius: 2px; }",
        ".key-print-box img, .key-image-box { background: #ffffff; }",
        "img { max-width: 100%; height: auto; }",
        "table { border-collapse: collapse; margin-top: 8px; width: 100%; }",
        "td, th { border: 1px solid #666; padding: 6px 8px; }",
        "</style></head>",
        "<body>",
    ]

    # Bez horní hlavičky. Volitelně jen malý název šifry, když si ho uživatel zapne.
    if show_cipher_name and self.selected_cipher:
        html_parts.append(f"<div class='cipher-name'>Šifra: {_print_options_escape_html(self.selected_cipher)}</div>")

    if options.get("key"):
        key_image_path = _print_options_get_key_image_path(self)
        if key_image_path and os.path.exists(key_image_path):
            image = QImage(key_image_path)
            if not image.isNull():
                document.addResource(QTextDocument.ImageResource, QUrl("cipher_key_image"), image)

        _print_section_start(html_parts, key_title, show_headings, show_frames, heading_size)
        html_parts.append(
            "<table width='100%' cellspacing='0' cellpadding='10' "
            "style='background-color:#ffffff; border:0; margin:0;'>"
            "<tr><td style='background-color:#ffffff; border:0;'>"
            f"<div style='font-size:{key_font_size}pt; background-color:#ffffff;'>"
        )
        previous_key_image_width = getattr(self, "_print_current_key_image_width", None)
        self._print_current_key_image_width = key_image_width
        try:
            html_parts.append(_print_options_key_html(self))
        finally:
            if previous_key_image_width is None:
                try:
                    delattr(self, "_print_current_key_image_width")
                except Exception:
                    pass
            else:
                self._print_current_key_image_width = previous_key_image_width
        html_parts.append("</div></td></tr></table>")
        _print_section_end(html_parts)

    if options.get("story"):
        _print_section_start(html_parts, story_title, show_headings, show_frames, heading_size)
        box_class = "textbox" if show_frames else "textbox no-border"
        html_parts.append(f"<div class='{box_class}' style='font-size:{story_font_size}pt;'>")
        html_parts.append(_print_options_escape_html(story_text))
        html_parts.append("</div>")
        _print_section_end(html_parts)

    if options.get("input"):
        input_text = self.input_text.toPlainText() if hasattr(self, "input_text") else ""
        _print_section_start(html_parts, input_title, show_headings, show_frames, heading_size)
        box_class = "textbox" if show_frames else "textbox no-border"
        html_parts.append(f"<div class='{box_class}' style='font-size:{input_font_size}pt;'>")
        html_parts.append(_print_options_escape_html(input_text))
        html_parts.append("</div>")
        _print_section_end(html_parts)

    if options.get("output"):
        _print_section_start(html_parts, output_title, show_headings, show_frames, heading_size)
        if self.output_text.isVisible() and self.output_text.toPlainText().strip():
            output_text = self.output_text.toPlainText()
            box_class = "textbox" if show_frames else "textbox no-border"
            html_parts.append(f"<div class='{box_class}' style='font-size:{output_font_size}pt;'>")
            html_parts.append(_print_options_escape_html(output_text))
            html_parts.append("</div>")
        elif drawn_widget is not None:
            if cipher_scale <= 0:
                # 0 % = kreslený výsledek se do tisku vůbec nevloží.
                pass
            else:
                # NEMRZNOUCÍ NÁHLED: tady už nikdy nevoláme drawn_widget.grab().
                # Grabování velkých kreslených widgetů umí na pár sekund zastavit celé GUI.
                # Náhled použije jen hotovou cache; když není hotová, zobrazí informaci
                # a cache se připraví později přes QTimer.
                image = QImage()
                try:
                    if '_print_cache_peek_drawn_result_image' in globals():
                        image = _print_cache_peek_drawn_result_image(self, drawn_widget)
                    elif '_print_cache_get_drawn_result_image' in globals():
                        image = _print_cache_get_drawn_result_image(self, drawn_widget, force=False)
                except Exception:
                    image = QImage()
                if not image.isNull():
                    document.addResource(QTextDocument.ImageResource, QUrl("drawn_result_image"), image)
                    html_parts.append(f"<p><img src='drawn_result_image' width='{image_width}'></p>")
                else:
                    html_parts.append("<p><i>Kreslený výsledek se připravuje na pozadí. Náhled se po chvíli obnoví.</i></p>")
                    try:
                        self._print_preview_needs_refresh = True
                    except Exception:
                        pass
                    try:
                        if '_schedule_print_cache_preload' in globals():
                            _schedule_print_cache_preload(self, 120)
                    except Exception:
                        pass
        else:
            html_parts.append("<p><i>Není vytvořený žádný výsledek.</i></p>")
        _print_section_end(html_parts)

    html_parts.append("</body></html>")
    document.setHtml("".join(html_parts))
    return document


def _print_current_result_with_options(self):
    """Zobrazí okno s volbami a potom vlastní náhled tisku.

    Windowsové systémové okno tisku neumí u PySide/Qt aplikace vždy zobrazit náhled.
    Proto se nejdřív otevře QPrintPreviewDialog přímo z aplikace. V něm je možné
    výsledek zkontrolovat a potom kliknout na ikonu tiskárny pro skutečný tisk.
    """
    try:
        from PySide6.QtPrintSupport import QPrinter, QPrintPreviewDialog
    except Exception as error:
        QMessageBox.warning(
            self,
            "Tisk není dostupný",
            f"Nepodařilo se načíst podporu tisku:\n{error}",
        )
        return

    input_text = self.input_text.toPlainText().strip() if hasattr(self, "input_text") else ""
    has_input = bool(input_text)
    has_plain_output = bool(self.output_text.toPlainText().strip()) if hasattr(self, "output_text") else False
    drawn_widget = self.print_find_visible_draw_widget() if hasattr(self, "print_find_visible_draw_widget") else None
    has_output = bool(has_plain_output or drawn_widget is not None)

    if not has_input and not has_output:
        QMessageBox.information(self, "Není co tisknout", "Nejdřív zadej text nebo vytvoř výsledek šifrování.")
        return

    options = _print_options_dialog(self, has_input=has_input, has_output=has_output)
    if options is None:
        return

    if not any(options.values()):
        QMessageBox.information(self, "Nic není vybráno", "Zaškrtni alespoň jednu položku k tisku.")
        return

    document = _print_options_build_document(self, options, drawn_widget)

    printer = QPrinter(QPrinter.HighResolution)
    preview = QPrintPreviewDialog(printer, self)
    preview.setWindowTitle("Náhled tisku")
    try:
        screen = self.window().screen() if hasattr(self, "window") and self.window() is not None else QApplication.primaryScreen()
        available = screen.availableGeometry() if screen is not None else None
        if available is not None:
            preview.resize(min(1100, int(available.width() * 0.92)), min(800, int(available.height() * 0.88)))
            preview.setMinimumSize(min(760, preview.width()), min(540, preview.height()))
        else:
            preview.resize(1100, 800)
    except Exception:
        preview.resize(1100, 800)

    def render_preview(target_printer):
        document.print_(target_printer)

    preview.paintRequested.connect(render_preview)
    preview.exec()


SifratorSkinWidget.print_current_result = _print_current_result_with_options
SifratorSkinWidget.print_options_dialog = _print_options_dialog
SifratorSkinWidget.print_options_build_document = _print_options_build_document_base


# ============================================================
# TISK – PŘÍMÝ NÁHLED S VOLBAMI
# Po kliknutí na TISK se rovnou otevře jedno okno, kde jsou:
#   - zaškrtávátka vlevo
#   - živý náhled vpravo
# Náhled se okamžitě mění podle toho, co je zaškrtnuté.
# ============================================================



def _print_current_result_with_live_preview(self):
    """Otevře tematický náhled tisku s možností úpravy rozvržení před tiskem."""
    try:
        from PySide6.QtPrintSupport import QPrinter, QPrintDialog
        from PySide6.QtGui import QPageSize, QPageLayout
        from PySide6.QtWidgets import (
            QDialog,
            QVBoxLayout,
            QHBoxLayout,
            QLabel,
            QCheckBox,
            QTextEdit,
            QPushButton,
            QFrame,
            QComboBox,
            QFormLayout,
            QLineEdit,
            QSpinBox,
            QScrollArea,
        )
    except Exception as error:
        QMessageBox.warning(
            self,
            "Tisk není dostupný",
            f"Nepodařilo se načíst podporu tisku:\n{error}",
        )
        return

    input_text = self.input_text.toPlainText().strip() if hasattr(self, "input_text") else ""
    has_input = bool(input_text)
    has_plain_output = bool(self.output_text.toPlainText().strip()) if hasattr(self, "output_text") else False
    drawn_widget = self.print_find_visible_draw_widget() if hasattr(self, "print_find_visible_draw_widget") else None
    has_output = bool(has_plain_output or drawn_widget is not None)

    if not has_input and not has_output:
        QMessageBox.information(self, "Není co tisknout", "Nejdřív zadej text nebo vytvoř výsledek šifrování.")
        return

    dialog = QDialog(self)
    dialog.setWindowTitle("Náhled tisku – Šifrátor Mraveniště")
    dialog.setModal(True)

    # Responzivní velikost náhledu tisku podle aktuální obrazovky.
    try:
        screen = self.window().screen() if hasattr(self, "window") and self.window() is not None else QApplication.primaryScreen()
        available = screen.availableGeometry() if screen is not None else None
    except Exception:
        available = None

    if available is not None:
        dialog_w = min(1280, max(760, int(available.width() * 0.94)))
        dialog_h = min(850, max(560, int(available.height() * 0.90)))
    else:
        dialog_w, dialog_h = 1280, 850

    compact_print_preview = dialog_w < 1080 or dialog_h < 720
    dialog.resize(dialog_w, dialog_h)
    dialog.setMinimumSize(min(760, dialog_w), min(540, dialog_h))
    dialog.setStyleSheet("""
        QDialog {
            background-color: #0a1626;
            color: #f3ddaa;
        }
        QLabel {
            color: #f3ddaa;
            background: transparent;
        }
        QFrame#sidePanel, QFrame#pageFrame {
            background-color: rgba(10, 20, 34, 230);
            border: 1px solid #9b6b2f;
            border-radius: 12px;
        }
        QFrame#sectionFrame {
            background-color: rgba(11, 31, 52, 200);
            border: 1px solid rgba(179, 130, 55, 160);
            border-radius: 10px;
        }
        QCheckBox {
            spacing: 10px;
            padding: 3px 0;
            font-size: 13px;
        }
        QCheckBox::indicator {
            width: 16px;
            height: 16px;
        }
        QCheckBox::indicator:unchecked {
            border: 1px solid #d4ae67;
            background: #10263e;
        }
        QCheckBox::indicator:checked {
            border: 1px solid #d4ae67;
            background: #0f8aa8;
        }
        QComboBox, QSpinBox, QLineEdit {
            background-color: #10263e;
            color: #f6e7bf;
            border: 1px solid #c49344;
            border-radius: 8px;
            padding: 6px 8px;
            min-height: 20px;
        }
        QSpinBox {
            min-width: 112px;
        }
        QWidget#spinControl {
            background-color: #10263e;
            border: 1px solid #c49344;
            border-radius: 8px;
        }
        QWidget#spinControl QSpinBox {
            background: transparent;
            border: none;
            padding: 6px 8px;
            min-height: 20px;
        }
        QComboBox QAbstractItemView {
            background-color: #10263e;
            color: #f6e7bf;
            border: 1px solid #c49344;
            selection-background-color: #155d75;
        }
        QTextEdit#storyEditor {
            background-color: #071626;
            color: #fff0bd;
            border: 1px solid #c49344;
            border-radius: 8px;
            padding: 8px;
            selection-background-color: #155d75;
        }
        QScrollArea#paperPreviewArea {
            background: #07111f;
            border: 1px solid #c49344;
            border-radius: 8px;
        }
        QScrollArea {
            background: transparent;
            border: none;
        }
        QPushButton {
            background-color: #10263e;
            color: #f6e7bf;
            border: 1px solid #c49344;
            border-radius: 10px;
            padding: 10px 18px;
            font-weight: bold;
            min-height: 22px;
        }
        QPushButton:hover {
            background-color: #144a63;
        }
        QPushButton#cancelButton {
            background-color: #4a1520;
            border-color: #b56a5c;
        }
        QPushButton#cancelButton:hover {
            background-color: #6a1b2c;
        }
        QPushButton#spinArrowButton {
            background-color: #17344f;
            color: #f6e7bf;
            border: none;
            border-left: 1px solid #c49344;
            border-radius: 0px;
            padding: 0px;
            font-size: 10px;
            font-weight: bold;
            min-width: 26px;
            min-height: 0px;
        }
        QPushButton#spinArrowButton:hover {
            background-color: #1e5573;
        }
    """)

    main_layout = QVBoxLayout(dialog)
    main_layout.setContentsMargins(16, 16, 16, 16)
    main_layout.setSpacing(12)

    title = QLabel("Uprav si tisk – náhled se mění hned:", dialog)
    title.setStyleSheet("font-weight: 700; font-size: 18px; color: #f8e8c2;")
    main_layout.addWidget(title)

    content_layout = QVBoxLayout() if compact_print_preview else QHBoxLayout()
    content_layout.setSpacing(10 if compact_print_preview else 14)
    main_layout.addLayout(content_layout, 1)

    side_scroll = QScrollArea(dialog)
    side_scroll.setWidgetResizable(True)
    side_scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
    if compact_print_preview:
        side_scroll.setMinimumWidth(0)
        side_scroll.setMaximumWidth(16777215)
        side_scroll.setMinimumHeight(min(220, max(170, int(dialog_h * 0.28))))
        side_scroll.setMaximumHeight(min(330, max(220, int(dialog_h * 0.38))))
    else:
        side_w_min = min(350, max(300, int(dialog_w * 0.27)))
        side_w_max = min(410, max(330, int(dialog_w * 0.32)))
        side_scroll.setMinimumWidth(side_w_min)
        side_scroll.setMaximumWidth(side_w_max)

    side_panel = QFrame()
    side_panel.setObjectName("sidePanel")
    side_layout = QVBoxLayout(side_panel)
    side_layout.setContentsMargins(14, 14, 14, 14)
    side_layout.setSpacing(12)
    side_scroll.setWidget(side_panel)

    def make_block(title_text):
        frame = QFrame(side_panel)
        frame.setObjectName("sectionFrame")
        layout = QVBoxLayout(frame)
        layout.setContentsMargins(12, 12, 12, 12)
        layout.setSpacing(8)
        label = QLabel(title_text, frame)
        label.setStyleSheet("font-weight: 700; font-size: 15px;")
        layout.addWidget(label)
        return frame, layout

    block_print, block_print_layout = make_block("Co tisknout")
    key_check = QCheckBox("Klíč šifry", block_print)
    story_check = QCheckBox("Vlastní text / příběh", block_print)
    input_check = QCheckBox("Text k zašifrování", block_print)
    output_check = QCheckBox("Zašifrovaný text", block_print)
    key_check.setChecked(True)
    story_check.setChecked(False)
    input_check.setChecked(has_input)
    output_check.setChecked(has_output)
    input_check.setEnabled(has_input)
    output_check.setEnabled(has_output)
    for widget in (key_check, story_check, input_check, output_check):
        block_print_layout.addWidget(widget)
    side_layout.addWidget(block_print)

    block_titles, block_titles_layout = make_block("Nadpisy a názvy")
    show_headings_check = QCheckBox("Zobrazit nadpisy", block_titles)
    show_frames_check = QCheckBox("Zobrazit rámečky sekcí", block_titles)
    show_cipher_name_check = QCheckBox("Zobrazit malý název šifry nahoře", block_titles)
    show_headings_check.setChecked(True)
    show_frames_check.setChecked(True)
    show_cipher_name_check.setChecked(False)
    block_titles_layout.addWidget(show_headings_check)
    block_titles_layout.addWidget(show_frames_check)
    block_titles_layout.addWidget(show_cipher_name_check)

    title_form = QFormLayout()
    title_form.setHorizontalSpacing(8)
    title_form.setVerticalSpacing(8)
    key_title_edit = QLineEdit("Klíč šifry", block_titles)
    story_title_edit = QLineEdit("Vlastní text / příběh", block_titles)
    input_title_edit = QLineEdit("Text k zašifrování", block_titles)
    output_title_edit = QLineEdit("Zašifrovaný text", block_titles)
    title_form.addRow("Klíč:", key_title_edit)
    title_form.addRow("Příběh:", story_title_edit)
    title_form.addRow("Text:", input_title_edit)
    title_form.addRow("Šifra:", output_title_edit)
    block_titles_layout.addLayout(title_form)
    side_layout.addWidget(block_titles)

    block_story, block_story_layout = make_block("Vlastní text / příběh")
    story_edit = QTextEdit(block_story)
    story_edit.setObjectName("storyEditor")
    story_edit.setPlaceholderText("Sem napiš příběh, poznámku nebo instrukce...")
    story_edit.setMinimumHeight(95)
    block_story_layout.addWidget(story_edit)
    side_layout.addWidget(block_story)

    block_page, block_page_layout = make_block("Nastavení stránky")
    page_form = QFormLayout()
    page_form.setHorizontalSpacing(10)
    page_form.setVerticalSpacing(10)
    paper_combo = QComboBox(block_page)
    paper_combo.addItems(["A4", "A5", "A3", "Letter", "Legal"])
    paper_combo.setCurrentText("A4")
    orientation_combo = QComboBox(block_page)
    orientation_combo.addItems(["Na výšku", "Na šířku"])
    orientation_combo.setCurrentText("Na výšku")
    page_form.addRow("Papír:", paper_combo)
    page_form.addRow("Orientace:", orientation_combo)
    block_page_layout.addLayout(page_form)
    side_layout.addWidget(block_page)

    block_sizes, block_sizes_layout = make_block("Velikost a měřítko")
    size_form = QFormLayout()
    size_form.setHorizontalSpacing(8)
    size_form.setVerticalSpacing(8)

    def make_spin(minimum, maximum, value, suffix=" bodů", step=1, tooltip=""):
        container = QWidget(block_sizes)
        container.setObjectName("spinControl")
        container_layout = QHBoxLayout(container)
        container_layout.setContentsMargins(0, 0, 0, 0)
        container_layout.setSpacing(0)

        spin = QSpinBox(container)
        spin.setRange(minimum, maximum)
        spin.setValue(value)
        spin.setSuffix(suffix)
        spin.setSingleStep(step)
        spin.setAccelerated(True)
        spin.setKeyboardTracking(True)
        spin.setCorrectionMode(QAbstractSpinBox.CorrectToNearestValue)
        spin.setButtonSymbols(QAbstractSpinBox.NoButtons)
        spin.setAlignment(Qt.AlignRight)
        spin.setMinimumWidth(118 if len(suffix) <= 7 else 140)
        if tooltip:
            spin.setToolTip(tooltip)
            container.setToolTip(tooltip)
        container_layout.addWidget(spin, 1)

        arrows = QWidget(container)
        arrows_layout = QVBoxLayout(arrows)
        arrows_layout.setContentsMargins(0, 0, 0, 0)
        arrows_layout.setSpacing(0)
        up_button = QPushButton("▲", arrows)
        down_button = QPushButton("▼", arrows)
        for button in (up_button, down_button):
            button.setObjectName("spinArrowButton")
            button.setCursor(Qt.PointingHandCursor)
            button.setFocusPolicy(Qt.NoFocus)
            button.setAutoRepeat(True)
            button.setAutoRepeatDelay(280)
            button.setAutoRepeatInterval(70)
            button.setFixedWidth(28)
        up_button.clicked.connect(lambda _checked=False, s=spin: s.stepUp())
        down_button.clicked.connect(lambda _checked=False, s=spin: s.stepDown())
        arrows_layout.addWidget(up_button)
        arrows_layout.addWidget(down_button)
        container_layout.addWidget(arrows)

        spin._print_control_widget = container
        return spin

    heading_size_spin = make_spin(8, 40, 20, " bodů", 1, "Velikost nadpisů v tisku.")
    key_size_spin = make_spin(30, 100, 100, " % stránky", 5, "Šířka klíče vůči papíru.")
    story_size_spin = make_spin(8, 40, 12, " bodů", 1, "Velikost vlastního textu nebo příběhu.")
    input_size_spin = make_spin(8, 40, 13, " bodů", 1, "Velikost původního textu.")
    output_size_spin = make_spin(8, 60, 13, " bodů", 1, "Velikost zašifrovaného textu.")
    cipher_scale_spin = make_spin(0, 100, 85, " % stránky", 5, "Šířka kreslené šifry vůči papíru.")

    heading_size_label = QLabel("Nadpisy:", block_sizes)
    key_size_label = QLabel("Klíč:", block_sizes)
    story_size_label = QLabel("Příběh:", block_sizes)
    input_size_label = QLabel("Text:", block_sizes)
    output_size_label = QLabel("Zašif. text:", block_sizes)
    cipher_scale_label = QLabel("Kreslená šifra:", block_sizes)

    size_form.addRow(heading_size_label, heading_size_spin._print_control_widget)
    size_form.addRow(key_size_label, key_size_spin._print_control_widget)
    size_form.addRow(story_size_label, story_size_spin._print_control_widget)
    size_form.addRow(input_size_label, input_size_spin._print_control_widget)
    size_form.addRow(output_size_label, output_size_spin._print_control_widget)
    size_form.addRow(cipher_scale_label, cipher_scale_spin._print_control_widget)
    block_sizes_layout.addLayout(size_form)

    text_output_has_size = bool(has_plain_output and hasattr(self, "output_text") and self.output_text.isVisible() and drawn_widget is None)
    drawn_output_has_size = bool(drawn_widget is not None)

    output_size_label.setVisible(text_output_has_size)
    output_size_spin._print_control_widget.setVisible(text_output_has_size)
    cipher_scale_label.setVisible(drawn_output_has_size)
    cipher_scale_spin._print_control_widget.setVisible(drawn_output_has_size)

    hint_size = QLabel("Text se měří v bodech. Klíč a kreslená šifra jsou procento šířky stránky.", block_sizes)
    hint_size.setWordWrap(True)
    hint_size.setStyleSheet("color: #d8c392; font-size: 12px;")
    block_sizes_layout.addWidget(hint_size)
    side_layout.addWidget(block_sizes)

    side_hint = QLabel(
        "Všechno vlevo můžeš upravit. Náhled vpravo ukazuje finální tisk.",
        side_panel,
    )
    side_hint.setWordWrap(True)
    side_hint.setStyleSheet("color: #cdb57e; font-size: 12px;")
    side_layout.addWidget(side_hint)
    side_layout.addStretch(1)

    content_layout.addWidget(side_scroll)

    page_frame = QFrame(dialog)
    page_frame.setObjectName("pageFrame")
    page_layout = QVBoxLayout(page_frame)
    page_layout.setContentsMargins(10 if compact_print_preview else 18, 10 if compact_print_preview else 18, 10 if compact_print_preview else 18, 10 if compact_print_preview else 18)
    page_layout.setSpacing(8 if compact_print_preview else 10)
    preview_header = QLabel("Náhled stránky", page_frame)
    preview_header.setStyleSheet("font-size: 16px; font-weight: 700;")
    page_layout.addWidget(preview_header)
    preview_detail = QLabel(_print_page_setup_label("A4", "Na výšku"), page_frame)
    preview_detail.setStyleSheet("color: #d8c392; font-size: 12px;")
    page_layout.addWidget(preview_detail)
    preview_scroll = QScrollArea(page_frame)
    preview_scroll.setObjectName("paperPreviewArea")
    preview_scroll.setWidgetResizable(False)
    preview_scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarAsNeeded)
    preview_scroll.setVerticalScrollBarPolicy(Qt.ScrollBarAsNeeded)
    preview_scroll.setAlignment(Qt.AlignHCenter | Qt.AlignTop)
    preview = _PrintPaperPreviewWidget(preview_scroll)
    preview_scroll.setWidget(preview)
    page_layout.addWidget(preview_scroll, 1)
    content_layout.addWidget(page_frame, 1)

    buttons_layout = QHBoxLayout()
    buttons_layout.addStretch(1)
    print_button = QPushButton("Tisknout", dialog)
    cancel_button = QPushButton("Zrušit", dialog)
    cancel_button.setObjectName("cancelButton")
    print_button.setMinimumWidth(150)
    cancel_button.setMinimumWidth(120)
    buttons_layout.addWidget(print_button)
    buttons_layout.addWidget(cancel_button)
    main_layout.addLayout(buttons_layout)

    def current_options():
        return {
            "key": key_check.isChecked(),
            "story": story_check.isChecked(),
            "input": input_check.isChecked(),
            "output": output_check.isChecked(),
        }

    def current_page_setup():
        return paper_combo.currentText(), orientation_combo.currentText()

    def current_settings():
        # DŮLEŽITÉ: když uživatel napíše číslo ručně, QSpinBox ho nemusí
        # hned převést na hodnotu. interpretText() ho vynutí ještě před
        # sestavením náhledu i před tiskem.
        for spin in [heading_size_spin, key_size_spin, story_size_spin, input_size_spin, output_size_spin, cipher_scale_spin]:
            try:
                spin.interpretText()
            except Exception:
                pass

        return {
            "show_headings": show_headings_check.isChecked(),
            "show_frames": show_frames_check.isChecked(),
            "show_cipher_name": show_cipher_name_check.isChecked(),
            "key_title": key_title_edit.text().strip(),
            "story_title": story_title_edit.text().strip(),
            "input_title": input_title_edit.text().strip(),
            "output_title": output_title_edit.text().strip(),
            "story_text": story_edit.toPlainText(),
            "heading_size": heading_size_spin.value(),
            "key_font_size": 12,
            "key_scale": key_size_spin.value(),
            "story_font_size": story_size_spin.value(),
            "input_font_size": input_size_spin.value(),
            "output_font_size": output_size_spin.value(),
            "cipher_scale": cipher_scale_spin.value(),
        }

    def update_preview():
        try:
            self._print_preview_needs_refresh = False
        except Exception:
            pass
        options = current_options()
        print_button.setEnabled(any(options.values()))
        heading_size_spin._print_control_widget.setEnabled(show_headings_check.isChecked())
        key_size_spin._print_control_widget.setEnabled(bool(options.get("key")))
        story_size_spin._print_control_widget.setEnabled(bool(options.get("story")))
        input_size_spin._print_control_widget.setEnabled(bool(options.get("input")))
        output_size_spin._print_control_widget.setEnabled(bool(options.get("output") and text_output_has_size))
        cipher_scale_spin._print_control_widget.setEnabled(bool(options.get("output") and drawn_output_has_size))
        paper_name, orientation_name = current_page_setup()
        preview_detail.setText(_print_page_setup_label(paper_name, orientation_name))

        if not any(options.values()):
            preview.set_message("Zaškrtni alespoň jednu položku vlevo.", paper_name, orientation_name)
            return

        document = _print_options_build_document(
            self,
            options,
            drawn_widget,
            paper_name=paper_name,
            orientation_name=orientation_name,
            settings=current_settings(),
        )
        preview.set_preview_document(document, paper_name, orientation_name)
        try:
            if getattr(self, "_print_preview_needs_refresh", False):
                preview_timer.start(1600)
        except Exception:
            pass

    # Náhled má reagovat hned. Krátký timer jen spojí rychlé psaní
    # do jednoho překreslení, aby dialog při držení klávesy necukal.
    preview_timer = QTimer(dialog)
    preview_timer.setSingleShot(True)
    preview_timer.timeout.connect(update_preview)

    def schedule_update_preview(*_args, delay_ms: int = 70):
        preview_timer.start(max(0, int(delay_ms)))

    def update_preview_now(*_args):
        preview_timer.stop()
        update_preview()

    widgets_to_update = [
        key_check,
        story_check,
        input_check,
        output_check,
        show_headings_check,
        show_frames_check,
        show_cipher_name_check,
    ]
    for widget in widgets_to_update:
        widget.toggled.connect(update_preview_now)

    for edit in [key_title_edit, story_title_edit, input_title_edit, output_title_edit]:
        edit.textChanged.connect(schedule_update_preview)

    story_edit.textChanged.connect(schedule_update_preview)
    paper_combo.currentTextChanged.connect(update_preview_now)
    orientation_combo.currentTextChanged.connect(update_preview_now)

    for spin in [heading_size_spin, key_size_spin, story_size_spin, input_size_spin, output_size_spin, cipher_scale_spin]:
        spin.valueChanged.connect(update_preview_now)
        try:
            spin.lineEdit().textEdited.connect(lambda _text: schedule_update_preview(delay_ms=90))
            spin.editingFinished.connect(lambda s=spin: (s.interpretText(), update_preview_now()))
        except Exception:
            pass

    cancel_button.clicked.connect(dialog.reject)

    def do_print():
        options = current_options()
        if not any(options.values()):
            QMessageBox.information(dialog, "Nic není vybráno", "Zaškrtni alespoň jednu položku k tisku.")
            return

        paper_name, orientation_name = current_page_setup()
        document = _print_options_build_document(
            self,
            options,
            drawn_widget,
            paper_name=paper_name,
            orientation_name=orientation_name,
            settings=current_settings(),
        )

        printer = QPrinter(QPrinter.HighResolution)
        _print_apply_page_setup_to_printer(printer, paper_name, orientation_name)

        print_dialog = QPrintDialog(printer, dialog)
        print_dialog.setWindowTitle("Tisk")
        if print_dialog.exec() != QDialog.Accepted:
            return

        try:
            _sifrator_log_from_widget(
                self,
                f"Tisk potvrzen: šifra={self.selected_cipher or 'žádná'}, "
                f"papír={paper_name}, orientace={orientation_name}, volby={options}",
            )
        except Exception:
            pass
        document.print_(printer)
        dialog.accept()

    print_button.clicked.connect(do_print)
    preview.set_message("Připravuji náhled...", "A4", "Na výšku")
    update_preview_now()
    dialog.exec()


# Přepíše staré chování: už se neotevře nejdřív malé okno se zaškrtávátky,
# ale rovnou velký dialog s volbami a živým náhledem.
SifratorSkinWidget.print_current_result = _print_current_result_with_live_preview



# ============================================================
# OPRAVA BINÁRNÍCH ČTVERCŮ + KLÍČE V TISKU
# - Binární čtverce se zobrazují jako skutečná kreslená mřížka,
#   přesněji jako na vzoru: písmena ve slově na sebe navazují.
# - Tlačítko klíče je jen "KLÍČ" a je menší, aby se nekřížilo s TISK.
# - Do tisku se propisuje skutečný generovaný klíč z get_key_data().
# ============================================================

def _fixed_get_logic_module_for_print_key(self):
    """Vrátí modul logiky shodný s tím, který používá generování klíče šifry."""
    if hasattr(self, "get_selected_logic_module"):
        try:
            return self.get_selected_logic_module()
        except Exception as error:
            print(f"CHYBA při načítání logiky pro tisk klíče: {error}")
    return None


def _fixed_key_image_file_url(path: str) -> str:
    """Převede cestu k PNG souboru na URL vhodnou pro vložení do QTextDocument."""
    try:
        from pathlib import Path as _Path
        return _Path(path).resolve().as_uri()
    except Exception:
        return "file:///" + str(path).replace("\\", "/")




def get_binary_squares_widget_class():
    module = get_cipher_logic("Binární čtverce")
    if module is None:
        return None
    return getattr(module, "BinarySquaresOutputWidget", None)


def _binary_squares_create_canvas(self):
    widget_class = get_binary_squares_widget_class()
    if widget_class is not None:
        return widget_class()

    fallback = QLabel("Chybí kreslicí modul Binární čtverce")
    fallback.setStyleSheet("background: transparent; color: #f3d79a;")
    return fallback


def _binary_squares_resize_canvas_to_content(self, text: str | None = None):
    if not hasattr(self, "binary_squares_scroll") or not hasattr(self, "binary_squares_canvas"):
        return

    viewport_w = max(200, self.binary_squares_scroll.viewport().width())
    if hasattr(self.binary_squares_canvas, "set_available_width"):
        self.binary_squares_canvas.set_available_width(viewport_w)

    if text is None:
        text = self.get_input_text() if hasattr(self, "get_input_text") else ""

    if hasattr(self.binary_squares_canvas, "set_plain_text"):
        self.binary_squares_canvas.set_plain_text(text)
    elif hasattr(self.binary_squares_canvas, "set_cipher_text"):
        self.binary_squares_canvas.set_cipher_text(text)
    elif hasattr(self.binary_squares_canvas, "setText"):
        self.binary_squares_canvas.setText(text)

    hint = self.binary_squares_canvas.sizeHint()
    self.binary_squares_canvas.resize(max(viewport_w, hint.width()), max(hint.height(), self.binary_squares_scroll.viewport().height()))
    self.binary_squares_canvas.update()


def _binary_squares_hide_all_draw_outputs(self, keep_name: str | None = None):
    for name, widget in list(self.__dict__.items()):
        if name == keep_name:
            continue
        if name == "scroll_area":
            continue
        if name.endswith("_scroll") and isinstance(widget, QScrollArea):
            widget.hide()


_BINARY_ORIGINAL_CREATE_WIDGETS = SifratorSkinWidget.create_widgets
_BINARY_ORIGINAL_UPDATE_LAYOUT_POSITIONS = SifratorSkinWidget.update_layout_positions
_BINARY_ORIGINAL_APPLY_RESPONSIVE_FONTS = SifratorSkinWidget.apply_responsive_fonts
_BINARY_ORIGINAL_UPDATE_OUTPUT_WIDGET_MODE = SifratorSkinWidget.update_output_widget_mode
_BINARY_ORIGINAL_SET_RESULT_OUTPUT = SifratorSkinWidget.set_result_output


def _binary_squares_create_widgets(self):
    _BINARY_ORIGINAL_CREATE_WIDGETS(self)

    # Vlastní grafický výstup pro Binární čtverce.
    self.binary_squares_canvas = _binary_squares_create_canvas(self)
    self.binary_squares_scroll = QScrollArea(self)
    self.binary_squares_scroll.setWidgetResizable(False)
    self.binary_squares_scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
    self.binary_squares_scroll.setVerticalScrollBarPolicy(Qt.ScrollBarAsNeeded)
    self.binary_squares_scroll.setFrameShape(QScrollArea.NoFrame)
    self.binary_squares_scroll.setWidget(self.binary_squares_canvas)
    self.binary_squares_scroll.setStyleSheet("""
        QScrollArea {
            background: #fffdf8;
            border: none;
        }
        QScrollArea > QWidget > QWidget {
            background: #fffdf8;
        }
        QScrollBar:vertical {
            background: rgba(20, 17, 12, 130);
            width: 11px;
            margin: 4px 2px 4px 2px;
            border-radius: 5px;
        }
        QScrollBar::handle:vertical {
            background: #b89b68;
            min-height: 42px;
            border-radius: 5px;
        }
        QScrollBar::add-line:vertical,
        QScrollBar::sub-line:vertical {
            height: 0px;
        }
    """)
    self.binary_squares_scroll.hide()

    # Text tlačítka klíče zkrátíme hned po vytvoření.
    if hasattr(self, "key_button"):
        self.key_button.setText("KLÍČ")
        self.key_button.setToolTip("Zobrazit klíč vybrané šifry")
        self.key_button.setMinimumWidth(0)


def _binary_squares_update_layout_positions(self):
    _BINARY_ORIGINAL_UPDATE_LAYOUT_POSITIONS(self)

    if hasattr(self, "binary_squares_scroll"):
        self.binary_squares_scroll.setGeometry(self.output_text.geometry())
        self.binary_squares_scroll.raise_()
        _binary_squares_resize_canvas_to_content(self)

    # Oprava překrývání KLÍČ / TISK. Obě tlačítka držíme na pravé straně nad výsledkem.
    if hasattr(self, "key_button"):
        out_rect = self.output_text.geometry()
        title_rect = self.result_title.geometry()
        key_w = max(72, int(92 * self.sx()))
        btn_h = max(26, int(36 * self.sy()))
        gap = max(7, int(10 * self.sx()))

        if hasattr(self, "print_button"):
            print_w = max(86, int(112 * self.sx()))
            print_x = out_rect.right() - print_w + 1
            y = title_rect.y() - max(2, int(2 * self.sy()))
            self.print_button.setGeometry(print_x, y, print_w, btn_h)
            self.print_button.raise_()
            self.print_button.show()

            key_x = max(out_rect.x(), print_x - gap - key_w)
            self.key_button.setGeometry(key_x, y, key_w, btn_h)
        else:
            self.key_button.setGeometry(out_rect.right() - key_w + 1, title_rect.y(), key_w, btn_h)

        self.key_button.raise_()
        self.key_button.show()


def _binary_squares_apply_responsive_fonts(self):
    _BINARY_ORIGINAL_APPLY_RESPONSIVE_FONTS(self)
    if hasattr(self, "key_button"):
        self.key_button.setFont(QFont("Georgia", self.fs(12), QFont.Bold))
    if hasattr(self, "print_button"):
        self.print_button.setFont(QFont("Georgia", self.fs(12), QFont.Bold))
        self.print_button.setIconSize(QSize(max(16, self.fs(21)), max(16, self.fs(21))))


def _binary_squares_update_output_widget_mode(self):
    _BINARY_ORIGINAL_UPDATE_OUTPUT_WIDGET_MODE(self)

    if self.is_binary_squares_selected() and self.result_mode == "encrypt" and hasattr(self, "binary_squares_scroll"):
        self.output_text.hide()
        _binary_squares_hide_all_draw_outputs(self, "binary_squares_scroll")
        self.binary_squares_scroll.show()
        self.binary_squares_scroll.raise_()
        _binary_squares_resize_canvas_to_content(self)


def _binary_squares_set_result_output(self, result: str):
    if self.is_binary_squares_selected() and self.result_mode == "encrypt" and hasattr(self, "binary_squares_scroll"):
        self.output_text.clear()
        self.output_text.hide()
        _binary_squares_hide_all_draw_outputs(self, "binary_squares_scroll")
        self.binary_squares_scroll.show()
        self.binary_squares_scroll.raise_()
        self.binary_squares_scroll.verticalScrollBar().setValue(0)
        _binary_squares_resize_canvas_to_content(self, self.get_input_text())
        QTimer.singleShot(0, lambda: _binary_squares_resize_canvas_to_content(self, self.get_input_text()))
        return

    return _BINARY_ORIGINAL_SET_RESULT_OUTPUT(self, result)


SifratorSkinWidget.create_widgets = _binary_squares_create_widgets
SifratorSkinWidget.update_layout_positions = _binary_squares_update_layout_positions
SifratorSkinWidget.apply_responsive_fonts = _binary_squares_apply_responsive_fonts
SifratorSkinWidget.update_output_widget_mode = _binary_squares_update_output_widget_mode
SifratorSkinWidget.set_result_output = _binary_squares_set_result_output
SifratorSkinWidget.resize_binary_squares_canvas_to_content = _binary_squares_resize_canvas_to_content


def set_windows_app_id():
    """Nastaví vlastní Windows AppUserModelID pro správné zobrazení ikony aplikace."""
    if sys.platform.startswith("win"):
        try:
            ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID(
                "Komarek.SifratorMraveniste.PiratiZKaribiku"
            )
        except Exception:
            pass


def make_app_icon_from_logo(icons_path):
    """Vytvoří icons/app_icon.ico z icons/logo.png a vrátí cestu k ikoně.

    Windows pro ikonu okna / hlavního panelu nejlépe pracuje s .ico souborem.
    Pokud app_icon.ico už existuje, použije se rovnou.
    """
    logo_png = os.path.join(icons_path, "logo.png")
    logo_png_alt = os.path.join(icons_path, "Logo.png")
    ico_path = os.path.join(icons_path, "app_icon.ico")

    if os.path.exists(ico_path):
        return ico_path

    if not os.path.exists(logo_png) and os.path.exists(logo_png_alt):
        logo_png = logo_png_alt

    if not os.path.exists(logo_png):
        return ""

    if Image is None:
        return logo_png

    try:
        img = Image.open(logo_png).convert("RGBA")

        bbox = img.getchannel("A").getbbox()
        if bbox:
            img = img.crop(bbox)

        size = max(img.size)
        canvas = Image.new("RGBA", (size, size), (0, 0, 0, 0))
        canvas.paste(img, ((size - img.width) // 2, (size - img.height) // 2), img)

        canvas.save(
            ico_path,
            format="ICO",
            sizes=[
                (16, 16),
                (24, 24),
                (32, 32),
                (48, 48),
                (64, 64),
                (128, 128),
                (256, 256),
            ],
        )
        return ico_path

    except Exception:
        return logo_png


class SifratorWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle(f"TÁBOROVÁ PALUBA MRAVENIŠTĚ - PIRÁTI Z KARIBIKU v{APP_VERSION}")

        # Responzivní obal pro celé pirátské UI.
        # Důvod: samotný skin má poměr 1672 × 941 a na menších MacBoocích / noteboocích
        # se původně okno nevešlo na obrazovku. ScrollArea nechá zachovat vzhled, ale
        # zároveň dovolí použití i na 800 × 600 bez oříznuté pravé části.
        self.central = SifratorSkinWidget()
        self.central.setSizePolicy(QSizePolicy.Fixed, QSizePolicy.Fixed)

        self.ui_scroll_area = QScrollArea(self)
        self.ui_scroll_area.setWidget(self.central)
        self.ui_scroll_area.setWidgetResizable(False)
        self.ui_scroll_area.setFrameShape(QScrollArea.NoFrame)
        self.ui_scroll_area.setAlignment(Qt.AlignCenter)
        self.ui_scroll_area.setHorizontalScrollBarPolicy(Qt.ScrollBarAsNeeded)
        self.ui_scroll_area.setVerticalScrollBarPolicy(Qt.ScrollBarAsNeeded)
        self.ui_scroll_area.setStyleSheet("""
            QScrollArea {
                background: #06131b;
                border: none;
            }
            QScrollBar:vertical, QScrollBar:horizontal {
                background: rgba(20, 17, 12, 160);
                border-radius: 5px;
                margin: 2px;
            }
            QScrollBar:vertical { width: 12px; }
            QScrollBar:horizontal { height: 12px; }
            QScrollBar::handle:vertical, QScrollBar::handle:horizontal {
                background: #b89b68;
                border-radius: 5px;
                min-height: 34px;
                min-width: 34px;
            }
            QScrollBar::add-line, QScrollBar::sub-line {
                width: 0px;
                height: 0px;
            }
        """)
        # Aplikace se otevírá do samostatné domovské obrazovky. Původní
        # šifrátor zůstává beze změny na druhé stránce zásobníku.
        self.home_page = PirateHomeWidget(self.central.icons_path, APP_VERSION, self)
        self.home_page.navigate_requested.connect(self._open_home_destination)

        self.page_stack = QStackedWidget(self)
        self.page_stack.setStyleSheet("background: #06131b;")
        self.page_stack.addWidget(self.home_page)
        self.page_stack.addWidget(self.ui_scroll_area)

        # Společná horní navigace pro všechny stránky otevřené z dlaždic.
        # Domovská obrazovka ji nepotřebuje, v obsahu nahrazuje samostatná okna.
        self.app_shell = QWidget(self)
        self._content_entry_animation = None
        self._content_entry_effect = None
        shell_layout = QVBoxLayout(self.app_shell)
        shell_layout.setContentsMargins(0, 0, 0, 0)
        shell_layout.setSpacing(0)

        self.content_navbar = QWidget(self.app_shell)
        self.content_navbar.setFixedHeight(54)
        self.content_navbar.setStyleSheet("""
            QWidget {
                background-color: #061923;
                border-bottom: 1px solid #8d6935;
            }
        """)
        nav_layout = QHBoxLayout(self.content_navbar)
        nav_layout.setContentsMargins(12, 6, 18, 6)
        nav_layout.setSpacing(14)

        self.home_button = QPushButton("VELITELSKÁ PALUBA", self.content_navbar)
        self.home_button.setCursor(Qt.PointingHandCursor)
        self.home_button.setFixedHeight(40)
        self.home_button.setIcon(QIcon(os.path.join(self.central.icons_path, "anchor.png")))
        self.home_button.setIconSize(QSize(25, 25))
        self.home_button.setToolTip("Vrátit se na velitelskou palubu")
        self.home_button.setStyleSheet("""
            QPushButton {
                color: #f3d79a;
                background-color: rgba(5, 25, 35, 235);
                border: 1px solid #c89a4c;
                border-radius: 10px;
                padding: 5px 11px;
                font-family: Georgia;
                font-size: 12px;
                font-weight: bold;
            }
            QPushButton:hover {
                color: #fff0bd;
                background-color: rgba(13, 63, 72, 245);
                border: 2px solid #f3d79a;
            }
            QPushButton:pressed {
                background-color: rgba(8, 42, 52, 250);
            }
        """)
        self.home_button.clicked.connect(self.show_home_page)
        nav_layout.addWidget(self.home_button)

        self.content_title = QLabel(self.content_navbar)
        self.content_title.setAlignment(Qt.AlignVCenter | Qt.AlignLeft)
        self.content_title.setStyleSheet(
            "color: #f3d79a; background: transparent; border: none; "
            "font-family: Georgia; font-size: 16px; font-weight: bold; letter-spacing: 1px;"
        )
        nav_layout.addWidget(self.content_title, 1)

        shell_layout.addWidget(self.content_navbar)
        shell_layout.addWidget(self.page_stack, 1)
        self.setCentralWidget(self.app_shell)
        self.content_navbar.hide()

        self.setMinimumSize(800, 600)
        self.resize(*self._initial_window_size())
        self._apply_responsive_canvas_size()

        # Ikona celé aplikace z icons/logo.png / icons/app_icon.ico.
        app_icon_path = make_app_icon_from_logo(self.central.icons_path)
        if app_icon_path and os.path.exists(app_icon_path):
            app_icon = QIcon(app_icon_path)
            self.setWindowIcon(app_icon)
            app = QApplication.instance()
            if app is not None:
                app.setWindowIcon(app_icon)

        # Na malých obrazovkách se aplikace otevře rovnou maximalizovaná, aby se
        # využila celá dostupná plocha mimo horní lištu a Dock.
        QTimer.singleShot(0, self._maximize_on_small_screen_if_needed)

        # Kontrola aktualizací až po zobrazení okna.
        # Když není internet nebo není novější verze, nic nevyskočí.
        QTimer.singleShot(1500, self.check_updates_after_start)

    def show_home_page(self):
        self.content_navbar.hide()
        self.page_stack.setCurrentWidget(self.home_page)
        self.home_page.show_main_menu(animated=True)
        self.setWindowTitle(f"TÁBOROVÁ PALUBA MRAVENIŠTĚ - PIRÁTI Z KARIBIKU v{APP_VERSION}")

    def _show_content_page(self, page, title: str):
        self.content_title.setText(str(title or "").upper())
        self.content_navbar.show()
        self.page_stack.setCurrentWidget(page)
        self.setWindowTitle(f"{title} - TÁBOROVÁ PALUBA MRAVENIŠTĚ v{APP_VERSION}")
        self._animate_content_page_entry()

    def _animate_content_page_entry(self):
        """Plynule zobrazi celou cilovou obrazovku po zmizeni dlazdic menu."""
        if self._content_entry_animation is not None:
            self._content_entry_animation.stop()
            self._content_entry_animation = None
        self.app_shell.setGraphicsEffect(None)

        effect = QGraphicsOpacityEffect(self.app_shell)
        effect.setOpacity(0.0)
        self.app_shell.setGraphicsEffect(effect)

        animation = QPropertyAnimation(effect, b"opacity", self)
        animation.setDuration(360)
        animation.setStartValue(0.0)
        animation.setEndValue(1.0)
        animation.setEasingCurve(QEasingCurve.OutCubic)
        self._content_entry_effect = effect
        self._content_entry_animation = animation

        def finish_entry():
            if self._content_entry_animation is not animation:
                return
            self.app_shell.setGraphicsEffect(None)
            self._content_entry_effect = None
            self._content_entry_animation = None

        animation.finished.connect(finish_entry)
        animation.start()

    def show_cipher_page(self):
        self._show_content_page(self.ui_scroll_area, "Šifrátor")
        QTimer.singleShot(0, self._apply_responsive_canvas_size)

    def _show_embedded_dialog_page(self, dialog, title: str):
        page = getattr(dialog, "_embedded_page", None)
        if page is None:
            dialog.setWindowFlags(Qt.Widget)
            dialog.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Expanding)
            page = QScrollArea(self.page_stack)
            page.setWidgetResizable(True)
            page.setFrameShape(QScrollArea.NoFrame)
            page.setHorizontalScrollBarPolicy(Qt.ScrollBarAsNeeded)
            page.setVerticalScrollBarPolicy(Qt.ScrollBarAsNeeded)
            page.setStyleSheet("""
                QScrollArea {
                    background: #06131b;
                    border: none;
                }
                QScrollBar:vertical, QScrollBar:horizontal {
                    background: rgba(20, 17, 12, 160);
                    border-radius: 5px;
                    margin: 2px;
                }
                QScrollBar:vertical { width: 12px; }
                QScrollBar:horizontal { height: 12px; }
                QScrollBar::handle:vertical, QScrollBar::handle:horizontal {
                    background: #b89b68;
                    border-radius: 5px;
                    min-height: 34px;
                    min-width: 34px;
                }
            """)
            page.setWidget(dialog)
            dialog._embedded_page = page
            self.page_stack.addWidget(page)
            dialog.finished.connect(
                lambda _result, embedded_page=page: (
                    self.show_home_page() if self.page_stack.currentWidget() is embedded_page else None
                )
            )
        dialog.show()
        self._show_content_page(page, title)

    def _open_home_destination(self, route: str):
        if route == "cipher":
            self.show_cipher_page()
            return
        if route == "planner":
            self.show_camp_planner_window()
            return
        if route == "batch":
            self.show_batch_encrypt_window()
            return
        if route == "overview":
            self.show_cipher_overview_window()
            return
        if route == "history":
            self.show_history_window()
            return
        if route == "backup":
            self.show_user_data_backup_window()
            return
        if route == "sports":
            self.show_sports_day_window()
            return
        if route == "groups":
            self.show_groups_window()
            return
        if route == "diploma":
            self.show_diploma_window()
            return
        if route == "diploma_sports":
            self.show_diploma_window("sports")
            return
        if route == "diploma_camp":
            self.show_diploma_window("camp")
            return
        if route == "diploma_cleaning":
            self.show_diploma_window("cleaning")
            return
        if route == "diploma_cleaning_award":
            self.show_diploma_window("cleaning_award")
            return
        if route == "diploma_daily":
            self.show_diploma_window("daily")
            return
        if route == "diploma_meal":
            self.show_diploma_window("meal")

    def show_sports_day_window(self):
        dialog = getattr(self, "_sports_day_dialog", None)
        if dialog is None:
            dialog = SportsDayDialog(self, self.central.icons_path)
            self._sports_day_dialog = dialog
        else:
            dialog.refresh_all()
        self._show_embedded_dialog_page(dialog, "Sportovní den")

    def show_groups_window(self):
        dialog = getattr(self, "_groups_dialog", None)
        if dialog is None:
            dialog = GroupsDialog(self, self.central.icons_path)
            self._groups_dialog = dialog
        else:
            dialog.refresh_data()
        self._show_embedded_dialog_page(dialog, "Oddíly/Ubytování")

    def show_diploma_window(self, diploma_kind=None):
        dialog = getattr(self, "_diploma_dialog", None)
        if dialog is None:
            dialog = DiplomaDialog(self, self.central.icons_path)
            self._diploma_dialog = dialog
        if diploma_kind == "sports":
            dialog.show_sports_diploma()
        elif diploma_kind == "camp":
            dialog.show_camp_diploma()
        elif diploma_kind == "cleaning":
            dialog.show_cleaning_sheet()
        elif diploma_kind == "cleaning_award":
            dialog.show_cleaning_award()
        elif diploma_kind == "daily":
            dialog.show_daily_program()
        elif diploma_kind == "meal":
            dialog.show_meal_plan()
        else:
            dialog.show_choices()
        self._show_embedded_dialog_page(dialog, "Diplom")

    def _available_screen_geometry(self):
        app = QApplication.instance()
        screen = app.primaryScreen() if app is not None else None
        if screen is None:
            return QRect(0, 0, BASE_W, BASE_H)
        return screen.availableGeometry()

    def _initial_window_size(self):
        available = self._available_screen_geometry()
        # Okno se na startu nikdy nesnaží být větší než dostupná plocha obrazovky.
        # Na 4K a větších monitorech zůstane rozumně velké, aby UI nebylo obří.
        target_w = min(BASE_W, max(800, int(available.width() * 0.96)))
        target_h = min(BASE_H, max(600, int(available.height() * 0.92)))
        return target_w, target_h

    def _responsive_canvas_size(self):
        viewport = self.ui_scroll_area.viewport().size() if hasattr(self, "ui_scroll_area") else self.size()
        viewport_w = max(1, viewport.width())
        viewport_h = max(1, viewport.height())

        # Malý režim: UI se drží čitelnější minimální pracovní plochy a okno dostane scroll.
        if viewport_w < 1000 or viewport_h < 620:
            return QSize(1050, 650)

        # Kompaktní notebookový režim: vejde se na menší MacBooky / HD displeje bez oříznutí.
        if viewport_w < 1300 or viewport_h < 760:
            return QSize(max(1050, viewport_w), max(650, viewport_h))

        # Normální režim: UI se přizpůsobí oknu až do původního návrhu.
        if viewport_w < BASE_W or viewport_h < BASE_H:
            return QSize(max(1200, viewport_w), max(675, viewport_h))

        # Velké a 4K monitory: držíme původní rozumný návrh uprostřed obrazovky.
        return QSize(BASE_W, BASE_H)

    def _apply_responsive_canvas_size(self):
        if not hasattr(self, "central"):
            return
        size = self._responsive_canvas_size()
        if self.central.size() != size:
            self.central.setFixedSize(size)
        self.central.update_layout_positions()
        self.central.update()

    def _maximize_on_small_screen_if_needed(self):
        available = self._available_screen_geometry()
        if available.width() < 1300 or available.height() < 760:
            self.showMaximized()
        self._apply_responsive_canvas_size()

    def resizeEvent(self, event):
        super().resizeEvent(event)
        if hasattr(self, "page_stack") and self.page_stack.currentWidget() is self.ui_scroll_area:
            QTimer.singleShot(0, self._apply_responsive_canvas_size)

    def check_updates_after_start(self):
        update_data = update_manager.check_for_update(APP_VERSION)

        if not update_data:
            return

        remote_version = update_data.get("version", "")
        notes = update_data.get("notes", "")

        message = (
            f"Je dostupná nová verze {remote_version}.\n\n"
            f"Aktuální verze: {APP_VERSION}\n\n"
            f"{notes}\n\n"
            "Chceš aplikaci aktualizovat?"
        )

        answer = QMessageBox.question(
            self,
            "Dostupná aktualizace",
            message,
            QMessageBox.Yes | QMessageBox.No,
            QMessageBox.Yes,
        )

        if answer == QMessageBox.Yes:
            try:
                update_manager.download_and_install_update(update_data)
            except Exception as error:
                QMessageBox.critical(
                    self,
                    "Chyba aktualizace",
                    f"Aktualizaci se nepodařilo dokončit:\n\n{error}",
                )

# ============================================================
# FINÁLNÍ OPRAVA: čitelné klíče + průhledné UI + bílý tisk
# ============================================================

# Přesměrování na nový renderer vedle main.py.
def get_pirate_key_renderer_file():
    """Vyhledá společný modul pro generování grafických klíčů šifer.

    Cesty jsou řazené tak, aby fungovaly ve vývoji, ve Windows onedir buildu
    i v macOS .app balíčku.
    """
    candidates = [
        os.path.join(get_app_dir(), "pirate_key_renderer.py"),
        os.path.join(get_script_dir(), "pirate_key_renderer.py"),
        os.path.join(get_app_dir(), "logika_sifer", "spolecne", "pirate_key_renderer.py"),
        os.path.join(get_script_dir(), "logika_sifer", "spolecne", "pirate_key_renderer.py"),
    ]

    bundle_dir = get_pyinstaller_bundle_dir()
    if bundle_dir:
        candidates.append(os.path.join(bundle_dir, "pirate_key_renderer.py"))
        candidates.append(os.path.join(bundle_dir, "logika_sifer", "spolecne", "pirate_key_renderer.py"))

    if getattr(sys, "frozen", False) and sys.platform == "darwin":
        macos_dir = os.path.dirname(sys.executable)
        contents_dir = os.path.dirname(macos_dir)
        resources_dir = os.path.join(contents_dir, "Resources")
        candidates.append(os.path.join(resources_dir, "pirate_key_renderer.py"))
        candidates.append(os.path.join(resources_dir, "logika_sifer", "spolecne", "pirate_key_renderer.py"))

    for path in candidates:
        if path and os.path.exists(path):
            return path

    return os.path.join(get_app_dir(), "pirate_key_renderer.py")


def _reload_pirate_key_renderer():
    global PIRATE_KEY_RENDERER
    PIRATE_KEY_RENDERER = None
    return get_pirate_key_renderer()


# Binární čtverce: když přepnu na jinou šifru, starý binární widget se musí vždy schovat.
def _binary_squares_update_output_widget_mode_fixed(self):
    if self.is_binary_squares_selected() and self.result_mode == "encrypt" and hasattr(self, "binary_squares_scroll"):
        self.output_text.hide()
        _binary_squares_hide_all_draw_outputs(self, "binary_squares_scroll")
        self.binary_squares_scroll.show()
        self.binary_squares_scroll.raise_()
        _binary_squares_resize_canvas_to_content(self)
        return

    if hasattr(self, "binary_squares_scroll"):
        self.binary_squares_scroll.hide()

    return _BINARY_ORIGINAL_UPDATE_OUTPUT_WIDGET_MODE(self)


def _binary_squares_set_result_output_fixed(self, result: str):
    if self.is_binary_squares_selected() and self.result_mode == "encrypt" and hasattr(self, "binary_squares_scroll"):
        self.output_text.clear()
        self.output_text.hide()
        _binary_squares_hide_all_draw_outputs(self, "binary_squares_scroll")
        self.binary_squares_scroll.show()
        self.binary_squares_scroll.raise_()
        self.binary_squares_scroll.verticalScrollBar().setValue(0)
        _binary_squares_resize_canvas_to_content(self, self.get_input_text())
        QTimer.singleShot(0, lambda: _binary_squares_resize_canvas_to_content(self, self.get_input_text()))
        return

    if hasattr(self, "binary_squares_scroll"):
        self.binary_squares_scroll.hide()

    return _BINARY_ORIGINAL_SET_RESULT_OUTPUT(self, result)


# Binární čtverce: v UI má být průhledná plocha, ne bílý papír.
_PREV_CREATE_WIDGETS_TRANSPARENT_UI = SifratorSkinWidget.create_widgets


def _create_widgets_transparent_binary_ui(self):
    _PREV_CREATE_WIDGETS_TRANSPARENT_UI(self)
    if hasattr(self, "binary_squares_scroll"):
        self.binary_squares_scroll.setStyleSheet("""
            QScrollArea {
                background: transparent;
                border: none;
            }
            QScrollArea > QWidget > QWidget {
                background: transparent;
            }
            QScrollBar:vertical {
                background: rgba(20, 17, 12, 130);
                width: 11px;
                margin: 4px 2px 4px 2px;
                border-radius: 5px;
            }
            QScrollBar::handle:vertical {
                background: #b89b68;
                min-height: 42px;
                border-radius: 5px;
            }
            QScrollBar::add-line:vertical,
            QScrollBar::sub-line:vertical {
                height: 0px;
            }
        """)
        viewport = self.binary_squares_scroll.viewport()
        if viewport is not None:
            viewport.setStyleSheet("background: transparent;")


SifratorSkinWidget.create_widgets = _create_widgets_transparent_binary_ui
SifratorSkinWidget.update_output_widget_mode = _binary_squares_update_output_widget_mode_fixed
SifratorSkinWidget.set_result_output = _binary_squares_set_result_output_fixed





# ============================================================
# RYCHLÉ PŘEDNAČÍTÁNÍ KLÍČŮ A NÁHLEDU TISKU NA POZADÍ
# ============================================================

# Důležité: QPixmap a QWidget renderování musí běžet v hlavním Qt vlákně.
# Proto se nepoužívá Python thread, ale QTimer. Po šifrování se UI nejdřív
# překreslí a až potom se v krátké pauze připraví cache pro KLÍČ a TISK.


def _print_cache_safe_name(value: str) -> str:
    return "".join(ch if ch.isalnum() else "_" for ch in str(value or "")).strip("_") or "cache"


def _print_cache_freeze(value):
    if isinstance(value, dict):
        return tuple(sorted((str(k), _print_cache_freeze(v)) for k, v in value.items()))
    if isinstance(value, (list, tuple)):
        return tuple(_print_cache_freeze(v) for v in value)
    return str(value)


def _print_cache_context_signature(self):
    try:
        context = self.get_current_key_context_for_cipher(self.selected_cipher) if hasattr(self, "get_current_key_context_for_cipher") else {}
    except Exception:
        context = {}
    return _print_cache_freeze(context)


def _print_cache_logic_signature(self):
    logic_module = None
    try:
        logic_module = self.get_selected_logic_module() if hasattr(self, "get_selected_logic_module") else None
    except Exception:
        logic_module = None

    logic_file = getattr(logic_module, "__file__", "") if logic_module is not None else ""
    logic_mtime = 0
    try:
        if logic_file and os.path.exists(logic_file):
            logic_mtime = int(os.path.getmtime(logic_file))
    except Exception:
        logic_mtime = 0

    renderer_file = ""
    renderer_mtime = 0
    try:
        renderer_file = get_pirate_key_renderer_file()
        if renderer_file and os.path.exists(renderer_file):
            renderer_mtime = int(os.path.getmtime(renderer_file))
    except Exception:
        pass

    return logic_module, (logic_file, logic_mtime, renderer_file, renderer_mtime)


def _print_cache_current_signature(self, drawn_widget=None):
    input_text = self.input_text.toPlainText() if hasattr(self, "input_text") else ""
    output_text = self.output_text.toPlainText() if hasattr(self, "output_text") else ""
    drawn_size = ""
    if drawn_widget is not None:
        try:
            drawn_size = f"{drawn_widget.width()}x{drawn_widget.height()}"
        except Exception:
            drawn_size = "drawn"
    _logic_module, logic_sig = _print_cache_logic_signature(self)
    return (
        str(self.selected_cipher or ""),
        str(getattr(self, "result_mode", "")),
        input_text,
        output_text,
        drawn_size,
        _print_cache_context_signature(self),
        logic_sig,
    )


def _print_cache_get_key_image_path(self, force: bool = False, print_mode: bool = True) -> str:
    if not getattr(self, "selected_cipher", None):
        return ""

    renderer = get_pirate_key_renderer()
    logic_module, logic_sig = _print_cache_logic_signature(self)
    if renderer is None or logic_module is None:
        return ""

    cipher_name = self.selected_cipher or "Nevybraná šifra"
    key_context = self.get_current_key_context_for_cipher(cipher_name) if hasattr(self, "get_current_key_context_for_cipher") else None
    sig = ("key_png", cipher_name, _print_cache_context_signature(self), logic_sig, 1400, bool(print_mode))

    cache = getattr(self, "_print_key_image_cache", None)
    if not isinstance(cache, dict):
        cache = {}
        self._print_key_image_cache = cache

    cached_path = cache.get(sig)
    if not force and cached_path and os.path.exists(cached_path):
        return cached_path

    import tempfile
    safe_name = _print_cache_safe_name(cipher_name)
    suffix = "print" if print_mode else "ui"
    image_path = os.path.join(tempfile.gettempdir(), f"sifrator_klic_cache_{safe_name}_{suffix}_{abs(hash(sig))}.png")

    try:
        if hasattr(renderer, "save_key_png_for_module"):
            saved = bool(renderer.save_key_png_for_module(cipher_name, logic_module, image_path, width=1400, context=key_context, print_mode=print_mode))
        else:
            saved = False
        if saved and os.path.exists(image_path):
            # Necháme v cache více verzí, aby přepnutí náhledu nebo opakované otevření bylo okamžité.
            if len(cache) > 30:
                cache.clear()
            cache[sig] = image_path
            return image_path
    except Exception as error:
        print(f"CHYBA cache klíče: {error}")

    return ""


def _print_cache_get_drawn_result_image(self, drawn_widget=None, force: bool = False) -> QImage:
    if drawn_widget is None:
        drawn_widget = self.print_find_visible_draw_widget() if hasattr(self, "print_find_visible_draw_widget") else None
    if drawn_widget is None:
        return QImage()

    sig = ("drawn", _print_cache_current_signature(self, drawn_widget))
    cache = getattr(self, "_print_drawn_image_cache", None)
    if not isinstance(cache, dict):
        cache = {}
        self._print_drawn_image_cache = cache

    cached = cache.get(sig)
    if not force and isinstance(cached, QImage) and not cached.isNull():
        return cached

    try:
        pixmap = drawn_widget.grab()
        if pixmap.isNull():
            return QImage()
        image = _print_crop_image_to_content(pixmap.toImage())
        image = _print_image_black_on_white(image)
        if len(cache) > 20:
            cache.clear()
        cache[sig] = image
        return image
    except Exception as error:
        print(f"CHYBA cache kresleného výsledku: {error}")
        return QImage()




# Cache celého QTextDocumentu pro první náhled. Těžké obrázky jsou už uložené,
# takže otevření okna TISK nemusí znovu generovat klíč ani kreslenou šifru.
try:
    _PRINT_PRELOAD_ORIGINAL_BUILD_DOCUMENT = _print_options_build_document_base

    def _print_options_build_document(self, options: dict, drawn_widget, paper_name="A4", orientation_name="Na výšku", settings=None):
        settings = settings or {}
        sig = (
            "document",
            _print_cache_freeze(options),
            str(paper_name),
            str(orientation_name),
            _print_cache_freeze(settings),
            _print_cache_current_signature(self, drawn_widget),
        )
        cache = getattr(self, "_print_document_cache", None)
        if not isinstance(cache, dict):
            cache = {}
            self._print_document_cache = cache

        cached = cache.get(sig)
        if cached is not None:
            return cached

        document = _PRINT_PRELOAD_ORIGINAL_BUILD_DOCUMENT(
            self,
            options,
            drawn_widget,
            paper_name=paper_name,
            orientation_name=orientation_name,
            settings=settings,
        )
        if not getattr(self, "_print_preview_needs_refresh", False):
            if len(cache) > 8:
                cache.clear()
            cache[sig] = document
        return document
except Exception:
    pass

try:
    SifratorSkinWidget.print_options_build_document = _print_options_build_document
except Exception:
    pass










def _clear_print_preload_cache(self):
    for attr in ("_print_key_image_cache", "_print_drawn_image_cache", "_print_document_cache"):
        try:
            cache = getattr(self, attr, None)
            if isinstance(cache, dict):
                cache.clear()
        except Exception:
            pass


# Otevření tisku ještě před oknem TISK zkusí cache dotáhnout. Když už je hotová,
# je to okamžité. Když hotová není, připraví se alespoň teď a náhled už ji použije.
try:
    _PRELOAD_ORIGINAL_PRINT_CURRENT_RESULT = SifratorSkinWidget.print_current_result

    def _preload_print_current_result(self):
        try:
            _print_cache_preload_now(self)
        except Exception:
            pass
        return _PRELOAD_ORIGINAL_PRINT_CURRENT_RESULT(self)

    SifratorSkinWidget.print_current_result = _preload_print_current_result
except Exception:
    pass


# Po výběru jiné šifry cache smažeme a připravíme novou.
try:
    _PRELOAD_ORIGINAL_SELECT_CIPHER = SifratorSkinWidget.select_cipher

    def _preload_select_cipher(self, name):
        result = _PRELOAD_ORIGINAL_SELECT_CIPHER(self, name)
        _clear_print_preload_cache(self)
        _schedule_print_cache_preload(self, 250)
        return result

    SifratorSkinWidget.select_cipher = _preload_select_cipher
except Exception:
    pass


# Po každém šifrování/dešifrování se cache připraví sama s malým zpožděním.
try:
    _PRELOAD_ORIGINAL_AUTO_ENCRYPT_ACTION = SifratorSkinWidget.auto_encrypt_action

    def _preload_auto_encrypt_action(self):
        result = _PRELOAD_ORIGINAL_AUTO_ENCRYPT_ACTION(self)
        _schedule_print_cache_preload(self, 250)
        return result

    SifratorSkinWidget.auto_encrypt_action = _preload_auto_encrypt_action
except Exception:
    pass

try:
    _PRELOAD_ORIGINAL_ENCRYPT_ACTION = SifratorSkinWidget.encrypt_action

    def _preload_encrypt_action(self):
        result = _PRELOAD_ORIGINAL_ENCRYPT_ACTION(self)
        _schedule_print_cache_preload(self, 250)
        return result

    SifratorSkinWidget.encrypt_action = _preload_encrypt_action
except Exception:
    pass

try:
    _PRELOAD_ORIGINAL_DECRYPT_ACTION = SifratorSkinWidget.decrypt_action

    def _preload_decrypt_action(self):
        result = _PRELOAD_ORIGINAL_DECRYPT_ACTION(self)
        _schedule_print_cache_preload(self, 250)
        return result

    SifratorSkinWidget.decrypt_action = _preload_decrypt_action
except Exception:
    pass



# ============================================================
# OPRAVA ZASEKÁVÁNÍ – šetrné přednačítání cache
# ============================================================
# Důvod:
# Qt neumí bezpečně renderovat QWidget/QPixmap v běžném Python vlákně.
# Proto se přednačítání dělá v hlavním GUI vlákně přes QTimer.
# Aby se aplikace nezasekávala, těžké části se rozdělí do menších kroků
# a spouští se až po delší pauze po psaní / změně šifry.

_PRINT_PRELOAD_MIN_DELAY_MS = 1200
_PRINT_PRELOAD_STEP_DELAY_MS = 80


def _print_cache_preload_now(self, token=None):
    """Šetrné přednačtení cache bez dlouhého jednorázového záseku GUI.

    Nepředgenerovává celý QTextDocument, protože to je nejčastější příčina
    zasekávání při psaní. Dokument se vytvoří až při skutečném otevření náhledu.
    """
    if token is not None and token != getattr(self, "_print_cache_preload_token", None):
        return
    if getattr(self, "_print_cache_preload_running", False):
        return
    if not getattr(self, "selected_cipher", None):
        return

    self._print_cache_preload_running = True

    def token_valid():
        return token is None or token == getattr(self, "_print_cache_preload_token", None)

    def finish():
        try:
            self._print_cache_preload_running = False
        except Exception:
            pass

    def get_drawn_widget():
        try:
            return self.print_find_visible_draw_widget() if hasattr(self, "print_find_visible_draw_widget") else None
        except Exception:
            return None

    steps = []

    # 1) Tiskový klíč – nejdůležitější pro náhled tisku.
    def preload_print_key():
        if not token_valid():
            finish()
            return
        try:
            _print_cache_get_key_image_path(self, force=False, print_mode=True)
        except Exception as error:
            print(f"CHYBA přednačítání tiskového klíče: {error}")

    steps.append(preload_print_key)

    # 2) Kreslený výsledek pro tisk – jen pokud je vidět grafický výstup.
    def preload_drawn_result():
        if not token_valid():
            finish()
            return
        try:
            drawn_widget = get_drawn_widget()
            if drawn_widget is not None:
                _print_cache_get_drawn_result_image(self, drawn_widget, force=False)
        except Exception as error:
            print(f"CHYBA přednačítání grafické šifry: {error}")

    steps.append(preload_drawn_result)

    # 3) Běžný klíč pro dialog KLÍČ – nízká priorita, až po tisku.
    def preload_ui_key():
        if not token_valid():
            finish()
            return
        try:
            _print_cache_get_key_image_path(self, force=False, print_mode=False)
        except Exception:
            pass

    steps.append(preload_ui_key)

    def run_step(index=0):
        if not token_valid():
            finish()
            return
        if index >= len(steps):
            finish()
            return
        try:
            steps[index]()
        finally:
            if token_valid():
                QTimer.singleShot(_PRINT_PRELOAD_STEP_DELAY_MS, lambda: run_step(index + 1))
            else:
                finish()

    QTimer.singleShot(0, lambda: run_step(0))


def _schedule_print_cache_preload(self, delay_ms: int = 1200):
    """Naplánuje přednačítání až po pauze.

    Původních cca 250 ms bylo moc krátké: při psaní se preload spouštěl
    téměř hned po každém znaku a tím blokoval GUI.
    """
    try:
        self._print_cache_preload_token = int(getattr(self, "_print_cache_preload_token", 0)) + 1
        token = self._print_cache_preload_token
        requested_delay = int(delay_ms or 0)
        if requested_delay <= 150:
            # Tiskový náhled už je otevřený a čeká na chybějící obrázek klíče.
            # V tom případě začneme rychle, ale běžné psaní dál chrání delší pauza.
            delay = max(40, requested_delay)
        else:
            delay = max(requested_delay, _PRINT_PRELOAD_MIN_DELAY_MS)
        QTimer.singleShot(delay, lambda: _print_cache_preload_now(self, token))
    except Exception:
        pass



# ============================================================
# NEMRZNOUCÍ TISK A KLÍČ – finální bezpečná vrstva
# ============================================================
# Důležité:
# - kliknutí na TISK už nesmí nejdřív generovat klíč ani grabovat grafický výstup.
# - kliknutí na KLÍČ otevře dialog hned a teprve potom se klíč doplní.
# - náhled tisku používá jen hotovou cache; když cache není, zobrazí text „připravuje se“.


def _print_cache_peek_key_image_path(self, print_mode: bool = True) -> str:
    """Vrátí pouze již existující obrázek klíče z cache bez spouštění generování."""
    try:
        if not getattr(self, "selected_cipher", None):
            return ""
        logic_module, logic_sig = _print_cache_logic_signature(self)
        if logic_module is None:
            return ""
        cipher_name = self.selected_cipher or "Nevybraná šifra"
        sig = ("key_png", cipher_name, _print_cache_context_signature(self), logic_sig, 1400, bool(print_mode))
        cache = getattr(self, "_print_key_image_cache", None)
        if isinstance(cache, dict):
            path = cache.get(sig)
            if path and os.path.exists(path):
                return path
    except Exception:
        pass
    return ""


def _print_drawn_widget_can_grab_now(drawn_widget) -> bool:
    """Malý už zobrazený výsledek je levné převést do obrázku hned."""
    try:
        if drawn_widget is None or not drawn_widget.isVisible():
            return False
        width = max(0, int(drawn_widget.width()))
        height = max(0, int(drawn_widget.height()))
        if width <= 0 or height <= 0:
            return False
        return (width * height) <= 1800000
    except Exception:
        return False


def _print_cache_peek_drawn_result_image(self, drawn_widget=None) -> QImage:
    """Vrátí hotový grafický výsledek; malé viditelné výsledky zachytí hned."""
    try:
        if drawn_widget is None:
            drawn_widget = self.print_find_visible_draw_widget() if hasattr(self, "print_find_visible_draw_widget") else None
        if drawn_widget is None:
            return QImage()
        sig = ("drawn", _print_cache_current_signature(self, drawn_widget))
        cache = getattr(self, "_print_drawn_image_cache", None)
        if isinstance(cache, dict):
            image = cache.get(sig)
            if isinstance(image, QImage) and not image.isNull():
                return image
        if _print_drawn_widget_can_grab_now(drawn_widget):
            return _print_cache_get_drawn_result_image(self, drawn_widget, force=False)
    except Exception:
        pass
    return QImage()


# Přepíšeme HTML klíče: použije hotový obrázek, nebo jen placeholder.
def _print_options_key_html(self):
    cipher_name = self.selected_cipher or "Nevybraná šifra"
    parts = [f"<p><b>Vybraná šifra:</b> {_print_options_escape_html(cipher_name)}</p>"]

    image_path = _print_cache_peek_key_image_path(self, print_mode=True)
    if image_path and os.path.exists(image_path):
        url = _fixed_key_image_file_url(image_path) if '_fixed_key_image_file_url' in globals() else image_path
        parts.append(
            "<p class='key-image-box' style='margin-top:10px;'>"
            f"<img src='{_print_options_escape_html(url)}' width='{_print_current_key_image_width(self)}'>"
            "</p>"
        )
        return "".join(parts)

    parts.append("<p><i>Klíč se připravuje na pozadí. Náhled se po chvíli obnoví.</i></p>")
    try:
        self._print_preview_needs_refresh = True
    except Exception:
        pass
    try:
        if '_schedule_print_cache_preload' in globals():
            _schedule_print_cache_preload(self, 120)
    except Exception:
        pass
    return "".join(parts)


# Otevření TISK nesmí spouštět synchronní preload před otevřením okna.
try:
    SifratorSkinWidget.print_current_result = _print_current_result_with_live_preview
except Exception:
    pass


# Otevření KLÍČ: použij async dialog z rendereru, pokud je dostupný.
try:
    _NONBLOCK_ORIGINAL_SHOW_CIPHER_KEY = SifratorSkinWidget.show_cipher_key

    def _nonblocking_show_cipher_key(self):
        if not self.selected_cipher:
            QMessageBox.information(self, "Klíč šifry", "Nejdřív vyber šifru.")
            return

        renderer = get_pirate_key_renderer()
        if renderer is None:
            QMessageBox.critical(self, "Chybí generátor klíčů", "Chybí soubor:\npirate_key_renderer.py")
            return

        logic_module = self.get_selected_logic_module() if hasattr(self, "get_selected_logic_module") else None
        if logic_module is None:
            QMessageBox.information(self, "Klíč šifry", "Pro tuto šifru se nepodařilo načíst soubor logiky.")
            return

        key_context = self.get_current_key_context() if hasattr(self, "get_current_key_context") else None

        try:
            if hasattr(renderer, "show_key_dialog_nonblocking"):
                renderer.show_key_dialog_nonblocking(self, self.selected_cipher, logic_module, key_context)
            else:
                # Fallback: otevře původní dialog.
                renderer.show_key_dialog(self, self.selected_cipher, logic_module, key_context)
        except Exception as error:
            QMessageBox.critical(self, "Chyba klíče", f"Klíč se nepodařilo zobrazit:\n\n{error}")

    SifratorSkinWidget.show_cipher_key = _nonblocking_show_cipher_key
except Exception:
    pass




# ============================================================
# PATCHE: spodní klikací LOGOVÁNÍ + AKTUALIZACE + live log okno
# ============================================================

_FOOTER_ORIGINAL_CREATE_WIDGETS = SifratorSkinWidget.create_widgets
_FOOTER_ORIGINAL_UPDATE_STATUS = SifratorSkinWidget.update_status
_FOOTER_ORIGINAL_UPDATE_LAYOUT_POSITIONS = SifratorSkinWidget.update_layout_positions


def _sifrator_debug_log(message: str) -> None:
    """Zapíše diagnostickou zprávu do aktualizačního logu, pokud je dostupný."""
    try:
        if hasattr(update_manager, "_debug_log"):
            update_manager._debug_log(str(message))
            return
    except Exception:
        pass

    try:
        path = os.path.join(tempfile.gettempdir(), "sifrator_update_debug.log")
        timestamp = time.strftime("%Y-%m-%d %H:%M:%S")
        with open(path, "a", encoding="utf-8") as file:
            file.write(f"[{timestamp}] {message}\n")
    except Exception:
        pass


def _sifrator_debug_log_path() -> str:
    try:
        if hasattr(update_manager, "get_debug_log_path"):
            return update_manager.get_debug_log_path()
    except Exception:
        pass
    try:
        return os.path.join(tempfile.gettempdir(), "sifrator_update_debug.log")
    except Exception:
        return "sifrator_update_debug.log"


class LiveLogDialog(QDialog):
    """Jednoduché okno s live výpisem diagnostických logů aplikace/updateru."""

    def __init__(self, owner_window):
        super().__init__(owner_window)
        self.owner_window = owner_window
        self._last_loaded_content = None
        self.setWindowTitle("Živé logování – Šifrátor Mraveniště")
        self.resize(880, 560)

        layout = QVBoxLayout(self)
        layout.setContentsMargins(12, 12, 12, 12)
        layout.setSpacing(8)

        self.info_label = QLabel(self)
        self.info_label.setText(
            "Logy aktualizací, cache a důležitých akcí aplikace. "
            "Text můžeš označit, kopírovat nebo vyexportovat do TXT."
        )
        self.info_label.setStyleSheet("color: #ead8b3;")
        layout.addWidget(self.info_label)

        control_row = QHBoxLayout()
        self.follow_tail_check = QCheckBox("Sledovat nejnovější záznam", self)
        self.follow_tail_check.setChecked(True)
        self.follow_tail_check.setToolTip("Když je zapnuto, log se při novém záznamu posune na konec.")
        self.follow_tail_check.setStyleSheet("color: #ead8b3; background: transparent;")
        control_row.addWidget(self.follow_tail_check)
        control_row.addStretch(1)
        layout.addLayout(control_row)

        self.log_edit = QTextEdit(self)
        self.log_edit.setReadOnly(True)
        self.log_edit.setAcceptRichText(False)
        self.log_edit.setTextInteractionFlags(Qt.TextSelectableByMouse | Qt.TextSelectableByKeyboard)
        self.log_edit.setLineWrapMode(QTextEdit.NoWrap)
        self.log_edit.setStyleSheet("""
            QTextEdit {
                color: #f0e2c0;
                background: #071018;
                border: 1px solid #8a6938;
                border-radius: 6px;
                padding: 8px;
                font-family: Consolas, Menlo, Monaco, monospace;
                font-size: 12px;
            }
        """)
        layout.addWidget(self.log_edit, 1)

        button_row = QHBoxLayout()
        self.path_label = QLabel(self)
        self.path_label.setText(_sifrator_debug_log_path())
        self.path_label.setStyleSheet("color: #a8a295;")
        button_row.addWidget(self.path_label, 1)

        self.clear_button = QPushButton("Vyčistit log", self)
        self.copy_button = QPushButton("Kopírovat", self)
        self.export_button = QPushButton("Export TXT", self)
        self.refresh_button = QPushButton("Obnovit", self)
        self.close_button = QPushButton("Zavřít", self)
        button_row.addWidget(self.clear_button)
        button_row.addWidget(self.copy_button)
        button_row.addWidget(self.export_button)
        button_row.addWidget(self.refresh_button)
        button_row.addWidget(self.close_button)
        layout.addLayout(button_row)

        self.clear_button.clicked.connect(self.clear_log)
        self.copy_button.clicked.connect(self.copy_log)
        self.export_button.clicked.connect(self.export_log)
        self.refresh_button.clicked.connect(self.refresh_log)
        self.close_button.clicked.connect(self.close)
        self.follow_tail_check.toggled.connect(self.follow_tail_changed)

        self.timer = QTimer(self)
        self.timer.timeout.connect(self.refresh_log)
        self.timer.start(700)
        self.refresh_log()

    def read_full_log_text(self) -> str:
        path = _sifrator_debug_log_path()
        if not os.path.exists(path):
            return ""
        with open(path, "r", encoding="utf-8", errors="replace") as file:
            return file.read()

    def selected_or_visible_text(self) -> str:
        cursor = self.log_edit.textCursor()
        if cursor is not None and cursor.hasSelection():
            return cursor.selectedText().replace("\u2029", "\n")
        return self.log_edit.toPlainText()

    def copy_log(self):
        text = self.selected_or_visible_text()
        if not text.strip():
            QMessageBox.information(self, "Kopírování logu", "Log je zatím prázdný.")
            return
        clipboard = QApplication.clipboard()
        if clipboard is not None:
            clipboard.setText(text)
        _sifrator_debug_log("Log byl zkopírován do schránky.")

    def export_log(self):
        try:
            text = self.read_full_log_text()
        except Exception:
            text = self.log_edit.toPlainText()

        if not text.strip():
            QMessageBox.information(self, "Export logu", "Log je zatím prázdný.")
            return

        default_name = f"sifrator_log_{time.strftime('%Y%m%d_%H%M%S')}.txt"
        path, _ = QFileDialog.getSaveFileName(
            self,
            "Exportovat log jako TXT",
            default_name,
            "Textový soubor (*.txt);;Všechny soubory (*.*)",
        )
        if not path:
            return
        if not path.lower().endswith(".txt"):
            path += ".txt"

        try:
            with open(path, "w", encoding="utf-8") as file:
                file.write(text)
            _sifrator_debug_log(f"Log byl exportován do TXT: {path}")
            QMessageBox.information(self, "Export logu", f"Log byl uložen do:\n{path}")
            self.refresh_log()
        except Exception as error:
            QMessageBox.warning(self, "Export logu", f"Log se nepodařilo uložit:\n{error}")

    def follow_tail_changed(self, checked: bool):
        if checked:
            self.log_edit.verticalScrollBar().setValue(self.log_edit.verticalScrollBar().maximum())

    def clear_log(self):
        path = _sifrator_debug_log_path()
        try:
            with open(path, "w", encoding="utf-8") as file:
                file.write("")
            self._last_loaded_content = None
        except Exception as error:
            self.log_edit.setPlainText(f"Log se nepodařilo vyčistit:\n{error}")
            return
        _sifrator_debug_log("Log byl vyčištěn z okna LOGOVÁNÍ.")
        self.refresh_log()

    def refresh_log(self):
        path = _sifrator_debug_log_path()
        self.path_label.setText(path)
        try:
            if not os.path.exists(path):
                content = "Log zatím neexistuje. Spusť kontrolu aktualizací nebo proveď akci v aplikaci."
                if content != self._last_loaded_content:
                    self.log_edit.setPlainText(content)
                    self._last_loaded_content = content
                return

            content = self.read_full_log_text()

            if not content.strip():
                content = "Log je zatím prázdný."

            # Nezobrazujeme nekonečně dlouhý soubor, poslední část stačí pro živou diagnostiku.
            max_chars = 80000
            if len(content) > max_chars:
                content = "... zkráceno na posledních 80000 znaků ...\n" + content[-max_chars:]

            old_scroll = self.log_edit.verticalScrollBar().value()
            should_follow = bool(self.follow_tail_check.isChecked())
            if content == self._last_loaded_content:
                if should_follow:
                    self.log_edit.verticalScrollBar().setValue(self.log_edit.verticalScrollBar().maximum())
                return

            self.log_edit.setPlainText(content)
            self._last_loaded_content = content
            if should_follow:
                self.log_edit.verticalScrollBar().setValue(self.log_edit.verticalScrollBar().maximum())
            else:
                self.log_edit.verticalScrollBar().setValue(min(old_scroll, self.log_edit.verticalScrollBar().maximum()))
        except Exception as error:
            self.log_edit.setPlainText(f"Log se nepodařilo načíst:\n{error}")

    def closeEvent(self, event):
        try:
            self.timer.stop()
        except Exception:
            pass
        if self.owner_window is not None and hasattr(self.owner_window, "set_live_logging_enabled"):
            self.owner_window.set_live_logging_enabled(False)
        super().closeEvent(event)


_HISTORY_MAX_ITEMS = 120
_HISTORY_MAX_TEXT_CHARS = 20000
_USER_DATA_ITEM_NAMES = [
    "historie_zprav.json",
    "plan_tabor_sifer.json",
    "poznamky_sifer.json",
    "history_images",
    "planner_attachments",
]
_HISTORY_STORAGE_DIR_CACHE: str | None = None
_HISTORY_STORAGE_MIGRATED = False


def _legacy_user_data_roots() -> list[str]:
    roots = []
    base_roots = [get_app_dir(), get_script_dir(), os.path.join(tempfile.gettempdir(), APP_NAME)]
    subdirs = ["", APP_NAME, "user_data", "data", "sifrator_data"]

    for base_root in base_roots:
        if not base_root:
            continue
        if os.path.basename(os.path.normpath(base_root)) == APP_NAME:
            roots.append(base_root)
            continue
        for subdir in subdirs:
            roots.append(os.path.join(base_root, subdir) if subdir else base_root)

    unique = []
    seen = set()
    for root in roots:
        try:
            key = os.path.normcase(os.path.abspath(root))
        except Exception:
            continue
        if key in seen:
            continue
        seen.add(key)
        unique.append(root)
    return unique


def _migrate_history_storage_if_needed(target_dir: str) -> None:
    global _HISTORY_STORAGE_MIGRATED
    if _HISTORY_STORAGE_MIGRATED:
        return

    _HISTORY_STORAGE_MIGRATED = True
    try:
        copied = migrate_user_data_items(_legacy_user_data_roots(), _USER_DATA_ITEM_NAMES, target_dir=target_dir)
        if copied:
            _sifrator_debug_log(f"Migrace uživatelských dat dokončena: {copied} souborů -> {target_dir}")
    except Exception as error:
        _sifrator_debug_log(f"Migrace uživatelských dat selhala: {type(error).__name__}: {error}")


def _history_storage_dir() -> str:
    """Vrátí zapisovatelnou složku pro uživatelskou historii."""
    global _HISTORY_STORAGE_DIR_CACHE
    if _HISTORY_STORAGE_DIR_CACHE:
        return _HISTORY_STORAGE_DIR_CACHE

    try:
        path = get_user_data_dir()
        _migrate_history_storage_if_needed(path)
        _HISTORY_STORAGE_DIR_CACHE = path
        return path
    except Exception:
        path = os.path.join(tempfile.gettempdir(), APP_NAME)
        os.makedirs(path, exist_ok=True)
        _migrate_history_storage_if_needed(path)
        _HISTORY_STORAGE_DIR_CACHE = path
        return path


def _history_file_path() -> str:
    return os.path.join(_history_storage_dir(), "historie_zprav.json")


def _history_image_dir() -> str:
    path = os.path.join(_history_storage_dir(), "history_images")
    os.makedirs(path, exist_ok=True)
    return path


def _history_image_path(image_name: str) -> str:
    image_name = str(image_name or "").strip()
    if not image_name:
        return ""

    candidates = []
    if os.path.isabs(image_name):
        candidates.append(image_name)
    candidates.append(os.path.join(_history_image_dir(), os.path.basename(image_name)))

    for path in candidates:
        if path and os.path.exists(path):
            return path
    return ""


def _history_delete_image(image_name: str) -> None:
    path = _history_image_path(image_name)
    if not path:
        return
    try:
        if os.path.dirname(os.path.abspath(path)) == os.path.abspath(_history_image_dir()):
            os.remove(path)
    except Exception:
        pass


def _history_prune_images(entries: list[dict]) -> None:
    try:
        keep = {
            os.path.basename(str(entry.get("image", "") or ""))
            for entry in entries
            if isinstance(entry, dict) and entry.get("image")
        }
        image_dir = _history_image_dir()
        for file_name in os.listdir(image_dir):
            if not file_name.lower().endswith(".png") or file_name in keep:
                continue
            try:
                os.remove(os.path.join(image_dir, file_name))
            except Exception:
                pass
    except Exception:
        pass


def _history_limit_text(value: str) -> str:
    text = str(value or "")
    if len(text) <= _HISTORY_MAX_TEXT_CHARS:
        return text
    return text[:_HISTORY_MAX_TEXT_CHARS] + "\n\n... zkráceno kvůli velikosti záznamu ..."


def _history_load() -> list[dict]:
    path = _history_file_path()
    if not os.path.exists(path):
        return []

    try:
        with open(path, "r", encoding="utf-8", errors="replace") as file:
            data = json.load(file)
    except Exception:
        return []

    if isinstance(data, dict):
        data = data.get("items", [])
    if not isinstance(data, list):
        return []

    entries = []
    for item in data:
        if isinstance(item, dict):
            entries.append(item)
    return entries[:_HISTORY_MAX_ITEMS]


def _history_save(entries: list[dict]) -> None:
    path = _history_file_path()
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8") as file:
        json.dump(entries[:_HISTORY_MAX_ITEMS], file, ensure_ascii=False, indent=2)


def _history_context_signature(context) -> str:
    try:
        return json.dumps(context or {}, ensure_ascii=False, sort_keys=True, default=str)
    except Exception:
        return str(context or {})


def _history_entry_signature(entry: dict) -> tuple:
    return (
        str(entry.get("cipher", "")),
        str(entry.get("input", "")),
        str(entry.get("output", "")),
        _history_context_signature(entry.get("context", {})),
    )


def _history_add_entry(entry: dict) -> bool:
    entries = _history_load()
    signature = _history_entry_signature(entry)

    if entries and _history_entry_signature(entries[0]) == signature:
        new_image = str(entry.get("image", "") or "")
        if new_image and not entries[0].get("image"):
            entries[0]["image"] = new_image
            _history_save(entries)
            _history_prune_images(entries)
            return True
        _history_delete_image(new_image)
        return False

    entries = [item for item in entries if _history_entry_signature(item) != signature]
    entries.insert(0, entry)
    _history_save(entries)
    _history_prune_images(entries)
    return True


def _history_count() -> int:
    return len(_history_load())


def _history_context_label(context) -> str:
    if not isinstance(context, dict) or not context:
        return ""

    labels = []
    if "caesar_shift" in context:
        labels.append(f"posun={context.get('caesar_shift')}")
    if "caesar_direction" in context:
        labels.append(f"směr={context.get('caesar_direction')}")

    for key, value in sorted(context.items()):
        if key in ("caesar_shift", "caesar_direction"):
            continue
        labels.append(f"{key}={value}")

    return ", ".join(labels)


def _history_format_entries(entries: list[dict]) -> str:
    if not entries:
        return "Historie je zatím prázdná."

    blocks = []
    for index, entry in enumerate(entries, start=1):
        created_at = str(entry.get("created_at", "")).strip()
        cipher = str(entry.get("cipher", "Nevybraná šifra")).strip() or "Nevybraná šifra"
        context = _history_context_label(entry.get("context", {}))
        input_text = str(entry.get("input", ""))
        output_text = str(entry.get("output", ""))
        image_path = _history_image_path(entry.get("image", ""))

        title = f"{index}. {created_at} | {cipher}" if created_at else f"{index}. {cipher}"
        if context:
            title += f" ({context})"

        block = (
            f"{title}\n"
            f"{'-' * min(72, max(16, len(title)))}\n"
            f"Vstup:\n{input_text}\n\n"
            f"Zašifrováno:\n{output_text}"
        )
        if image_path:
            block += f"\n\nObrázkový náhled:\n{image_path}"
        blocks.append(block)

    return ("\n\n" + "=" * 72 + "\n\n").join(blocks)


def _history_escape_html(value: str) -> str:
    import html

    return html.escape(str(value or "")).replace("\n", "<br>")


def _history_format_entries_html(entries: list[dict]) -> str:
    if not entries:
        return """
        <html><body style="color:#f0e2c0; font-family:Consolas, Menlo, Monaco, monospace; font-size:12px;">
        Historie je zatím prázdná.
        </body></html>
        """

    blocks = []
    for index, entry in enumerate(entries, start=1):
        created_at = str(entry.get("created_at", "")).strip()
        cipher = str(entry.get("cipher", "Nevybraná šifra")).strip() or "Nevybraná šifra"
        context = _history_context_label(entry.get("context", {}))
        input_text = _history_escape_html(entry.get("input", ""))
        output_text = _history_escape_html(entry.get("output", ""))

        title = f"{index}. {created_at} | {cipher}" if created_at else f"{index}. {cipher}"
        if context:
            title += f" ({context})"

        image_html = ""
        image_path = _history_image_path(entry.get("image", ""))
        if image_path:
            src = _history_escape_html(QUrl.fromLocalFile(image_path).toString())
            image = QImage(image_path)
            display_width = 760
            if not image.isNull():
                display_width = min(display_width, max(1, image.width()))
            image_html = (
                "<div style='margin-top:10px; padding:8px; border:1px solid #8a6938; "
                "background:#101923;'>"
                f"<img src='{src}' width='{display_width}'>"
                "</div>"
            )

        blocks.append(
            "<div style='margin-bottom:18px;'>"
            f"<div style='color:#f3d79a; font-weight:bold;'>{_history_escape_html(title)}</div>"
            "<div style='height:1px; background:#8a6938; margin:4px 0 8px 0;'></div>"
            "<div style='color:#ead8b3; font-weight:bold;'>Vstup:</div>"
            f"<div>{input_text}</div>"
            "<div style='height:8px;'></div>"
            "<div style='color:#ead8b3; font-weight:bold;'>Zašifrováno:</div>"
            f"<div>{output_text}</div>"
            f"{image_html}"
            "</div>"
        )

    separator = "<div style='height:1px; background:#6f5734; margin:12px 0 18px 0;'></div>"
    return (
        "<html><body style='color:#f0e2c0; font-family:Consolas, Menlo, Monaco, monospace; "
        "font-size:12px; background:transparent;'>"
        + separator.join(blocks)
        + "</body></html>"
    )


def _history_serializable_context(widget) -> dict:
    try:
        context = widget.get_current_key_context() if hasattr(widget, "get_current_key_context") else {}
        if not isinstance(context, dict):
            return {}
        return json.loads(json.dumps(context, ensure_ascii=False, default=str))
    except Exception:
        return {}


def _history_capture_preview_image(widget, entry: dict) -> str:
    try:
        drawn_widget = None
        if hasattr(widget, "print_find_visible_draw_widget"):
            drawn_widget = widget.print_find_visible_draw_widget()
        if drawn_widget is None:
            return ""

        pixmap = drawn_widget.grab()
        if pixmap.isNull():
            return ""

        max_width = 1200
        if pixmap.width() > max_width:
            pixmap = pixmap.scaledToWidth(max_width, Qt.SmoothTransformation)

        payload = json.dumps(
            {
                "created_at": entry.get("created_at", ""),
                "cipher": entry.get("cipher", ""),
                "context": entry.get("context", {}),
                "input": entry.get("input", ""),
                "output": entry.get("output", ""),
            },
            ensure_ascii=False,
            sort_keys=True,
            default=str,
        )
        digest = hashlib.sha1(payload.encode("utf-8", errors="replace")).hexdigest()[:16]
        file_name = f"historie_{time.strftime('%Y%m%d_%H%M%S')}_{digest}.png"
        path = os.path.join(_history_image_dir(), file_name)
        if pixmap.save(path, "PNG"):
            return file_name
    except Exception as error:
        _sifrator_debug_log(f"Historie: uložení obrázkového náhledu selhalo: {type(error).__name__}: {error}")

    return ""


def _history_current_output(widget) -> str:
    result = str(getattr(widget, "_history_last_result_output", "") or "")
    if result.strip():
        return result

    try:
        output_text = getattr(widget, "output_text", None)
        if output_text is not None and output_text.isVisible():
            return output_text.toPlainText()
    except Exception:
        pass

    return ""


def _history_capture_now(widget, token=None) -> None:
    try:
        if token is not None and token != getattr(widget, "_history_capture_token", None):
            return
        if getattr(widget, "result_mode", None) != "encrypt":
            return

        cipher = str(getattr(widget, "selected_cipher", "") or "").strip()
        if not cipher:
            return

        input_widget = getattr(widget, "input_text", None)
        input_text = input_widget.toPlainText() if input_widget is not None else ""
        if not input_text.strip():
            return

        output_text = _history_current_output(widget)
        if not output_text.strip():
            return

        entry = {
            "created_at": time.strftime("%Y-%m-%d %H:%M:%S"),
            "cipher": cipher,
            "context": _history_serializable_context(widget),
            "input": _history_limit_text(input_text),
            "output": _history_limit_text(output_text),
        }
        image_name = _history_capture_preview_image(widget, entry)
        if image_name:
            entry["image"] = image_name

        if _history_add_entry(entry):
            _sifrator_debug_log(f"Historie: uložen záznam šifrování: šifra={cipher}")
            try:
                widget.update_status()
            except Exception:
                pass

            try:
                window = widget.window()
                dialog = getattr(window, "_history_dialog", None)
                if dialog is not None and dialog.isVisible():
                    dialog.refresh_history()
                overview_dialog = getattr(window, "_cipher_overview_dialog", None)
                if overview_dialog is not None and overview_dialog.isVisible():
                    overview_dialog.refresh_overview()
            except Exception:
                pass
    except Exception as error:
        _sifrator_debug_log(f"Historie: uložení selhalo: {type(error).__name__}: {error}")


def _history_schedule_capture(widget, delay_ms: int = 900) -> None:
    try:
        if getattr(widget, "result_mode", None) != "encrypt":
            return
        if not getattr(widget, "selected_cipher", None):
            return
        input_widget = getattr(widget, "input_text", None)
        if input_widget is None or not input_widget.toPlainText().strip():
            return

        widget._history_capture_token = int(getattr(widget, "_history_capture_token", 0)) + 1
        token = widget._history_capture_token
        QTimer.singleShot(max(250, int(delay_ms)), lambda: _history_capture_now(widget, token))
    except Exception:
        pass


class HistoryDialog(PirateModuleDialog):
    """Okno s historií zašifrovaných zpráv."""

    def __init__(self, owner_window):
        super().__init__(
            owner_window,
            "history_BG.png",
            ((0.065, 0.287, 0.95), (0.871, 0.237, 0.30), (0.945, 0.837, 1.20)),
        )
        self.owner_window = owner_window
        self.setWindowTitle("Historie zašifrovaných zpráv")
        self.resize(900, 620)

        layout = QVBoxLayout(self)
        layout.setContentsMargins(12, 12, 12, 12)
        layout.setSpacing(8)

        self.info_label = QLabel("Poslední zašifrované zprávy", self)
        self.info_label.setStyleSheet("color: #ead8b3;")
        layout.addWidget(self.info_label)

        self.history_edit = QTextEdit(self)
        self.history_edit.setReadOnly(True)
        self.history_edit.setAcceptRichText(False)
        self.history_edit.setTextInteractionFlags(Qt.TextSelectableByMouse | Qt.TextSelectableByKeyboard)
        self.history_edit.setLineWrapMode(QTextEdit.WidgetWidth)
        self.history_edit.setWordWrapMode(QTextOption.WrapAnywhere)
        self.history_edit.setStyleSheet("""
            QTextEdit {
                color: #f0e2c0;
                background: #071018;
                border: 1px solid #8a6938;
                border-radius: 6px;
                padding: 8px;
                font-family: Consolas, Menlo, Monaco, monospace;
                font-size: 12px;
            }
        """)
        layout.addWidget(self.history_edit, 1)

        button_row = QHBoxLayout()
        self.path_label = QLabel(self)
        self.path_label.setStyleSheet("color: #a8a295;")
        button_row.addWidget(self.path_label, 1)

        self.clear_button = QPushButton("Vyčistit historii", self)
        self.copy_button = QPushButton("Kopírovat", self)
        self.export_button = QPushButton("Export TXT", self)
        self.refresh_button = QPushButton("Obnovit", self)
        self.close_button = QPushButton("Zavřít", self)
        button_row.addWidget(self.clear_button)
        button_row.addWidget(self.copy_button)
        button_row.addWidget(self.export_button)
        button_row.addWidget(self.refresh_button)
        button_row.addWidget(self.close_button)
        layout.addLayout(button_row)

        self.clear_button.clicked.connect(self.clear_history)
        self.copy_button.clicked.connect(self.copy_history)
        self.export_button.clicked.connect(self.export_history)
        self.refresh_button.clicked.connect(self.refresh_history)
        self.close_button.clicked.connect(self.close)

        self.apply_pirate_glass()
        self.refresh_history()

    def selected_or_visible_text(self) -> str:
        cursor = self.history_edit.textCursor()
        if cursor is not None and cursor.hasSelection():
            return cursor.selectedText().replace("\u2029", "\n")
        return self.history_edit.toPlainText()

    def refresh_history(self):
        entries = _history_load()
        self.path_label.setText(_history_file_path())
        self.info_label.setText(f"Poslední zašifrované zprávy ({len(entries)})")
        self.history_edit.setHtml(_history_format_entries_html(entries))
        self.history_edit.moveCursor(QTextCursor.Start)

    def copy_history(self):
        text = self.selected_or_visible_text()
        if not text.strip() or text.strip() == "Historie je zatím prázdná.":
            QMessageBox.information(self, "Kopírování historie", "Historie je zatím prázdná.")
            return
        clipboard = QApplication.clipboard()
        if clipboard is not None:
            clipboard.setText(text)

    def export_history(self):
        entries = _history_load()
        if not entries:
            QMessageBox.information(self, "Export historie", "Historie je zatím prázdná.")
            return

        default_name = f"sifrator_historie_{time.strftime('%Y%m%d_%H%M%S')}.txt"
        path, _ = QFileDialog.getSaveFileName(
            self,
            "Exportovat historii jako TXT",
            default_name,
            "Textový soubor (*.txt);;Všechny soubory (*.*)",
        )
        if not path:
            return
        if not path.lower().endswith(".txt"):
            path += ".txt"

        try:
            with open(path, "w", encoding="utf-8") as file:
                file.write(_history_format_entries(entries).strip() + "\n")
            QMessageBox.information(self, "Export historie", f"Historie byla uložena do:\n{path}")
        except Exception as error:
            QMessageBox.warning(self, "Export historie", f"Historii se nepodařilo uložit:\n{error}")

    def clear_history(self):
        if not _history_load():
            QMessageBox.information(self, "Historie", "Historie je zatím prázdná.")
            return

        answer = QMessageBox.question(
            self,
            "Vyčistit historii",
            "Opravdu chceš smazat historii zašifrovaných zpráv?",
            QMessageBox.Yes | QMessageBox.No,
            QMessageBox.No,
        )
        if answer != QMessageBox.Yes:
            return

        try:
            _history_save([])
            _history_prune_images([])
            _sifrator_debug_log("Historie zašifrovaných zpráv byla vyčištěna.")
            self.refresh_history()
            if self.owner_window is not None and hasattr(self.owner_window, "central"):
                self.owner_window.central.update_status()
        except Exception as error:
            QMessageBox.warning(self, "Historie", f"Historii se nepodařilo vyčistit:\n{error}")


def _planner_file_path() -> str:
    return os.path.join(_history_storage_dir(), "plan_tabor_sifer.json")


_PLANNER_MAX_DAYS = 60
_PLANNER_SEGMENTS = (
    ("morning", "Dopoledne"),
    ("afternoon", "Odpoledne"),
    ("evening", "Večer"),
    ("whole_day", "Celý den"),
)
_PLANNER_DEFAULT_SEGMENT_KEY = "whole_day"
_PLANNER_WEEKDAY_NAMES = (
    "pondělí",
    "úterý",
    "středa",
    "čtvrtek",
    "pátek",
    "sobota",
    "neděle",
)


def _planner_clamp_day_count(count: int | str | None, default: int = 7) -> int:
    try:
        value = int(count if count is not None else default)
    except Exception:
        value = default
    return max(1, min(_PLANNER_MAX_DAYS, value))


def _planner_parse_date(value, fallback: date | None = None) -> date | None:
    if isinstance(value, date):
        return value

    text = str(value or "").strip()
    if not text:
        return fallback

    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d. %m. %Y", "%d.%m.%y", "%d. %m. %y"):
        try:
            return datetime.strptime(text, fmt).date()
        except Exception:
            pass
    return fallback


def _planner_date_to_iso(day_date: date | None) -> str:
    safe_date = _planner_parse_date(day_date, date.today())
    return safe_date.isoformat() if safe_date else date.today().isoformat()


def _planner_range_end(start_date: date, day_count: int) -> date:
    return start_date + timedelta(days=_planner_clamp_day_count(day_count) - 1)


def _planner_inclusive_day_count(start_date: date, end_date: date) -> int:
    if end_date < start_date:
        return 1
    return _planner_clamp_day_count((end_date - start_date).days + 1)


def _planner_format_date(value) -> str:
    day_date = _planner_parse_date(value)
    if day_date is None:
        return ""
    return f"{day_date.day}. {day_date.month}. {day_date.year}"


def _planner_weekday_name(value) -> str:
    day_date = _planner_parse_date(value)
    if day_date is None:
        return ""
    return _PLANNER_WEEKDAY_NAMES[day_date.weekday()]


def _planner_empty_segment() -> dict:
    return {"note": "", "encrypted_text": "", "ciphers": [], "attachments": []}


def _planner_default_segments() -> dict:
    return {key: _planner_empty_segment() for key, _label in _PLANNER_SEGMENTS}


def _planner_default_day(index: int, day_date: date | None = None) -> dict:
    safe_date = _planner_parse_date(day_date, date.today() + timedelta(days=max(0, int(index) - 1)))
    return {
        "name": f"Den {index}",
        "date": _planner_date_to_iso(safe_date),
        "weekday": _planner_weekday_name(safe_date),
        "segments": _planner_default_segments(),
    }


def _planner_default_days(count: int = 7, start_date: date | str | None = None) -> list[dict]:
    safe_count = _planner_clamp_day_count(count)
    safe_start = _planner_parse_date(start_date, date.today()) or date.today()
    return [
        _planner_default_day(index, safe_start + timedelta(days=index - 1))
        for index in range(1, safe_count + 1)
    ]


def _planner_known_cipher_names() -> list[str]:
    try:
        return list(list_cipher_names())
    except Exception:
        return []


def _planner_normalize_ciphers(raw_ciphers, known: set[str] | None = None) -> list[str]:
    ciphers = []
    seen = set()
    if isinstance(raw_ciphers, str):
        raw_ciphers = [raw_ciphers]
    if not isinstance(raw_ciphers, list):
        raw_ciphers = []

    for cipher in raw_ciphers:
        cipher_name = str(cipher or "").strip()
        if not cipher_name or cipher_name in seen:
            continue
        if known and cipher_name not in known:
            continue
        seen.add(cipher_name)
        ciphers.append(cipher_name)
    return ciphers


def _planner_attachment_title(path: str, title: str | None = None) -> str:
    safe_title = str(title or "").strip()
    if safe_title:
        return safe_title
    return os.path.basename(str(path or "").strip()) or "Příloha"


def _planner_attachment_kind(path: str) -> str:
    suffix = os.path.splitext(str(path or "").strip())[1].lower().lstrip(".")
    if suffix == "pdf":
        return "pdf"
    if suffix in ("docx", "doc"):
        return suffix
    if suffix in ("txt", "md", "rtf"):
        return suffix
    return suffix or "soubor"


def _planner_attachment_storage_dir() -> str:
    path = os.path.join(_history_storage_dir(), "planner_attachments")
    os.makedirs(path, exist_ok=True)
    return path


def _planner_safe_attachment_name(path: str) -> str:
    file_name = os.path.basename(str(path or "").strip()) or "priloha"
    stem, suffix = os.path.splitext(file_name)
    safe_stem = "".join(ch if ch.isalnum() or ch in " ._-" else "_" for ch in stem).strip(" ._")
    safe_suffix = "".join(ch if ch.isalnum() or ch == "." else "" for ch in suffix).strip()
    return f"{(safe_stem or 'priloha')[:90]}{safe_suffix[:16]}"


def _planner_unique_attachment_path(source_path: str) -> str:
    folder = _planner_attachment_storage_dir()
    safe_name = _planner_safe_attachment_name(source_path)
    stem, suffix = os.path.splitext(safe_name)
    candidate = os.path.join(folder, safe_name)
    if not os.path.exists(candidate):
        return candidate

    for index in range(2, 10000):
        candidate = os.path.join(folder, f"{stem}_{index}{suffix}")
        if not os.path.exists(candidate):
            return candidate

    timestamp = time.strftime("%Y%m%d_%H%M%S")
    return os.path.join(folder, f"{stem}_{timestamp}{suffix}")


def _planner_import_attachment(path: str) -> dict | None:
    attachment = _planner_normalize_attachment(path)
    if attachment is None:
        return None

    source_path = str(path or "").strip()
    if not source_path or not os.path.isfile(source_path):
        return attachment

    try:
        target_path = _planner_unique_attachment_path(source_path)
        if os.path.normcase(os.path.abspath(source_path)) != os.path.normcase(os.path.abspath(target_path)):
            shutil.copy2(source_path, target_path)
            attachment["path"] = target_path
            attachment["source_path"] = os.path.abspath(source_path)
            attachment["kind"] = _planner_attachment_kind(target_path)
    except Exception as error:
        _sifrator_debug_log(f"Import přílohy selhal, ponechávám původní cestu: {type(error).__name__}: {error}")

    return attachment


def _planner_normalize_attachment(raw_attachment) -> dict | None:
    if isinstance(raw_attachment, str):
        raw_attachment = {"path": raw_attachment}
    if not isinstance(raw_attachment, dict):
        return None

    path = str(raw_attachment.get("path") or raw_attachment.get("file") or "").strip()
    if not path:
        return None

    title = _planner_attachment_title(path, raw_attachment.get("title") or raw_attachment.get("name"))
    attachment = {
        "path": path,
        "title": title,
        "kind": _planner_attachment_kind(path),
        "note": str(raw_attachment.get("note") or "").strip(),
    }
    source_path = str(raw_attachment.get("source_path") or "").strip()
    if source_path:
        attachment["source_path"] = source_path
    return attachment


def _planner_attachment_identity(attachment: dict) -> str:
    path = str(attachment.get("source_path") or attachment.get("path") or "").strip()
    try:
        return os.path.normcase(os.path.abspath(path))
    except Exception:
        return path


def _planner_normalize_attachments(raw_attachments) -> list[dict]:
    if isinstance(raw_attachments, (str, dict)):
        raw_attachments = [raw_attachments]
    if not isinstance(raw_attachments, list):
        raw_attachments = []

    attachments = []
    seen = set()
    for raw_attachment in raw_attachments:
        attachment = _planner_normalize_attachment(raw_attachment)
        if attachment is None:
            continue
        key = _planner_attachment_identity(attachment)
        if key in seen:
            continue
        seen.add(key)
        attachments.append(attachment)
    return attachments


def _planner_normalize_segment(raw_segment, known: set[str] | None = None) -> dict:
    if not isinstance(raw_segment, dict):
        raw_segment = {}
    return {
        "note": str(raw_segment.get("note") or raw_segment.get("notes") or "").strip(),
        "encrypted_text": str(
            raw_segment.get("encrypted_text")
            or raw_segment.get("cipher_text")
            or raw_segment.get("encrypted")
            or raw_segment.get("output")
            or ""
        ).strip(),
        "ciphers": _planner_normalize_ciphers(raw_segment.get("ciphers", []), known),
        "attachments": _planner_normalize_attachments(raw_segment.get("attachments", [])),
    }


def _planner_segment_has_content(segment: dict) -> bool:
    return bool(
        segment.get("ciphers")
        or segment.get("attachments")
        or str(segment.get("note") or "").strip()
        or str(segment.get("encrypted_text") or "").strip()
    )


def _planner_normalize_segments(raw_day: dict, known: set[str] | None = None) -> dict:
    raw_segments = raw_day.get("segments", {})
    if not isinstance(raw_segments, dict):
        raw_segments = {}

    segments = {}
    for key, label in _PLANNER_SEGMENTS:
        raw_segment = raw_segments.get(key)
        if raw_segment is None:
            raw_segment = raw_segments.get(label)
        segments[key] = _planner_normalize_segment(raw_segment, known)

    legacy_segment = _planner_normalize_segment(raw_day, known)
    if (
        _planner_segment_has_content(legacy_segment)
        and not _planner_segment_has_content(segments[_PLANNER_DEFAULT_SEGMENT_KEY])
    ):
        segments[_PLANNER_DEFAULT_SEGMENT_KEY] = legacy_segment

    return segments


def _planner_normalize_days(raw_days, target_count: int | None = None, start_date: date | str | None = None) -> list[dict]:
    count = _planner_clamp_day_count(target_count)
    known = set(_planner_known_cipher_names())
    safe_start = _planner_parse_date(start_date, date.today()) or date.today()
    days = []

    if isinstance(raw_days, list):
        source_days = raw_days
    else:
        source_days = []

    for index in range(1, count + 1):
        raw_day = source_days[index - 1] if index - 1 < len(source_days) else {}
        if not isinstance(raw_day, dict):
            raw_day = {}

        name = str(raw_day.get("name") or f"Den {index}").strip() or f"Den {index}"
        day_date = safe_start + timedelta(days=index - 1)
        days.append({
            "name": name,
            "date": _planner_date_to_iso(day_date),
            "weekday": _planner_weekday_name(day_date),
            "segments": _planner_normalize_segments(raw_day, known),
        })

    return days


def _planner_default_plan(count: int = 7, start_date: date | str | None = None) -> dict:
    safe_count = _planner_clamp_day_count(count)
    safe_start = _planner_parse_date(start_date, date.today()) or date.today()
    safe_end = _planner_range_end(safe_start, safe_count)
    return {
        "day_count": safe_count,
        "start_date": _planner_date_to_iso(safe_start),
        "end_date": _planner_date_to_iso(safe_end),
        "days": _planner_default_days(safe_count, safe_start),
    }


def _planner_plan_range(data: dict, raw_days: list, default_count: int) -> tuple[date, date, int]:
    start_date = _planner_parse_date(data.get("start_date"))
    if start_date is None and raw_days and isinstance(raw_days[0], dict):
        start_date = _planner_parse_date(raw_days[0].get("date"))
    if start_date is None:
        start_date = date.today()

    stored_count = _planner_clamp_day_count(data.get("day_count") or default_count)
    end_date = _planner_parse_date(data.get("end_date"))
    if end_date is None:
        end_date = _planner_range_end(start_date, stored_count)
    if end_date < start_date:
        end_date = start_date

    day_count = _planner_inclusive_day_count(start_date, end_date)
    end_date = _planner_range_end(start_date, day_count)
    return start_date, end_date, day_count


def _planner_load() -> dict:
    path = _planner_file_path()
    if not os.path.exists(path):
        return _planner_default_plan(7)

    try:
        with open(path, "r", encoding="utf-8", errors="replace") as file:
            data = json.load(file)
    except Exception:
        return _planner_default_plan(7)

    if not isinstance(data, dict):
        data = {}

    raw_days = data.get("days", [])
    if isinstance(raw_days, list) and raw_days:
        default_count = len(raw_days)
    else:
        default_count = 7
    start_date, end_date, day_count = _planner_plan_range(data, raw_days, default_count)
    days = _planner_normalize_days(raw_days, day_count, start_date)
    return {
        "day_count": day_count,
        "start_date": _planner_date_to_iso(start_date),
        "end_date": _planner_date_to_iso(end_date),
        "days": days,
    }


def _planner_save(plan: dict) -> None:
    raw_days = plan.get("days", [])
    default_count = len(raw_days) if isinstance(raw_days, list) and raw_days else plan.get("day_count") or 7
    start_date, end_date, day_count = _planner_plan_range(plan, raw_days if isinstance(raw_days, list) else [], default_count)
    days = _planner_normalize_days(raw_days, day_count, start_date)
    payload = {
        "day_count": day_count,
        "start_date": _planner_date_to_iso(start_date),
        "end_date": _planner_date_to_iso(end_date),
        "updated_at": time.strftime("%Y-%m-%d %H:%M:%S"),
        "days": days,
    }
    path = _planner_file_path()
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8") as file:
        json.dump(payload, file, ensure_ascii=False, indent=2)


def _planner_used_cipher_names(plan: dict | None = None) -> list[str]:
    if plan is None:
        plan = _planner_load()
    used = []
    seen = set()
    for day in plan.get("days", []):
        if not isinstance(day, dict):
            continue
        for _key, _label, segment in _planner_day_segments(day):
            for cipher in segment.get("ciphers", []):
                cipher_name = str(cipher or "").strip()
                if cipher_name and cipher_name not in seen:
                    seen.add(cipher_name)
                    used.append(cipher_name)
    return used


def _planner_used_count() -> int:
    return len(_planner_used_cipher_names())


def _planner_day_indexes(plan: dict, day_indexes: list[int] | None = None) -> list[int]:
    days = plan.get("days", [])
    count = len(days) if isinstance(days, list) else 0
    if not day_indexes:
        return list(range(1, count + 1))

    result = []
    seen = set()
    for raw_index in day_indexes:
        try:
            index = int(raw_index)
        except Exception:
            continue
        if 1 <= index <= count and index not in seen:
            seen.add(index)
            result.append(index)
    return result


def _planner_escape_html(value: str) -> str:
    import html

    return html.escape(str(value or "")).replace("\n", "<br>")


def _planner_segment_label(segment_key: str) -> str:
    for key, label in _PLANNER_SEGMENTS:
        if key == segment_key:
            return label
    return "Celý den"


def _planner_day_segments(day: dict) -> list[tuple[str, str, dict]]:
    segments = day.get("segments", {})
    if not isinstance(segments, dict):
        segments = {}
    result = []
    for key, label in _PLANNER_SEGMENTS:
        segment = segments.get(key)
        if not isinstance(segment, dict):
            segment = _planner_empty_segment()
        result.append((key, label, segment))
    return result


def _planner_day_segment(day: dict, segment_key: str) -> dict:
    segments = day.setdefault("segments", _planner_default_segments())
    if not isinstance(segments, dict):
        segments = _planner_default_segments()
        day["segments"] = segments
    if segment_key not in segments or not isinstance(segments.get(segment_key), dict):
        segments[segment_key] = _planner_empty_segment()
    return segments[segment_key]


def _planner_day_ciphers(day: dict) -> list[str]:
    ciphers = []
    for _key, _label, segment in _planner_day_segments(day):
        ciphers.extend(str(cipher) for cipher in segment.get("ciphers", []) if str(cipher).strip())
    return ciphers


def _planner_day_has_encrypted_text(day: dict) -> bool:
    return any(str(segment.get("encrypted_text") or "").strip() for _key, _label, segment in _planner_day_segments(day))


def _planner_day_attachment_count(day: dict) -> int:
    total = 0
    for _key, _label, segment in _planner_day_segments(day):
        total += len(_planner_normalize_attachments(segment.get("attachments", [])))
    return total


def _planner_day_first_note(day: dict) -> str:
    for _key, label, segment in _planner_day_segments(day):
        note = str(segment.get("note") or "").strip()
        if note:
            return f"{label}: {note}"
    return ""


def _planner_day_heading(day: dict, index: int) -> str:
    date_label = _planner_format_date(day.get("date"))
    weekday = str(day.get("weekday") or _planner_weekday_name(day.get("date")) or "").strip()
    if date_label and weekday:
        return f"Den {index} ({date_label}, {weekday})"
    if date_label:
        return f"Den {index} ({date_label})"
    return f"Den {index}"


def _planner_qdate_to_date(value: QDate) -> date:
    return date(value.year(), value.month(), value.day())


def _planner_date_to_qdate(value) -> QDate:
    day_date = _planner_parse_date(value, date.today()) or date.today()
    return QDate(day_date.year, day_date.month, day_date.day)


def _planner_export_text(plan: dict | None = None, day_indexes: list[int] | None = None) -> str:
    if plan is None:
        plan = _planner_load()

    known = _planner_known_cipher_names()
    used = _planner_used_cipher_names(plan)
    unused = [name for name in known if name not in set(used)]
    selected_indexes = _planner_day_indexes(plan, day_indexes)
    selected_label = "vybrané dny" if day_indexes else "všechny dny"
    lines = [
        "Plán šifer na tábor",
        "=" * 40,
        "",
        f"Použito: {len(used)} / {len(known)} šifer",
        f"Výpis: {selected_label}",
        "",
    ]

    for index in selected_indexes:
        day = plan.get("days", [])[index - 1]
        name = str(day.get("name") or f"Den {index}")
        lines.append(_planner_day_heading(day, index))
        if name and name != f"Den {index}":
            lines.append(f"  Název dne: {name}")

        for _segment_key, segment_label, segment in _planner_day_segments(day):
            note = str(segment.get("note") or "").strip()
            encrypted_text = str(segment.get("encrypted_text") or "").strip()
            ciphers = [str(cipher) for cipher in segment.get("ciphers", []) if str(cipher).strip()]
            attachments = _planner_normalize_attachments(segment.get("attachments", []))

            lines.append(f"  {segment_label}:")
            if note:
                lines.append("    Poznámka:")
                for note_line in note.splitlines():
                    lines.append(f"      {note_line}")
            if ciphers:
                lines.append("    Šifry:")
                for cipher in ciphers:
                    lines.append(f"      - {cipher}")
            else:
                lines.append("    Šifry: zatím žádná")
            if encrypted_text:
                lines.append("    Zašifrovaný text:")
                for output_line in encrypted_text.splitlines():
                    lines.append(f"      {output_line}")
            if attachments:
                lines.append("    Přílohy:")
                for attachment in attachments:
                    lines.append(f"      - {attachment['title']} ({attachment['path']})")
        lines.append("")

    if not day_indexes:
        lines.append("Ještě nepoužité šifry")
        lines.append("-" * 40)
        if unused:
            lines.extend(f"- {name}" for name in unused)
        else:
            lines.append("Všechny dostupné šifry už jsou v plánu.")

    return "\n".join(lines).strip() + "\n"


def _planner_print_html(plan: dict, day_indexes: list[int] | None = None) -> str:
    indexes = _planner_day_indexes(plan, day_indexes)
    if not indexes:
        return "<html><body><p>Nejsou vybrané žádné dny.</p></body></html>"

    blocks = []
    for position, index in enumerate(indexes):
        day = plan.get("days", [])[index - 1]
        name = str(day.get("name") or f"Den {index}")
        page_break = "page-break-before: always;" if position else ""
        segment_blocks = []
        for _segment_key, segment_label, segment in _planner_day_segments(day):
            note = str(segment.get("note") or "").strip()
            encrypted_text = str(segment.get("encrypted_text") or "").strip()
            ciphers = [str(cipher) for cipher in segment.get("ciphers", []) if str(cipher).strip()]
            attachments = _planner_normalize_attachments(segment.get("attachments", []))
            cipher_items = "".join(f"<li>{_planner_escape_html(cipher)}</li>" for cipher in ciphers)
            if not cipher_items:
                cipher_items = "<li>Zatím žádná šifra.</li>"
            attachment_items = "".join(
                f"<li>{_planner_escape_html(attachment['title'])}<br><span class='path'>{_planner_escape_html(attachment['path'])}</span></li>"
                for attachment in attachments
            )
            if not attachment_items:
                attachment_items = "<li>Bez příloh.</li>"

            note_html = (
                f"<div class='label'>Poznámka / program</div><div class='note'>{_planner_escape_html(note)}</div>"
                if note else
                "<div class='label'>Poznámka / program</div><div class='muted'>Bez poznámky.</div>"
            )
            encrypted_html = (
                "<div class='label'>Zašifrovaný text</div>"
                f"<div class='encrypted'>{_planner_escape_html(encrypted_text)}</div>"
                if encrypted_text else
                "<div class='label'>Zašifrovaný text</div><div class='muted'>Bez vloženého textu.</div>"
            )
            segment_blocks.append(
                "<div class='segment'>"
                f"<h2>{_planner_escape_html(segment_label)}</h2>"
                f"{note_html}"
                "<div class='label'>Šifry</div>"
                f"<ol>{cipher_items}</ol>"
                f"{encrypted_html}"
                "<div class='label'>Přílohy</div>"
                f"<ol>{attachment_items}</ol>"
                "</div>"
            )

        custom_name = (
            f"<div class='day-name'>Název dne: {_planner_escape_html(name)}</div>"
            if name and name != f"Den {index}" else
            ""
        )
        blocks.append(
            f"<section class='day' style='{page_break}'>"
            "<div class='day-number'>Plán dne</div>"
            f"<h1>{_planner_escape_html(_planner_day_heading(day, index))}</h1>"
            f"{custom_name}"
            + "".join(segment_blocks)
            + "</section>"
        )

    return (
        "<html><head><meta charset='utf-8'>"
        "<style>"
        "body{font-family:Georgia,serif;color:#111;background:#fff;font-size:12pt;}"
        ".day{border:1px solid #999;border-radius:8px;padding:18px;margin:0 0 18px 0;}"
        ".day-number{font-size:11pt;color:#666;text-transform:uppercase;letter-spacing:1px;}"
        "h1{font-size:24pt;margin:4px 0 14px 0;color:#10223a;}"
        "h2{font-size:16pt;margin:0 0 8px 0;color:#10223a;}"
        ".day-name{font-size:12pt;color:#555;margin:-8px 0 14px 0;}"
        ".segment{border-top:1px solid #ccc;padding-top:10px;margin-top:12px;}"
        ".label{font-weight:bold;margin:14px 0 5px 0;border-bottom:1px solid #bbb;padding-bottom:3px;}"
        ".note{white-space:pre-wrap;line-height:1.35;}"
        ".encrypted{white-space:pre-wrap;font-family:Consolas,monospace;font-size:11pt;line-height:1.35;border:1px solid #bbb;padding:10px;}"
        ".muted{color:#777;font-style:italic;}"
        ".path{font-size:9pt;color:#666;}"
        "ol{margin-top:8px;line-height:1.45;}"
        "li{margin-bottom:4px;}"
        "</style></head><body>"
        + "".join(blocks)
        + "</body></html>"
    )


def _planner_build_print_document(plan: dict, day_indexes: list[int] | None = None, paper_name: str = "A4", orientation_name: str = "Na výšku"):
    from PySide6.QtGui import QTextDocument

    document = QTextDocument()
    document.setDefaultFont(QFont("Georgia", 12))
    document.setPageSize(_print_page_size_points(paper_name, orientation_name))
    document.setHtml(_planner_print_html(plan, day_indexes))
    return document


def _planner_write_pdf(path: str, plan: dict, day_indexes: list[int] | None = None, paper_name: str = "A4", orientation_name: str = "Na výšku") -> None:
    writer = _pdf_writer_for_path(path, paper_name, orientation_name)
    document = _planner_build_print_document(plan, day_indexes, paper_name, orientation_name)
    document.print_(writer)


def _planner_print_days(parent, plan: dict, indexes: list[int], title: str = "Tisk vybraných dnů") -> None:
    try:
        from PySide6.QtPrintSupport import QPrinter, QPrintDialog
    except Exception as error:
        QMessageBox.warning(parent, "Tisk není dostupný", f"Nepodařilo se načíst podporu tisku:\n{error}")
        return

    printer = QPrinter(QPrinter.HighResolution)
    _print_apply_page_setup_to_printer(printer, "A4", "Na výšku")
    dialog = QPrintDialog(printer, parent)
    dialog.setWindowTitle(title)
    if dialog.exec() != QDialog.Accepted:
        return
    document = _planner_build_print_document(plan, indexes, "A4", "Na výšku")
    document.print_(printer)


def _planner_docx_read_text(path: str) -> str:
    namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    try:
        with zipfile.ZipFile(path, "r") as archive:
            xml_data = archive.read("word/document.xml")
    except Exception as error:
        return f"Soubor DOCX se nepodařilo načíst:\n{error}"

    try:
        root = ET.fromstring(xml_data)
    except Exception as error:
        return f"Text dokumentu se nepodařilo přečíst:\n{error}"

    paragraphs = []
    for paragraph in root.iter(f"{namespace}p"):
        pieces = []
        for text_node in paragraph.iter(f"{namespace}t"):
            pieces.append(text_node.text or "")
        if pieces:
            paragraphs.append("".join(pieces))
    return "\n".join(paragraphs).strip()


def _planner_docx_write_text(path: str, text: str) -> None:
    namespace = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    ET.register_namespace("w", namespace)
    ns = f"{{{namespace}}}"

    with zipfile.ZipFile(path, "r") as source:
        document_xml = source.read("word/document.xml")
        root = ET.fromstring(document_xml)
        body = root.find(f"{ns}body")
        if body is None:
            raise ValueError("V DOCX chybí word/document.xml/body.")

        section_properties = body.find(f"{ns}sectPr")
        body.clear()
        for line in str(text or "").splitlines() or [""]:
            paragraph = ET.SubElement(body, f"{ns}p")
            run = ET.SubElement(paragraph, f"{ns}r")
            text_node = ET.SubElement(run, f"{ns}t")
            if line.startswith(" ") or line.endswith(" "):
                text_node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            text_node.text = line
        if section_properties is not None:
            body.append(section_properties)

        new_document_xml = ET.tostring(root, encoding="utf-8", xml_declaration=True)
        temp_path = path + ".tmp"
        with zipfile.ZipFile(temp_path, "w", zipfile.ZIP_DEFLATED) as target:
            for item in source.infolist():
                if item.filename == "word/document.xml":
                    target.writestr(item, new_document_xml)
                else:
                    target.writestr(item, source.read(item.filename))

    os.replace(temp_path, path)


def _planner_read_text_file(path: str) -> str:
    try:
        with open(path, "r", encoding="utf-8", errors="replace") as file:
            return file.read()
    except Exception as error:
        return f"Soubor se nepodařilo načíst:\n{error}"


def _planner_write_text_file(path: str, text: str) -> None:
    with open(path, "w", encoding="utf-8") as file:
        file.write(text)


def _planner_encrypt_text(cipher_name: str, text: str, owner_window=None) -> str:
    cipher_name = str(cipher_name or "").strip()
    if not cipher_name:
        return "Nejdřív vyber šifru."

    logic = get_cipher_logic(cipher_name)
    if logic is None or not hasattr(logic, "encrypt"):
        return f"Tahle šifra zatím nemá dostupné šifrování:\n{cipher_name}"

    if cipher_name == "Caesarova šifra":
        shift = 3
        direction = "dopredu"
        central = getattr(owner_window, "central", None)
        if central is not None:
            try:
                shift = int(central.get_caesar_shift())
                direction = central.get_caesar_direction()
            except Exception:
                pass
        signed_shift = -shift if direction == "dozadu" else shift
        return logic.encrypt(text, signed_shift)

    return logic.encrypt(text)


class PlannerAttachmentDialog(QDialog):
    """Prohlížení a základní editace přílohy uložené u dne v plánovači."""

    def __init__(self, attachment: dict, parent=None):
        super().__init__(parent)
        self.attachment = _planner_normalize_attachment(attachment) or {}
        self.path = self.attachment.get("path", "")
        self.kind = _planner_attachment_kind(self.path)
        self.pdf_document = None

        self.setWindowTitle(f"Příloha - {_planner_attachment_title(self.path, self.attachment.get('title'))}")
        self.resize(980, 720)
        self.setMinimumSize(760, 520)
        self.setStyleSheet("""
            QDialog { background-color: #0a1626; color: #f3ddaa; }
            QLabel { color: #ead8b3; background: transparent; }
            QTextEdit {
                color: #f0e2c0;
                background: #071018;
                border: 1px solid #8a6938;
                border-radius: 6px;
                padding: 8px;
                font-family: Consolas, Menlo, Monaco, monospace;
                font-size: 12px;
            }
            QPushButton {
                color: #f6e7bf;
                background-color: #10263e;
                border: 1px solid #c49344;
                border-radius: 8px;
                padding: 8px 12px;
                font-weight: bold;
            }
            QPushButton:hover { background-color: #144a63; }
        """)

        layout = QVBoxLayout(self)
        layout.setContentsMargins(12, 12, 12, 12)
        layout.setSpacing(8)

        title = QLabel(_planner_attachment_title(self.path, self.attachment.get("title")), self)
        title.setStyleSheet("color: #f3d79a; font-size: 17px; font-weight: bold;")
        layout.addWidget(title)

        self.path_label = QLabel(self.path, self)
        self.path_label.setWordWrap(True)
        self.path_label.setStyleSheet("color: #a8a295; font-size: 11px;")
        layout.addWidget(self.path_label)

        self.info_label = QLabel(self)
        self.info_label.setWordWrap(True)
        layout.addWidget(self.info_label)

        self.editor = QTextEdit(self)
        self.editor.setLineWrapMode(QTextEdit.WidgetWidth)
        self.editor.setWordWrapMode(QTextOption.WrapAnywhere)
        layout.addWidget(self.editor, 1)

        self.pdf_container = None
        self.load_content(layout)

        buttons = QHBoxLayout()
        buttons.addStretch(1)
        self.open_system_button = QPushButton("Otevřít v systému", self)
        self.save_button = QPushButton("Uložit změny", self)
        self.close_button = QPushButton("Zavřít", self)
        buttons.addWidget(self.open_system_button)
        buttons.addWidget(self.save_button)
        buttons.addWidget(self.close_button)
        layout.addLayout(buttons)

        self.open_system_button.clicked.connect(self.open_in_system)
        self.save_button.clicked.connect(self.save_changes)
        self.close_button.clicked.connect(self.close)

    def load_content(self, layout: QVBoxLayout):
        if not self.path or not os.path.exists(self.path):
            self.info_label.setText("Soubor už na této cestě neexistuje.")
            self.editor.setReadOnly(True)
            self.editor.setPlainText("")
            return

        if self.kind == "pdf":
            self.load_pdf(layout)
            return

        if self.kind == "docx":
            self.info_label.setText("DOCX se v Šifrátoru otevře jako text. Uložení upraví text dokumentu bez zachování složitého formátování.")
            self.editor.setPlainText(_planner_docx_read_text(self.path))
            return

        if self.kind in ("txt", "md", "rtf"):
            self.info_label.setText("Textovou přílohu můžeš upravit přímo tady.")
            self.editor.setPlainText(_planner_read_text_file(self.path))
            return

        self.info_label.setText("Tento typ souboru nejde bezpečně editovat přímo v Šifrátoru. Otevři ho v systémové aplikaci.")
        self.editor.setReadOnly(True)
        self.editor.setPlainText("Pro starší .doc nebo jiné binární dokumenty použij tlačítko Otevřít v systému.")

    def load_pdf(self, layout: QVBoxLayout):
        self.info_label.setText("PDF se dá v Šifrátoru prohlížet. Pro úpravy použij editor PDF v systému.")
        self.editor.hide()
        try:
            from PySide6.QtPdf import QPdfDocument
            from PySide6.QtPdfWidgets import QPdfView

            self.pdf_document = QPdfDocument(self)
            self.pdf_document.load(self.path)
            pdf_view = QPdfView(self)
            pdf_view.setDocument(self.pdf_document)
            pdf_view.setZoomMode(QPdfView.ZoomMode.FitToWidth)
            layout.addWidget(pdf_view, 1)
            self.save_button = None
        except Exception as error:
            self.editor.show()
            self.editor.setReadOnly(True)
            self.editor.setPlainText(f"Interní PDF náhled není dostupný:\n{error}\n\nPoužij Otevřít v systému.")

    def open_in_system(self):
        if self.path:
            QDesktopServices.openUrl(QUrl.fromLocalFile(self.path))

    def save_changes(self):
        if self.kind == "pdf":
            QMessageBox.information(self, "Příloha", "PDF se tady prohlíží, ale neupravuje. Použij editor PDF v systému.")
            return
        if self.kind == "doc":
            QMessageBox.information(self, "Příloha", "Starý .doc formát nejde bezpečně uložit přímo v Šifrátoru.")
            return
        try:
            if self.kind == "docx":
                _planner_docx_write_text(self.path, self.editor.toPlainText())
            elif self.kind in ("txt", "md", "rtf"):
                _planner_write_text_file(self.path, self.editor.toPlainText())
            else:
                QMessageBox.information(self, "Příloha", "Tento typ souboru nejde uložit přímo v Šifrátoru.")
                return
            QMessageBox.information(self, "Příloha", "Změny byly uloženy.")
        except Exception as error:
            QMessageBox.warning(self, "Příloha", f"Soubor se nepodařilo uložit:\n{error}")


class PlannerDayDetailDialog(QDialog):
    """Detail jednoho dne se všemi částmi, přílohami a tiskem."""

    def __init__(self, planner, day_index: int):
        super().__init__(planner)
        self.planner = planner
        self.day_index = day_index
        self.day = planner.day(day_index)
        self._updating = False
        self.note_edits = {}
        self.encrypted_edits = {}
        self.attachment_lists = {}

        self.setWindowTitle(_planner_day_heading(self.day, self.day_index))
        self.resize(1080, 760)
        self.setMinimumSize(820, 560)
        self.setStyleSheet("""
            QDialog { background-color: #0a1626; color: #f3ddaa; }
            QLabel { color: #ead8b3; background: transparent; }
            QTabWidget::pane { border: 1px solid #8a6938; border-radius: 6px; }
            QTabBar::tab {
                color: #f0e2c0;
                background: #10263e;
                border: 1px solid #8a6938;
                padding: 7px 12px;
            }
            QTabBar::tab:selected { background: #144a63; color: #fff2cc; }
            QTextEdit, QListWidget {
                color: #f0e2c0;
                background: #071018;
                border: 1px solid #8a6938;
                border-radius: 6px;
                padding: 6px;
            }
            QTextEdit {
                font-family: Consolas, Menlo, Monaco, monospace;
                font-size: 12px;
            }
            QListWidget::item { padding: 4px; }
            QListWidget::item:selected { background: #144a63; color: #fff2cc; }
            QPushButton {
                color: #f6e7bf;
                background-color: #10263e;
                border: 1px solid #c49344;
                border-radius: 8px;
                padding: 8px 12px;
                font-weight: bold;
            }
            QPushButton:hover { background-color: #144a63; }
        """)

        layout = QVBoxLayout(self)
        layout.setContentsMargins(12, 12, 12, 12)
        layout.setSpacing(9)

        title = QLabel(_planner_day_heading(self.day, self.day_index), self)
        title.setStyleSheet("color: #f3d79a; font-size: 19px; font-weight: bold;")
        layout.addWidget(title)

        name = str(self.day.get("name") or f"Den {self.day_index}").strip()
        if name and name != f"Den {self.day_index}":
            name_label = QLabel(f"Název dne: {name}", self)
            name_label.setWordWrap(True)
            layout.addWidget(name_label)

        self.tabs = QTabWidget(self)
        layout.addWidget(self.tabs, 1)

        self.summary_edit = QTextEdit(self)
        self.summary_edit.setReadOnly(True)
        self.tabs.addTab(self.summary_edit, "Souhrn")

        for segment_key, segment_label, _segment in _planner_day_segments(self.day):
            self.tabs.addTab(self.create_segment_tab(segment_key), segment_label)

        buttons = QHBoxLayout()
        buttons.addStretch(1)
        self.print_day_button = QPushButton("Tisknout tento den", self)
        self.save_button = QPushButton("Uložit změny", self)
        self.close_button = QPushButton("Zavřít", self)
        buttons.addWidget(self.print_day_button)
        buttons.addWidget(self.save_button)
        buttons.addWidget(self.close_button)
        layout.addLayout(buttons)

        self.print_day_button.clicked.connect(self.print_day)
        self.save_button.clicked.connect(self.save_changes)
        self.close_button.clicked.connect(self.close)

        self.refresh_summary()

    def create_segment_tab(self, segment_key: str) -> QWidget:
        segment = _planner_day_segment(self.day, segment_key)
        tab = QWidget(self)
        layout = QVBoxLayout(tab)
        layout.setContentsMargins(10, 10, 10, 10)
        layout.setSpacing(8)

        ciphers = [str(cipher) for cipher in segment.get("ciphers", []) if str(cipher).strip()]
        ciphers_edit = QTextEdit(tab)
        ciphers_edit.setReadOnly(True)
        ciphers_edit.setMaximumHeight(72)
        ciphers_edit.setPlainText("\n".join(f"{i}. {cipher}" for i, cipher in enumerate(ciphers, start=1)) if ciphers else "Zatím žádná šifra.")
        layout.addWidget(QLabel("Šifry:", tab))
        layout.addWidget(ciphers_edit)

        note_edit = QTextEdit(tab)
        note_edit.setMinimumHeight(95)
        note_edit.setPlaceholderText("Poznámka / program této části dne.")
        note_edit.setPlainText(str(segment.get("note") or ""))
        self.note_edits[segment_key] = note_edit
        layout.addWidget(QLabel("Poznámka / program:", tab))
        layout.addWidget(note_edit)

        encrypted_edit = QTextEdit(tab)
        encrypted_edit.setMinimumHeight(120)
        encrypted_edit.setPlaceholderText("Zašifrovaný text této části dne.")
        encrypted_edit.setPlainText(str(segment.get("encrypted_text") or ""))
        self.encrypted_edits[segment_key] = encrypted_edit
        layout.addWidget(QLabel("Zašifrovaný text:", tab))
        layout.addWidget(encrypted_edit)

        layout.addWidget(QLabel("Přílohy a náhledy:", tab))
        attachment_list = QListWidget(tab)
        attachment_list.setMaximumHeight(110)
        self.attachment_lists[segment_key] = attachment_list
        layout.addWidget(attachment_list)

        attachment_buttons = QHBoxLayout()
        add_button = QPushButton("Přidat PDF/Word", tab)
        open_button = QPushButton("Náhled/upravit", tab)
        remove_button = QPushButton("Odebrat", tab)
        attachment_buttons.addWidget(add_button)
        attachment_buttons.addWidget(open_button)
        attachment_buttons.addWidget(remove_button)
        attachment_buttons.addStretch(1)
        layout.addLayout(attachment_buttons)

        note_edit.textChanged.connect(lambda key=segment_key: self.segment_text_changed(key))
        encrypted_edit.textChanged.connect(lambda key=segment_key: self.segment_text_changed(key))
        add_button.clicked.connect(lambda _checked=False, key=segment_key: self.add_attachment(key))
        open_button.clicked.connect(lambda _checked=False, key=segment_key: self.open_attachment(key))
        remove_button.clicked.connect(lambda _checked=False, key=segment_key: self.remove_attachment(key))
        attachment_list.itemDoubleClicked.connect(lambda _item, key=segment_key: self.open_attachment(key))

        self.refresh_attachment_list(segment_key)
        return tab

    def segment_text_changed(self, segment_key: str):
        if self._updating:
            return
        segment = _planner_day_segment(self.day, segment_key)
        segment["note"] = self.note_edits[segment_key].toPlainText().strip()
        segment["encrypted_text"] = self.encrypted_edits[segment_key].toPlainText().strip()
        self.refresh_summary()

    def refresh_summary(self):
        self.summary_edit.setPlainText(_planner_export_text(self.planner.plan, [self.day_index]).strip())

    def refresh_attachment_list(self, segment_key: str):
        attachment_list = self.attachment_lists.get(segment_key)
        if attachment_list is None:
            return
        segment = _planner_day_segment(self.day, segment_key)
        attachments = _planner_normalize_attachments(segment.get("attachments", []))
        attachment_list.clear()
        for index, attachment in enumerate(attachments):
            item = QListWidgetItem(f"{attachment['title']}  [{attachment['kind'].upper()}]")
            item.setToolTip(attachment["path"])
            item.setData(Qt.UserRole, index)
            attachment_list.addItem(item)

    def selected_attachment_index(self, segment_key: str) -> int:
        attachment_list = self.attachment_lists.get(segment_key)
        if attachment_list is None:
            return -1
        item = attachment_list.currentItem()
        if item is None:
            return -1
        try:
            return int(item.data(Qt.UserRole))
        except Exception:
            return -1

    def add_attachment(self, segment_key: str):
        path, _ = QFileDialog.getOpenFileName(
            self,
            "Přidat PDF / Word přílohu",
            "",
            "Dokumenty (*.pdf *.docx *.doc *.txt *.md *.rtf);;Všechny soubory (*.*)",
        )
        if not path:
            return

        attachment = _planner_import_attachment(path)
        if attachment is None:
            return

        segment = _planner_day_segment(self.day, segment_key)
        attachments = _planner_normalize_attachments(segment.get("attachments", []))
        new_key = _planner_attachment_identity(attachment)
        if any(_planner_attachment_identity(item) == new_key for item in attachments):
            QMessageBox.information(self, "Příloha", "Tahle příloha už je v této části dne.")
            return

        attachments.append(attachment)
        segment["attachments"] = attachments
        self.refresh_attachment_list(segment_key)
        self.refresh_summary()
        self.sync_parent(save=True)

    def open_attachment(self, segment_key: str):
        segment = _planner_day_segment(self.day, segment_key)
        attachments = _planner_normalize_attachments(segment.get("attachments", []))
        index = self.selected_attachment_index(segment_key)
        if not 0 <= index < len(attachments):
            QMessageBox.information(self, "Příloha", "Nejdřív vyber přílohu ze seznamu.")
            return

        dialog = PlannerAttachmentDialog(attachments[index], self)
        dialog.exec()

    def remove_attachment(self, segment_key: str):
        segment = _planner_day_segment(self.day, segment_key)
        attachments = _planner_normalize_attachments(segment.get("attachments", []))
        index = self.selected_attachment_index(segment_key)
        if not 0 <= index < len(attachments):
            return

        del attachments[index]
        segment["attachments"] = attachments
        self.refresh_attachment_list(segment_key)
        self.refresh_summary()
        self.sync_parent(save=True)

    def sync_parent(self, save: bool = True):
        self.planner.refresh_editor()
        self.planner.refresh_calendar_cards()
        self.planner.refresh_overview(save=save)

    def save_changes(self):
        for segment_key in self.note_edits:
            segment = _planner_day_segment(self.day, segment_key)
            segment["note"] = self.note_edits[segment_key].toPlainText().strip()
            segment["encrypted_text"] = self.encrypted_edits[segment_key].toPlainText().strip()
        self.refresh_summary()
        self.sync_parent(save=True)

    def print_day(self):
        self.save_changes()
        _planner_print_days(self, self.planner.plan, [self.day_index], f"Tisk - {_planner_day_heading(self.day, self.day_index)}")

    def closeEvent(self, event):
        self.save_changes()
        super().closeEvent(event)


class CampPlannerDialog(PirateModuleDialog):
    """Kalendářový plánovač šifer na jednotlivé dny tábora."""

    def __init__(self, owner_window):
        super().__init__(
            owner_window,
            "planner_BG.png",
            (
                (0.117, 0.223, 0.82),
                (0.227, 0.524, 0.36),
                (0.307, 0.539, 0.32),
                (0.957, 0.835, 1.18),
            ),
        )
        self.owner_window = owner_window
        self.available_names = _planner_known_cipher_names()
        self.plan = _planner_load()
        self.calendar_cards = []
        self.selected_day_indexes = set()
        self.current_day_index = 1
        self.current_segment_key = _PLANNER_DEFAULT_SEGMENT_KEY
        self._rebuilding = False
        self._updating_editor = False
        self._updating_dates = False

        self.setWindowTitle("Kalendář šifer na dny tábora")
        self.resize(1260, 780)
        self.setMinimumSize(980, 620)
        self.setStyleSheet("""
            QDialog {
                background-color: #0a1626;
                color: #f3ddaa;
            }
            QLabel {
                color: #ead8b3;
                background: transparent;
            }
            QLineEdit, QComboBox, QTextEdit, QSpinBox, QDateEdit {
                color: #f0e2c0;
                background: #071018;
                border: 1px solid #8a6938;
                border-radius: 6px;
                padding: 6px;
            }
            QTextEdit {
                font-family: Consolas, Menlo, Monaco, monospace;
                font-size: 12px;
            }
            QPushButton {
                color: #f6e7bf;
                background-color: #10263e;
                border: 1px solid #c49344;
                border-radius: 8px;
                padding: 8px 12px;
                font-weight: bold;
            }
            QPushButton:hover {
                background-color: #144a63;
            }
            QCheckBox {
                color: #ead8b3;
                background: transparent;
                spacing: 7px;
            }
        """)

        layout = QVBoxLayout(self)
        layout.setContentsMargins(12, 12, 12, 12)
        layout.setSpacing(10)

        title = QLabel("Kalendář šifer na tábor", self)
        title.setStyleSheet("color: #f3d79a; font-size: 20px; font-weight: bold;")
        layout.addWidget(title)

        hint = QLabel(
            "Vyber začátek a konec tábora, klikni na den v kalendáři a uprav dopoledne, odpoledne nebo celý den. Zaškrtni dny, které chceš tisknout nebo exportovat do PDF.",
            self,
        )
        hint.setWordWrap(True)
        layout.addWidget(hint)

        start_date = _planner_parse_date(self.plan.get("start_date"), date.today()) or date.today()
        end_date = _planner_parse_date(self.plan.get("end_date"), _planner_range_end(start_date, self.plan.get("day_count") or 7))
        if end_date is None or end_date < start_date:
            end_date = start_date

        controls = QHBoxLayout()
        controls.addWidget(QLabel("Od:", self))
        self.start_date_edit = QDateEdit(self)
        self.start_date_edit.setCalendarPopup(True)
        self.start_date_edit.setDisplayFormat("d. M. yyyy")
        self.start_date_edit.setDate(QDate(start_date.year, start_date.month, start_date.day))
        controls.addWidget(self.start_date_edit)

        controls.addWidget(QLabel("Do:", self))
        self.end_date_edit = QDateEdit(self)
        self.end_date_edit.setCalendarPopup(True)
        self.end_date_edit.setDisplayFormat("d. M. yyyy")
        self.end_date_edit.setDate(QDate(end_date.year, end_date.month, end_date.day))
        controls.addWidget(self.end_date_edit)

        controls.addWidget(QLabel("Počet dnů:", self))
        self.day_count_spin = QSpinBox(self)
        self.day_count_spin.setRange(1, _PLANNER_MAX_DAYS)
        self.day_count_spin.setValue(_planner_inclusive_day_count(start_date, end_date))
        self.day_count_spin.setSuffix(" dnů")
        self.day_count_spin.setReadOnly(True)
        self.day_count_spin.setButtonSymbols(QAbstractSpinBox.NoButtons)
        self.day_count_spin.setFocusPolicy(Qt.NoFocus)
        controls.addWidget(self.day_count_spin)

        self.prevent_repeat_check = QCheckBox("Hlídat opakování šifer", self)
        self.prevent_repeat_check.setChecked(True)
        controls.addWidget(self.prevent_repeat_check)

        self.select_all_button = QPushButton("Vybrat vše", self)
        self.select_none_button = QPushButton("Zrušit výběr", self)
        controls.addWidget(self.select_all_button)
        controls.addWidget(self.select_none_button)
        controls.addStretch(1)
        layout.addLayout(controls)

        body = QHBoxLayout()
        body.setSpacing(12)
        layout.addLayout(body, 1)

        self.calendar_scroll = QScrollArea(self)
        self.calendar_scroll.setWidgetResizable(True)
        self.calendar_scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
        self.calendar_scroll.setVerticalScrollBarPolicy(Qt.ScrollBarAsNeeded)
        self.calendar_scroll.setStyleSheet("""
            QScrollArea {
                background: transparent;
                border: 1px solid #8a6938;
                border-radius: 8px;
            }
        """)
        self.calendar_widget = QWidget(self.calendar_scroll)
        self.calendar_layout = QGridLayout(self.calendar_widget)
        self.calendar_layout.setContentsMargins(10, 10, 10, 10)
        self.calendar_layout.setHorizontalSpacing(10)
        self.calendar_layout.setVerticalSpacing(10)
        self.calendar_scroll.setWidget(self.calendar_widget)
        self.calendar_scroll.viewport().installEventFilter(self)
        body.addWidget(self.calendar_scroll, 3)

        right_panel = QWidget(self)
        right_panel.setObjectName("plannerRightPanel")
        right_panel.setStyleSheet("""
            QWidget#plannerRightPanel {
                background: rgba(7, 16, 24, 150);
                border: 1px solid #8a6938;
                border-radius: 8px;
            }
        """)
        right_col = QVBoxLayout(right_panel)
        right_col.setContentsMargins(12, 12, 12, 12)
        right_col.setSpacing(8)
        self.editor_scroll = QScrollArea(self)
        self.editor_scroll.setObjectName("plannerEditorScroll")
        self.editor_scroll.setWidgetResizable(True)
        self.editor_scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
        self.editor_scroll.setVerticalScrollBarPolicy(Qt.ScrollBarAsNeeded)
        self.editor_scroll.setFrameShape(QScrollArea.NoFrame)
        self.editor_scroll.setSizeAdjustPolicy(QAbstractScrollArea.AdjustIgnored)
        self.editor_scroll.setWidget(right_panel)
        body.addWidget(self.editor_scroll, 2)

        self.editor_title = QLabel(self)
        self.editor_title.setStyleSheet("color: #f3d79a; font-size: 16px; font-weight: bold;")
        right_col.addWidget(self.editor_title)

        self.open_day_detail_button = QPushButton("Detail dne", self)
        right_col.addWidget(self.open_day_detail_button)

        right_col.addWidget(QLabel("Název dne:", self))
        self.day_name_edit = QLineEdit(self)
        right_col.addWidget(self.day_name_edit)

        right_col.addWidget(QLabel("Část dne:", self))
        self.segment_combo = QComboBox(self)
        for segment_key, segment_label in _PLANNER_SEGMENTS:
            self.segment_combo.addItem(segment_label, segment_key)
        right_col.addWidget(self.segment_combo)

        right_col.addWidget(QLabel("Poznámka / program části dne:", self))
        self.day_note_edit = QTextEdit(self)
        self.day_note_edit.setMaximumHeight(100)
        self.day_note_edit.setPlaceholderText("Třeba téma části dne, stanoviště, pomůcky nebo kdo ji vede.")
        right_col.addWidget(self.day_note_edit)

        right_col.addWidget(QLabel("Šifrovat přímo v části dne:", self))
        self.planner_plain_edit = QTextEdit(self)
        self.planner_plain_edit.setMaximumHeight(76)
        self.planner_plain_edit.setPlaceholderText("Napiš text, vyber šifru a ulož výsledek do této části dne.")
        right_col.addWidget(self.planner_plain_edit)

        encrypt_row = QHBoxLayout()
        self.planner_cipher_combo = QComboBox(self)
        self.planner_cipher_combo.addItems(self.available_names)
        encrypt_row.addWidget(self.planner_cipher_combo, 1)
        self.encrypt_into_day_button = QPushButton("Zašifrovat do části", self)
        encrypt_row.addWidget(self.encrypt_into_day_button)
        right_col.addLayout(encrypt_row)

        add_row = QHBoxLayout()
        self.add_cipher_combo = QComboBox(self)
        self.add_cipher_combo.addItems(self.available_names)
        add_row.addWidget(self.add_cipher_combo, 1)
        self.add_cipher_button = QPushButton("Přidat šifru", self)
        add_row.addWidget(self.add_cipher_button)
        right_col.addLayout(add_row)

        assigned_row = QHBoxLayout()
        self.assigned_combo = QComboBox(self)
        assigned_row.addWidget(self.assigned_combo, 1)
        self.remove_cipher_button = QPushButton("Odebrat", self)
        self.clear_day_button = QPushButton("Vyčistit část", self)
        assigned_row.addWidget(self.remove_cipher_button)
        assigned_row.addWidget(self.clear_day_button)
        right_col.addLayout(assigned_row)

        self.day_ciphers_edit = QTextEdit(self)
        self.day_ciphers_edit.setReadOnly(True)
        self.day_ciphers_edit.setMaximumHeight(82)
        right_col.addWidget(self.day_ciphers_edit)

        right_col.addWidget(QLabel("Přílohy PDF / Word:", self))
        self.attachment_list = QListWidget(self)
        self.attachment_list.setMaximumHeight(84)
        self.attachment_list.setStyleSheet("""
            QListWidget {
                color: #f0e2c0;
                background: #071018;
                border: 1px solid #8a6938;
                border-radius: 6px;
                padding: 3px;
            }
            QListWidget::item { padding: 3px; }
            QListWidget::item:selected { background: #144a63; color: #fff2cc; }
        """)
        right_col.addWidget(self.attachment_list)

        attachment_row = QHBoxLayout()
        self.add_attachment_button = QPushButton("Přidat PDF/Word", self)
        self.open_attachment_button = QPushButton("Otevřít/upravit", self)
        self.remove_attachment_button = QPushButton("Odebrat", self)
        attachment_row.addWidget(self.add_attachment_button)
        attachment_row.addWidget(self.open_attachment_button)
        attachment_row.addWidget(self.remove_attachment_button)
        right_col.addLayout(attachment_row)

        encrypted_header = QHBoxLayout()
        encrypted_header.addWidget(QLabel("Zašifrovaný text:", self), 1)
        self.insert_current_button = QPushButton("Vložit z aplikace", self)
        encrypted_header.addWidget(self.insert_current_button)
        right_col.addLayout(encrypted_header)

        self.day_encrypted_edit = QTextEdit(self)
        self.day_encrypted_edit.setMaximumHeight(120)
        self.day_encrypted_edit.setPlaceholderText("Sem vlož zašifrovanou zprávu pro tuto část dne.")
        right_col.addWidget(self.day_encrypted_edit)

        self.summary_label = QLabel(self)
        self.summary_label.setStyleSheet("color: #f3d79a; font-weight: bold;")
        right_col.addWidget(self.summary_label)

        self.selection_label = QLabel(self)
        self.selection_label.setWordWrap(True)
        self.selection_label.setStyleSheet("color: #d8c392;")
        right_col.addWidget(self.selection_label)

        self.overview_edit = QTextEdit(self)
        self.overview_edit.setReadOnly(True)
        self.overview_edit.setAcceptRichText(False)
        self.overview_edit.setLineWrapMode(QTextEdit.WidgetWidth)
        self.overview_edit.setWordWrapMode(QTextOption.WrapAnywhere)
        right_col.addWidget(self.overview_edit, 1)

        self.path_label = QLabel(_planner_file_path(), self)
        self.path_label.setWordWrap(True)
        self.path_label.setStyleSheet("color: #a8a295; font-size: 11px;")
        right_col.addWidget(self.path_label)

        button_row = QHBoxLayout()
        button_row.addStretch(1)
        self.copy_button = QPushButton("Kopírovat vybrané", self)
        self.export_button = QPushButton("Export TXT", self)
        self.print_button = QPushButton("Tisk vybraných", self)
        self.pdf_button = QPushButton("PDF vybraných", self)
        self.clear_button = QPushButton("Vyčistit plán", self)
        self.close_button = QPushButton("Zavřít", self)
        button_row.addWidget(self.copy_button)
        button_row.addWidget(self.export_button)
        button_row.addWidget(self.print_button)
        button_row.addWidget(self.pdf_button)
        button_row.addWidget(self.clear_button)
        button_row.addWidget(self.close_button)
        layout.addLayout(button_row)

        self.start_date_edit.dateChanged.connect(lambda *_args: self.change_date_range())
        self.end_date_edit.dateChanged.connect(lambda *_args: self.change_date_range())
        self.select_all_button.clicked.connect(self.select_all_days)
        self.select_none_button.clicked.connect(self.clear_day_selection)
        self.open_day_detail_button.clicked.connect(lambda *_args: self.open_day_detail())
        self.day_name_edit.textChanged.connect(lambda *_args: self.rename_current_day())
        self.segment_combo.currentIndexChanged.connect(lambda *_args: self.change_current_segment())
        self.day_note_edit.textChanged.connect(self.update_current_note)
        self.encrypt_into_day_button.clicked.connect(self.encrypt_text_into_current_segment)
        self.day_encrypted_edit.textChanged.connect(self.update_current_encrypted_text)
        self.insert_current_button.clicked.connect(self.insert_current_app_result)
        self.add_cipher_button.clicked.connect(self.add_cipher_to_current_day)
        self.remove_cipher_button.clicked.connect(self.remove_cipher_from_current_day)
        self.clear_day_button.clicked.connect(self.clear_current_day)
        self.add_attachment_button.clicked.connect(self.add_attachment_to_current_segment)
        self.open_attachment_button.clicked.connect(self.open_selected_attachment)
        self.remove_attachment_button.clicked.connect(self.remove_selected_attachment)
        self.attachment_list.itemDoubleClicked.connect(lambda *_args: self.open_selected_attachment())
        self.copy_button.clicked.connect(self.copy_overview)
        self.export_button.clicked.connect(self.export_plan)
        self.print_button.clicked.connect(self.print_selected_days)
        self.pdf_button.clicked.connect(self.export_selected_pdf)
        self.clear_button.clicked.connect(self.clear_plan)
        self.close_button.clicked.connect(self.close)

        self.rebuild_calendar_cards()
        self.select_day(1, save=False)
        self.refresh_overview(save=False)
        self.apply_pirate_glass()
        self.editor_scroll.setStyleSheet(
            self.editor_scroll.styleSheet()
            + "QScrollArea#plannerEditorScroll { background: transparent; border: none; }"
        )
        self.editor_scroll.viewport().setStyleSheet("background: transparent;")

    def ensure_day_count(self):
        if hasattr(self, "start_date_edit") and hasattr(self, "end_date_edit"):
            start_date = _planner_qdate_to_date(self.start_date_edit.date())
            end_date = _planner_qdate_to_date(self.end_date_edit.date())
        else:
            start_date = _planner_parse_date(self.plan.get("start_date"), date.today()) or date.today()
            end_date = _planner_parse_date(
                self.plan.get("end_date"),
                _planner_range_end(start_date, self.plan.get("day_count") or 7),
            ) or start_date

        if end_date < start_date:
            end_date = start_date

        count = _planner_inclusive_day_count(start_date, end_date)
        end_date = _planner_range_end(start_date, count)
        self.plan["start_date"] = _planner_date_to_iso(start_date)
        self.plan["end_date"] = _planner_date_to_iso(end_date)
        self.plan["day_count"] = count

        self._updating_dates = True
        try:
            if hasattr(self, "start_date_edit"):
                self.start_date_edit.blockSignals(True)
                self.start_date_edit.setDate(_planner_date_to_qdate(start_date))
                self.start_date_edit.blockSignals(False)
            if hasattr(self, "end_date_edit"):
                self.end_date_edit.blockSignals(True)
                self.end_date_edit.setDate(_planner_date_to_qdate(end_date))
                self.end_date_edit.blockSignals(False)
            if hasattr(self, "day_count_spin"):
                self.day_count_spin.blockSignals(True)
                self.day_count_spin.setValue(count)
                self.day_count_spin.blockSignals(False)
        finally:
            self._updating_dates = False

        days = list(self.plan.get("days", []))
        if len(days) < count:
            for index in range(len(days) + 1, count + 1):
                days.append(_planner_default_day(index, start_date + timedelta(days=index - 1)))
        elif len(days) > count:
            days = days[:count]
        self.plan["days"] = _planner_normalize_days(days, count, start_date)
        self.selected_day_indexes = {index for index in self.selected_day_indexes if 1 <= index <= count}
        self.current_day_index = max(1, min(self.current_day_index, count))

    def day(self, day_index: int) -> dict:
        return self.plan["days"][day_index - 1]

    def clear_calendar_cards(self):
        while self.calendar_layout.count():
            item = self.calendar_layout.takeAt(0)
            widget = item.widget()
            if widget is not None:
                widget.deleteLater()
        self.calendar_cards = []

    def rebuild_calendar_cards(self):
        self._rebuilding = True
        self.ensure_day_count()
        self.clear_calendar_cards()

        for index, day in enumerate(self.plan.get("days", []), start=1):
            card = QWidget(self.calendar_widget)
            card.setObjectName("plannerDayCard")
            card.setMinimumSize(175, 190)
            card.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)
            card.setCursor(Qt.PointingHandCursor)
            card.setToolTip("Kliknutím vybereš den, dvojklikem otevřeš celý detail dne.")

            card_layout = QVBoxLayout(card)
            card_layout.setContentsMargins(10, 8, 10, 8)
            card_layout.setSpacing(5)

            header = QHBoxLayout()
            check = QCheckBox(f"{index}.", card)
            check.setToolTip("Zaškrtnuté dny se budou tisknout/exportovat.")
            header.addWidget(check)
            count_label = QLabel(card)
            count_label.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
            header.addWidget(count_label, 1)
            card_layout.addLayout(header)

            title_label = QLabel(_planner_day_heading(day, index), card)
            title_label.setWordWrap(True)
            title_label.setStyleSheet("font-weight: bold; color: #f3d79a;")
            card_layout.addWidget(title_label)

            note_label = QLabel(card)
            note_label.setWordWrap(True)
            note_label.setStyleSheet("color: #d8c392; font-size: 11px;")
            card_layout.addWidget(note_label)

            preview_label = QLabel(card)
            preview_label.setWordWrap(True)
            preview_label.setStyleSheet("color: #f0e2c0; font-size: 11px;")
            card_layout.addWidget(preview_label, 1)

            card_info = {
                "index": index,
                "widget": card,
                "check": check,
                "title": title_label,
                "count": count_label,
                "note": note_label,
                "preview": preview_label,
            }
            self.calendar_cards.append(card_info)

            card.mousePressEvent = lambda event, i=index: self.select_day(i)
            title_label.mousePressEvent = lambda event, i=index: self.select_day(i)
            note_label.mousePressEvent = lambda event, i=index: self.select_day(i)
            preview_label.mousePressEvent = lambda event, i=index: self.select_day(i)
            card.mouseDoubleClickEvent = lambda event, i=index: self.open_day_detail(i)
            title_label.mouseDoubleClickEvent = lambda event, i=index: self.open_day_detail(i)
            note_label.mouseDoubleClickEvent = lambda event, i=index: self.open_day_detail(i)
            preview_label.mouseDoubleClickEvent = lambda event, i=index: self.open_day_detail(i)
            check.toggled.connect(lambda checked, i=index: self.toggle_day_selection(i, checked))

        self._rebuilding = False
        self.reflow_calendar_cards()
        self.refresh_calendar_cards()

    def eventFilter(self, watched, event):
        if hasattr(self, "calendar_scroll") and watched is self.calendar_scroll.viewport() and event.type() == QEvent.Resize:
            QTimer.singleShot(0, self.reflow_calendar_cards)
        return super().eventFilter(watched, event)

    def calendar_column_count(self) -> int:
        if not self.calendar_cards:
            return 1

        margins = self.calendar_layout.contentsMargins()
        spacing = max(0, self.calendar_layout.horizontalSpacing())
        viewport_width = self.calendar_scroll.viewport().width()
        usable_width = max(150, viewport_width - margins.left() - margins.right())
        min_card_width = 185
        columns = max(1, int((usable_width + spacing) // (min_card_width + spacing)))
        return min(columns, len(self.calendar_cards))

    def reflow_calendar_cards(self):
        if not hasattr(self, "calendar_layout") or not self.calendar_cards:
            return

        while self.calendar_layout.count():
            self.calendar_layout.takeAt(0)

        columns = self.calendar_column_count()
        for position, card_info in enumerate(self.calendar_cards):
            row = position // columns
            col = position % columns
            self.calendar_layout.addWidget(card_info["widget"], row, col)

        for col in range(columns):
            self.calendar_layout.setColumnStretch(col, 1)

    def repeated_cipher_counts(self) -> dict:
        counts = {}
        for day in self.plan.get("days", []):
            for cipher in _planner_day_ciphers(day):
                counts[cipher] = counts.get(cipher, 0) + 1
        return counts

    def day_has_repeated_cipher(self, day_index: int) -> bool:
        counts = self.repeated_cipher_counts()
        return any(counts.get(cipher, 0) > 1 for cipher in _planner_day_ciphers(self.day(day_index)))

    def card_style(self, day_index: int) -> str:
        selected = day_index == self.current_day_index
        checked = day_index in self.selected_day_indexes
        repeated = self.day_has_repeated_cipher(day_index)
        border = "#f3d79a" if selected else "#d65f51" if repeated else "#0f8aa8" if checked else "#8a6938"
        background = "rgba(16, 67, 78, 142)" if selected else "rgba(13, 48, 61, 112)" if checked else "rgba(2, 13, 21, 54)"
        return (
            "QWidget#plannerDayCard {"
            f"background: {background};"
            f"border: 2px solid {border};"
            "border-radius: 8px;"
            "}"
        )

    def refresh_calendar_cards(self):
        if not hasattr(self, "calendar_cards"):
            return
        for card_info in self.calendar_cards:
            index = card_info["index"]
            day = self.day(index)
            ciphers = _planner_day_ciphers(day)
            name = str(day.get("name") or f"Den {index}").strip()
            first_note = _planner_day_first_note(day)
            has_encrypted_text = _planner_day_has_encrypted_text(day)
            attachment_count = _planner_day_attachment_count(day)
            repeated = self.day_has_repeated_cipher(index)

            card_info["widget"].setStyleSheet(self.card_style(index))
            card_info["title"].setText(_planner_day_heading(day, index))
            text_mark = " + text" if has_encrypted_text else ""
            attachment_mark = f" + {attachment_count} příl." if attachment_count else ""
            card_info["count"].setText(f"{len(ciphers)} šifer{text_mark}{attachment_mark}" + (" !" if repeated else ""))
            card_info["count"].setStyleSheet("color: #ffb0a0;" if repeated else "color: #d8c392;")
            if name and name != f"Den {index}":
                note_preview = name
            else:
                note_preview = first_note[:70] + ("..." if len(first_note) > 70 else "") if first_note else "Bez poznámky"
            card_info["note"].setText(note_preview)

            segment_lines = []
            for _segment_key, segment_label, segment in _planner_day_segments(day):
                segment_ciphers = [str(cipher) for cipher in segment.get("ciphers", []) if str(cipher).strip()]
                markers = []
                if str(segment.get("note") or "").strip():
                    markers.append("pozn.")
                if str(segment.get("encrypted_text") or "").strip():
                    markers.append("text")
                segment_attachments = _planner_normalize_attachments(segment.get("attachments", []))
                if segment_attachments:
                    markers.append(f"{len(segment_attachments)} příl.")
                suffix = f" ({', '.join(markers)})" if markers else ""
                segment_lines.append(f"{segment_label}: {len(segment_ciphers)} šifer{suffix}")

            preview = "\n".join(segment_lines)
            if ciphers:
                preview += "\n" + ", ".join(ciphers[:3])
                if len(ciphers) > 3:
                    preview += f" + {len(ciphers) - 3} další"
            card_info["preview"].setText(preview)

            check = card_info["check"]
            check.blockSignals(True)
            check.setChecked(index in self.selected_day_indexes)
            check.blockSignals(False)

    def select_day(self, day_index: int, save: bool = True):
        self.ensure_day_count()
        if not 1 <= day_index <= len(self.plan.get("days", [])):
            return
        self.current_day_index = day_index
        self.refresh_editor()
        self.refresh_calendar_cards()
        if save:
            self.refresh_overview(save=False)

    def open_day_detail(self, day_index: int | None = None):
        self.ensure_day_count()
        if day_index is None:
            day_index = self.current_day_index
        if not 1 <= day_index <= len(self.plan.get("days", [])):
            return
        self.select_day(day_index, save=False)
        dialog = PlannerDayDetailDialog(self, day_index)
        dialog.exec()
        self.refresh_editor()
        self.refresh_calendar_cards()
        self.refresh_overview(save=True)

    def refresh_editor(self):
        self._updating_editor = True
        day = self.day(self.current_day_index)
        name = str(day.get("name") or f"Den {self.current_day_index}")
        segment = _planner_day_segment(day, self.current_segment_key)
        note = str(segment.get("note") or "")
        encrypted_text = str(segment.get("encrypted_text") or "")
        ciphers = [str(cipher) for cipher in segment.get("ciphers", []) if str(cipher).strip()]
        attachments = _planner_normalize_attachments(segment.get("attachments", []))

        self.editor_title.setText(_planner_day_heading(day, self.current_day_index))
        self.day_name_edit.setText(name)
        if hasattr(self, "segment_combo"):
            segment_index = self.segment_combo.findData(self.current_segment_key)
            if segment_index < 0:
                segment_index = self.segment_combo.findData(_PLANNER_DEFAULT_SEGMENT_KEY)
                self.current_segment_key = _PLANNER_DEFAULT_SEGMENT_KEY
            self.segment_combo.blockSignals(True)
            self.segment_combo.setCurrentIndex(max(0, segment_index))
            self.segment_combo.blockSignals(False)
        self.day_note_edit.setPlainText(note)
        self.day_encrypted_edit.setPlainText(encrypted_text)
        self.assigned_combo.clear()
        self.assigned_combo.addItems(ciphers)
        if ciphers:
            self.day_ciphers_edit.setPlainText("\n".join(f"{i}. {cipher}" for i, cipher in enumerate(ciphers, start=1)))
        else:
            self.day_ciphers_edit.setPlainText("Zatím žádná šifra.")
        self.attachment_list.clear()
        for attachment_index, attachment in enumerate(attachments):
            item = QListWidgetItem(f"{attachment['title']}  [{attachment['kind'].upper()}]")
            item.setToolTip(attachment["path"])
            item.setData(Qt.UserRole, attachment_index)
            self.attachment_list.addItem(item)
        self._updating_editor = False

    def change_current_segment(self):
        if self._updating_editor:
            return
        segment_key = self.segment_combo.currentData()
        self.current_segment_key = str(segment_key or _PLANNER_DEFAULT_SEGMENT_KEY)
        self.refresh_editor()
        self.refresh_calendar_cards()
        self.refresh_overview(save=False)

    def toggle_day_selection(self, day_index: int, checked: bool):
        if self._rebuilding:
            return
        if checked:
            self.selected_day_indexes.add(day_index)
        else:
            self.selected_day_indexes.discard(day_index)
        self.select_day(day_index, save=False)
        self.refresh_overview(save=False)

    def select_all_days(self):
        self.ensure_day_count()
        self.selected_day_indexes = set(range(1, len(self.plan.get("days", [])) + 1))
        self.refresh_calendar_cards()
        self.refresh_overview(save=False)

    def clear_day_selection(self):
        self.selected_day_indexes.clear()
        self.refresh_calendar_cards()
        self.refresh_overview(save=False)

    def print_day_indexes(self) -> list[int]:
        indexes = sorted(self.selected_day_indexes)
        if indexes:
            return indexes
        return [self.current_day_index]

    def change_day_count(self):
        self.change_date_range()

    def change_date_range(self):
        if self._rebuilding:
            return
        if self._updating_dates:
            return
        self.ensure_day_count()
        self.rebuild_calendar_cards()
        self.select_day(self.current_day_index, save=False)
        self.refresh_overview(save=True)

    def rename_current_day(self):
        if self._updating_editor:
            return
        text = self.day_name_edit.text().strip()
        self.day(self.current_day_index)["name"] = text or f"Den {self.current_day_index}"
        self.refresh_calendar_cards()
        self.refresh_overview(save=True)

    def update_current_note(self):
        if self._updating_editor:
            return
        _planner_day_segment(self.day(self.current_day_index), self.current_segment_key)["note"] = self.day_note_edit.toPlainText().strip()
        self.refresh_calendar_cards()
        self.refresh_overview(save=True)

    def update_current_encrypted_text(self):
        if self._updating_editor:
            return
        _planner_day_segment(self.day(self.current_day_index), self.current_segment_key)["encrypted_text"] = self.day_encrypted_edit.toPlainText().strip()
        self.refresh_calendar_cards()
        self.refresh_overview(save=True)

    def cipher_location(
        self,
        cipher_name: str,
        ignore_day_index: int | None = None,
        ignore_segment_key: str | None = None,
    ) -> str:
        for index, day in enumerate(self.plan.get("days", []), start=1):
            for segment_key, segment_label, segment in _planner_day_segments(day):
                if (
                    ignore_day_index is not None
                    and index == ignore_day_index
                    and ignore_segment_key == segment_key
                ):
                    continue
                if cipher_name in segment.get("ciphers", []):
                    return f"{_planner_day_heading(day, index)} - {segment_label}"
        return ""

    def add_cipher_name_to_current_day(self, cipher_name: str, show_messages: bool = True) -> bool:
        if not self.available_names:
            if show_messages:
                QMessageBox.information(self, "Plánovač", "Nejsou dostupné žádné šifry.")
            return False

        cipher_name = str(cipher_name or "").strip()
        if not cipher_name:
            return False

        day = self.day(self.current_day_index)
        segment = _planner_day_segment(day, self.current_segment_key)
        if cipher_name in segment.get("ciphers", []):
            if show_messages:
                QMessageBox.information(self, "Plánovač", "Tahle šifra už je v této části dne.")
            return False

        location = self.cipher_location(
            cipher_name,
            ignore_day_index=self.current_day_index,
            ignore_segment_key=self.current_segment_key,
        )
        if location and self.prevent_repeat_check.isChecked():
            if show_messages:
                QMessageBox.warning(
                    self,
                    "Šifra už je použitá",
                    f"Šifra „{cipher_name}“ už je v plánu: {location}.\n\n"
                    "Když ji chceš opravdu použít znovu, vypni volbu „Hlídat opakování šifer“.",
                )
            return False

        segment.setdefault("ciphers", []).append(cipher_name)
        self.refresh_editor()
        self.refresh_calendar_cards()
        self.refresh_overview(save=True)
        return True

    def add_cipher_to_current_day(self):
        self.add_cipher_name_to_current_day(self.add_cipher_combo.currentText().strip(), show_messages=True)

    def encrypt_text_into_current_segment(self):
        plain_text = self.planner_plain_edit.toPlainText().strip()
        cipher_name = self.planner_cipher_combo.currentText().strip()
        if not plain_text:
            QMessageBox.information(self, "Šifrovat v kalendáři", "Nejdřív napiš text, který chceš zašifrovat.")
            return
        if not cipher_name:
            QMessageBox.information(self, "Šifrovat v kalendáři", "Nejdřív vyber šifru.")
            return

        try:
            encrypted_text = _planner_encrypt_text(cipher_name, plain_text, self.owner_window)
        except Exception as error:
            QMessageBox.warning(self, "Šifrovat v kalendáři", f"Šifrování selhalo:\n{error}")
            return

        segment = _planner_day_segment(self.day(self.current_day_index), self.current_segment_key)
        current_output = str(segment.get("encrypted_text") or "").strip()
        block = f"{cipher_name}:\n{encrypted_text}"
        segment["encrypted_text"] = current_output + "\n\n" + block if current_output else block
        added = True
        if cipher_name not in segment.get("ciphers", []):
            added = self.add_cipher_name_to_current_day(cipher_name, show_messages=True)
        if not added or cipher_name in segment.get("ciphers", []):
            self.refresh_editor()
            self.refresh_calendar_cards()
            self.refresh_overview(save=True)

    def add_attachment_to_current_segment(self):
        path, _ = QFileDialog.getOpenFileName(
            self,
            "Přidat PDF / Word přílohu",
            "",
            "Dokumenty (*.pdf *.docx *.doc *.txt *.md *.rtf);;Všechny soubory (*.*)",
        )
        if not path:
            return

        attachment = _planner_import_attachment(path)
        if attachment is None:
            return

        segment = _planner_day_segment(self.day(self.current_day_index), self.current_segment_key)
        attachments = _planner_normalize_attachments(segment.get("attachments", []))
        new_key = _planner_attachment_identity(attachment)
        if any(_planner_attachment_identity(item) == new_key for item in attachments):
            QMessageBox.information(self, "Příloha", "Tahle příloha už je v této části dne.")
            return
        attachments.append(attachment)
        segment["attachments"] = attachments
        self.refresh_editor()
        self.refresh_calendar_cards()
        self.refresh_overview(save=True)

    def selected_attachment_index(self) -> int:
        item = self.attachment_list.currentItem()
        if item is None:
            return -1
        try:
            return int(item.data(Qt.UserRole))
        except Exception:
            return -1

    def selected_attachment(self) -> dict | None:
        segment = _planner_day_segment(self.day(self.current_day_index), self.current_segment_key)
        attachments = _planner_normalize_attachments(segment.get("attachments", []))
        index = self.selected_attachment_index()
        if 0 <= index < len(attachments):
            return attachments[index]
        return None

    def open_selected_attachment(self):
        attachment = self.selected_attachment()
        if attachment is None:
            QMessageBox.information(self, "Příloha", "Nejdřív vyber přílohu ze seznamu.")
            return
        dialog = PlannerAttachmentDialog(attachment, self)
        dialog.exec()

    def remove_selected_attachment(self):
        segment = _planner_day_segment(self.day(self.current_day_index), self.current_segment_key)
        attachments = _planner_normalize_attachments(segment.get("attachments", []))
        index = self.selected_attachment_index()
        if not 0 <= index < len(attachments):
            return
        del attachments[index]
        segment["attachments"] = attachments
        self.refresh_editor()
        self.refresh_calendar_cards()
        self.refresh_overview(save=True)

    def current_app_cipher_and_output(self) -> tuple[str, str]:
        central = getattr(self.owner_window, "central", None)
        if central is None:
            return "", ""

        cipher_name = str(getattr(central, "selected_cipher", "") or "").strip()
        output_text = _history_current_output(central)
        return cipher_name, str(output_text or "").strip()

    def insert_current_app_result(self):
        cipher_name, output_text = self.current_app_cipher_and_output()
        if not cipher_name and not output_text:
            QMessageBox.information(
                self,
                "Vložit z aplikace",
                "V hlavním okně zatím není vybraná šifra ani zašifrovaný text.",
            )
            return

        if cipher_name:
            self.add_cipher_name_to_current_day(cipher_name, show_messages=False)

        if not output_text:
            QMessageBox.information(
                self,
                "Vložit z aplikace",
                "Šifra byla vložena do dne, ale v hlavním okně zatím není zašifrovaný text.",
            )
            return

        insert_text = f"{cipher_name}:\n{output_text}" if cipher_name else output_text
        current = self.day_encrypted_edit.toPlainText().strip()
        if current and insert_text not in current:
            new_text = current + "\n\n" + insert_text
        else:
            new_text = insert_text

        self.day_encrypted_edit.setPlainText(new_text)
        self.day_encrypted_edit.moveCursor(QTextCursor.End)

    def remove_cipher_from_current_day(self):
        cipher_name = self.assigned_combo.currentText().strip()
        if not cipher_name:
            return
        segment = _planner_day_segment(self.day(self.current_day_index), self.current_segment_key)
        segment["ciphers"] = [cipher for cipher in segment.get("ciphers", []) if cipher != cipher_name]
        self.refresh_editor()
        self.refresh_calendar_cards()
        self.refresh_overview(save=True)

    def clear_current_day(self):
        segment = _planner_day_segment(self.day(self.current_day_index), self.current_segment_key)
        if (
            not segment.get("ciphers")
            and not segment.get("attachments")
            and not str(segment.get("note") or "").strip()
            and not str(segment.get("encrypted_text") or "").strip()
        ):
            return
        segment["ciphers"] = []
        segment["attachments"] = []
        segment["note"] = ""
        segment["encrypted_text"] = ""
        self.refresh_editor()
        self.refresh_calendar_cards()
        self.refresh_overview(save=True)

    def refresh_overview(self, save: bool = True):
        self.ensure_day_count()
        if save:
            try:
                _planner_save(self.plan)
                if self.owner_window is not None and hasattr(self.owner_window, "central"):
                    self.owner_window.central.update_status()
            except Exception as error:
                _sifrator_debug_log(f"Plánovač: uložení selhalo: {type(error).__name__}: {error}")

        known = self.available_names
        used = _planner_used_cipher_names(self.plan)
        used_set = set(used)
        unused = [name for name in known if name not in used_set]
        selected = self.print_day_indexes()
        selected_text = ", ".join(str(index) for index in selected)
        selected_prefix = "Vybrané dny" if self.selected_day_indexes else "Bez zaškrtnutí se použije aktuální den"
        self.summary_label.setText(f"Použito {len(used)} / {len(known)} šifer, zbývá {len(unused)}")
        self.selection_label.setText(f"{selected_prefix}: {selected_text}")

        counts = self.repeated_cipher_counts()
        repeated = [name for name, count in counts.items() if count > 1]

        lines = []
        if repeated:
            lines.append("OPAKOVANÉ ŠIFRY")
            lines.append("-" * 36)
            lines.extend(repeated)
            lines.append("")

        lines.append("POUŽITÉ ŠIFRY")
        lines.append("-" * 36)
        if used:
            for cipher in used:
                location = self.cipher_location(cipher)
                lines.append(f"{cipher}  ({location})" if location else cipher)
        else:
            lines.append("Zatím žádná.")

        lines.append("")
        lines.append("JEŠTĚ NEPOUŽITÉ ŠIFRY")
        lines.append("-" * 36)
        if unused:
            lines.extend(unused)
        else:
            lines.append("Všechny dostupné šifry už jsou v plánu.")

        self.overview_edit.setPlainText("\n".join(lines))

    def copy_overview(self):
        text = _planner_export_text(self.plan, self.print_day_indexes())
        clipboard = QApplication.clipboard()
        if clipboard is not None:
            clipboard.setText(text)

    def export_plan(self):
        indexes = self.print_day_indexes()
        default_name = f"sifrator_plan_tabor_{time.strftime('%Y%m%d_%H%M%S')}.txt"
        path, _ = QFileDialog.getSaveFileName(
            self,
            "Exportovat vybrané dny jako TXT",
            default_name,
            "Textový soubor (*.txt);;Všechny soubory (*.*)",
        )
        if not path:
            return
        if not path.lower().endswith(".txt"):
            path += ".txt"

        try:
            with open(path, "w", encoding="utf-8") as file:
                file.write(_planner_export_text(self.plan, indexes))
            QMessageBox.information(self, "Export plánu", f"Plán byl uložen do:\n{path}")
        except Exception as error:
            QMessageBox.warning(self, "Export plánu", f"Plán se nepodařilo uložit:\n{error}")

    def print_selected_days(self):
        _planner_print_days(self, self.plan, self.print_day_indexes(), "Tisk vybraných dnů")

    def export_selected_pdf(self):
        indexes = self.print_day_indexes()
        default_name = f"sifrator_plan_tabor_{time.strftime('%Y%m%d_%H%M%S')}.pdf"
        path, _ = QFileDialog.getSaveFileName(
            self,
            "Exportovat vybrané dny jako PDF",
            default_name,
            "PDF soubor (*.pdf);;Všechny soubory (*.*)",
        )
        if not path:
            return
        if not path.lower().endswith(".pdf"):
            path += ".pdf"

        try:
            _planner_write_pdf(path, self.plan, indexes, "A4", "Na výšku")
            QMessageBox.information(self, "Export PDF", f"PDF bylo uloženo do:\n{path}")
        except Exception as error:
            QMessageBox.warning(self, "Export PDF", f"PDF se nepodařilo uložit:\n{error}")

    def clear_plan(self):
        answer = QMessageBox.question(
            self,
            "Vyčistit plán",
            "Opravdu chceš smazat celý plán šifer?",
            QMessageBox.Yes | QMessageBox.No,
            QMessageBox.No,
        )
        if answer != QMessageBox.Yes:
            return

        self.ensure_day_count()
        start_date = _planner_parse_date(self.plan.get("start_date"), date.today()) or date.today()
        count = _planner_clamp_day_count(self.plan.get("day_count") or 7)
        self.selected_day_indexes.clear()
        self.current_day_index = 1
        self.current_segment_key = _PLANNER_DEFAULT_SEGMENT_KEY
        self.plan = _planner_default_plan(count, start_date)
        self.rebuild_calendar_cards()
        self.select_day(1, save=False)
        self.refresh_overview(save=True)


_CIPHER_DEFAULT_METADATA = {
    _RANDOM_EASY_CIPHER_NAME: {"difficulty": 1, "age": "8+"},
    "Binární čtverce": {"difficulty": 2, "age": "10+"},
    "Brailovo písmo": {"difficulty": 2, "age": "9+"},
    "Britská vlajka": {"difficulty": 2, "age": "10+"},
    "Caesarova šifra": {"difficulty": 1, "age": "8+"},
    "Čtverec": {"difficulty": 1, "age": "8+"},
    "Hebrejský kříž": {"difficulty": 2, "age": "10+"},
    "Malý polský kříž": {"difficulty": 1, "age": "8+"},
    "Mobil": {"difficulty": 1, "age": "8+"},
    "Moonovo písmo": {"difficulty": 2, "age": "9+"},
    "Morseova abeceda": {"difficulty": 1, "age": "8+"},
    "Morseova abeceda – hory": {"difficulty": 2, "age": "9+"},
    "Morseova abeceda – pila": {"difficulty": 2, "age": "9+"},
    "Morseova abeceda – stromy": {"difficulty": 2, "age": "9+"},
    "Mříž": {"difficulty": 2, "age": "10+"},
    "Okno": {"difficulty": 2, "age": "10+"},
    "Pavoučí síť": {"difficulty": 3, "age": "12+"},
    "Posunková abeceda": {"difficulty": 2, "age": "9+"},
    "Pseudo-Čína": {"difficulty": 2, "age": "10+"},
    "Semafor": {"difficulty": 2, "age": "10+"},
    "SuperKrychle": {"difficulty": 3, "age": "12+"},
    "Tančící figurky": {"difficulty": 2, "age": "9+"},
    "Tančící figurky II": {"difficulty": 3, "age": "11+"},
    "Velký polský kříž": {"difficulty": 2, "age": "9+"},
    "Velký polský kříž (26 znaků)": {"difficulty": 3, "age": "11+"},
    "Vlčácká šifra": {"difficulty": 2, "age": "10+"},
    "Záměna písmen (A=Z)": {"difficulty": 1, "age": "8+"},
    "Záměna písmen za čísla (A=01, Z=26)": {"difficulty": 1, "age": "8+"},
    "Záměna písmen za čísla (A=26, Z=01)": {"difficulty": 1, "age": "8+"},
    "Zednářská šifra": {"difficulty": 2, "age": "10+"},
    "Zlomky": {"difficulty": 3, "age": "11+"},
}


def _cipher_notes_file_path() -> str:
    return os.path.join(_history_storage_dir(), "poznamky_sifer.json")


def _cipher_notes_load() -> dict:
    path = _cipher_notes_file_path()
    if not os.path.exists(path):
        return {}
    try:
        with open(path, "r", encoding="utf-8", errors="replace") as file:
            data = json.load(file)
    except Exception:
        return {}
    return data if isinstance(data, dict) else {}


def _cipher_notes_save(notes: dict) -> None:
    path = _cipher_notes_file_path()
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8") as file:
        json.dump(notes if isinstance(notes, dict) else {}, file, ensure_ascii=False, indent=2)


def _cipher_metadata(cipher_name: str) -> dict:
    data = dict(_CIPHER_DEFAULT_METADATA.get(str(cipher_name or ""), {}))
    data.setdefault("difficulty", 2)
    data.setdefault("age", "9+")
    return data


def _cipher_used_details() -> dict[str, list[str]]:
    used = {}
    known = set(_cipher_known_names())

    try:
        for cipher_name in _planner_used_cipher_names(_planner_load()):
            if known and cipher_name not in known:
                continue
            used.setdefault(cipher_name, [])
            if "plán" not in used[cipher_name]:
                used[cipher_name].append("plán")
    except Exception:
        pass

    try:
        for entry in _history_load():
            cipher_name = str(entry.get("cipher", "") or "").strip()
            if known and cipher_name not in known:
                continue
            if cipher_name:
                used.setdefault(cipher_name, [])
                if "historie" not in used[cipher_name]:
                    used[cipher_name].append("historie")
    except Exception:
        pass

    return used


def _cipher_known_names() -> list[str]:
    try:
        names = list(list_cipher_names())
    except Exception:
        names = list(_CIPHER_DEFAULT_METADATA)
    if _RANDOM_EASY_CIPHER_NAME not in names:
        names.append(_RANDOM_EASY_CIPHER_NAME)
    return names


def _cipher_overview_refresh_usage_marks(widget) -> None:
    used = _cipher_used_details()
    notes = _cipher_notes_load()
    known = _cipher_known_names()

    for btn in getattr(widget, "cipher_buttons", []):
        cipher_name = str(getattr(getattr(btn, "item", None), "name", "") or "")
        if not cipher_name:
            continue
        meta = _cipher_metadata(cipher_name)
        is_used = cipher_name in used
        mark = "✓" if is_used else "○"
        btn.full_text = f"{mark} {cipher_name}"
        try:
            btn.update_elided_text()
        except Exception:
            pass
        note = str(notes.get(cipher_name, {}).get("note", "") if isinstance(notes.get(cipher_name), dict) else "")
        status = "použitá" if is_used else "nepoužitá"
        source = ", ".join(used.get(cipher_name, [])) or "zatím nikde"
        tooltip = (
            f"{cipher_name}\n"
            f"Stav: {status} ({source})\n"
            f"Obtížnost: {meta['difficulty']}/3\n"
            f"Doporučený věk: {meta['age']}"
        )
        if note.strip():
            tooltip += f"\nPoznámka: {note.strip()}"
        btn.setToolTip(tooltip)

    title = getattr(widget, "title_left", None)
    if title is not None:
        title.setText(f"VYBER SI ŠIFRU ({len(used)}/{len(known)} použito)")


def _cipher_plain_ascii(text: str) -> str:
    normalized = unicodedata.normalize("NFKD", str(text or ""))
    normalized = "".join(ch for ch in normalized if not unicodedata.combining(ch))
    return normalized.upper()


def _cipher_shift_text(text: str, shift: int) -> str:
    alphabet = "ABCDEFGHIJKLMNOPQRSTUVWXYZ"
    result = []
    for ch in _cipher_plain_ascii(text):
        if ch in alphabet:
            result.append(alphabet[(alphabet.index(ch) + shift) % len(alphabet)])
        else:
            result.append(ch)
    return "".join(result)


def _cipher_vowel_number_text(text: str) -> str:
    mapping = {"A": "1", "E": "2", "I": "3", "O": "4", "U": "5", "Y": "6"}
    return "".join(mapping.get(ch, ch) for ch in _cipher_plain_ascii(text))


def _cipher_reverse_words_text(text: str) -> str:
    return " ".join(word[::-1] for word in _cipher_plain_ascii(text).split(" "))


def _cipher_even_odd_words_text(text: str) -> str:
    words = []
    for word in _cipher_plain_ascii(text).split(" "):
        letters = [ch for ch in word if ch.isalpha()]
        if len(letters) < 3:
            words.append(word)
            continue
        odd = letters[0::2]
        even = letters[1::2]
        words.append("".join(odd + even))
    return " ".join(words)


def _cipher_generate_easy_custom_cipher(text: str) -> dict:
    result = _cipher_apply_easy_recipe(_cipher_make_easy_recipe(), text)
    result["difficulty"] = 1
    result["age"] = "8+"
    result["plain_text"] = _cipher_plain_ascii(text)
    return result


def _cipher_easy_recipes() -> list[dict]:
    recipes = [{"kind": "shift", "shift": shift} for shift in range(1, 6)]
    recipes.extend([
        {"kind": "vowels"},
        {"kind": "reverse_words"},
        {"kind": "even_odd"},
    ])
    return recipes


def _cipher_make_easy_recipe(previous: dict | None = None) -> dict:
    recipes = _cipher_easy_recipes()
    if isinstance(previous, dict) and previous:
        recipes = [recipe for recipe in recipes if recipe != previous] or recipes
    return dict(random.choice(recipes))


def _cipher_apply_easy_recipe(recipe: dict, text: str) -> dict:
    kind = str((recipe or {}).get("kind") or "shift")
    if kind == "shift":
        shift = int((recipe or {}).get("shift") or 3)
        return {
            "name": f"Posun o {shift}",
            "cipher_text": _cipher_shift_text(text, shift),
            "key": f"Každé písmeno je posunuté v abecedě o {shift} dál. A→{_cipher_shift_text('A', shift)}, B→{_cipher_shift_text('B', shift)}.",
            "hint": "Hledej, jestli se písmena neposunula v abecedě vždy stejně.",
        }
    if kind == "vowels":
        return {
            "name": "Samohlásky jako čísla",
            "cipher_text": _cipher_vowel_number_text(text),
            "key": "A=1, E=2, I=3, O=4, U=5, Y=6. Souhlásky zůstávají stejné.",
            "hint": "Čísla se objevují hlavně tam, kde ve slovech bývají samohlásky.",
        }
    if kind == "reverse_words":
        return {
            "name": "Slova pozpátku",
            "cipher_text": _cipher_reverse_words_text(text),
            "key": "Každé slovo se čte odzadu, ale pořadí slov zůstává stejné.",
            "hint": "Zkus přečíst každé slovo zprava doleva.",
        }
    return {
        "name": "Lichá písmena před sudými",
        "cipher_text": _cipher_even_odd_words_text(text),
        "key": "V každém slově jsou nejdřív písmena z pozic 1, 3, 5... a za nimi písmena z pozic 2, 4, 6...",
        "hint": "Všimni si, že písmena jsou správná, jen v každém slově přeházená podle pozic.",
    }


def _random_easy_format_output(generated: dict) -> str:
    return (
        "NÁHODNÁ LEHKÁ ŠIFRA\n"
        "=" * 40 + "\n\n"
        f"Typ: {generated.get('name', '')}\n"
        f"Obtížnost: {generated.get('difficulty', 1)}/3\n"
        f"Doporučený věk: {generated.get('age', '8+')}\n\n"
        "ZAŠIFROVANÝ TEXT PRO DĚTI:\n"
        f"{generated.get('cipher_text', '')}\n\n"
        "KLÍČ PRO VEDOUCÍHO:\n"
        f"{generated.get('key', '')}\n\n"
        "NÁPOVĚDA PRO DĚTI:\n"
        f"{generated.get('hint', '')}\n\n"
        "Původní text bez diakritiky:\n"
        f"{generated.get('plain_text', '')}"
    )


def _random_easy_format_key(generated: dict) -> str:
    return (
        "KLÍČ ŠIFRY - NÁHODNÁ LEHKÁ ŠIFRA\n"
        "=" * 40 + "\n\n"
        f"Typ šifry: {generated.get('name', '')}\n"
        f"Obtížnost: {generated.get('difficulty', 1)}/3\n"
        f"Doporučený věk: {generated.get('age', '8+')}\n\n"
        "Jak se luští:\n"
        f"{generated.get('key', '')}\n\n"
        "Nápověda pro děti:\n"
        f"{generated.get('hint', '')}\n\n"
        "Zašifrovaný text:\n"
        f"{generated.get('cipher_text', '')}\n\n"
        "Původní text bez diakritiky:\n"
        f"{generated.get('plain_text', '')}"
    )


def _random_easy_key_data(generated: dict) -> dict:
    return {
        "title": f"Klíč šifry – {generated.get('name', _RANDOM_EASY_CIPHER_NAME)}",
        "subtitle": "Šifrátor Mraveniště – náhodná lehká šifra",
        "description": "Dětem dej jen zašifrovaný text. Klíč a nápovědu použij postupně podle potřeby.",
        "cipher_name": _RANDOM_EASY_CIPHER_NAME,
        "logo_path": os.path.join(get_icons_dir(), _RANDOM_EASY_CIPHER_ICON),
        "theme": "pirate_modern",
        "background_path": os.path.join(
            get_icons_dir(), "key_templates", "pirate_key_morse_template.png"
        ),
        "type": "generic",
        "columns": 2,
        "items": [
            ("Typ šifry", str(generated.get("name", ""))),
            ("Obtížnost / věk", f"{generated.get('difficulty', 1)}/3, {generated.get('age', '8+')}"),
            ("Jak se luští", str(generated.get("key", ""))),
            ("Nápověda", str(generated.get("hint", ""))),
            ("Zašifrovaný text", str(generated.get("cipher_text", ""))),
            ("Původní text", str(generated.get("plain_text", ""))),
        ],
    }


def _random_easy_current_data(widget, text: str, force_new: bool = False) -> dict:
    normalized_text = _cipher_plain_ascii(text)
    data = getattr(widget, "_random_easy_cipher_data", None)
    recipe = getattr(widget, "_random_easy_cipher_recipe", None)
    if (
        not force_new
        and isinstance(data, dict)
        and data.get("plain_text") == normalized_text
        and data.get("cipher_text")
    ):
        return data

    if force_new or not isinstance(recipe, dict) or not recipe:
        recipe = _cipher_make_easy_recipe(recipe if isinstance(recipe, dict) else None)
        widget._random_easy_cipher_recipe = recipe

    data = _cipher_apply_easy_recipe(recipe, text)
    data["difficulty"] = 1
    data["age"] = "8+"
    data["plain_text"] = normalized_text
    widget._random_easy_cipher_data = data
    return data


_RANDOM_EASY_ORIGINAL_ENCRYPT_SELECTED_CIPHER = SifratorSkinWidget.encrypt_selected_cipher
_RANDOM_EASY_ORIGINAL_DECRYPT_SELECTED_CIPHER = SifratorSkinWidget.decrypt_selected_cipher
_RANDOM_EASY_ORIGINAL_SHOW_CIPHER_KEY = SifratorSkinWidget.show_cipher_key


def _random_easy_encrypt_selected_cipher(self, text: str) -> str:
    if self.selected_cipher == _RANDOM_EASY_CIPHER_NAME:
        if not str(text or "").strip():
            return "Nejdřív zadej text, ze kterého mám vymyslet lehkou šifru."
        force_new = bool(getattr(self, "_random_easy_force_next_encrypt", False))
        self._random_easy_force_next_encrypt = False
        generated = _random_easy_current_data(self, text, force_new=force_new)
        return str(generated.get("cipher_text", "") or "")
    return _RANDOM_EASY_ORIGINAL_ENCRYPT_SELECTED_CIPHER(self, text)


def _random_easy_decrypt_selected_cipher(self, text: str) -> str:
    if self.selected_cipher == _RANDOM_EASY_CIPHER_NAME:
        data = getattr(self, "_random_easy_cipher_data", None)
        if isinstance(data, dict) and data.get("plain_text"):
            return (
                "Tahle vymyšlená šifra se řeší podle klíče.\n\n"
                f"Klíč: {data.get('key', '')}\n\n"
                f"Původní text:\n{data.get('plain_text', '')}"
            )
        return "Nejdřív si nech text zašifrovat náhodnou lehkou šifrou."
    return _RANDOM_EASY_ORIGINAL_DECRYPT_SELECTED_CIPHER(self, text)


def _random_easy_show_cipher_key(self):
    if self.selected_cipher != _RANDOM_EASY_CIPHER_NAME:
        return _RANDOM_EASY_ORIGINAL_SHOW_CIPHER_KEY(self)

    text = self.get_input_text() if hasattr(self, "get_input_text") else ""
    if not text.strip():
        QMessageBox.information(self, "Klíč náhodné šifry", "Nejdřív zadej text a nech ho zašifrovat.")
        return

    generated = _random_easy_current_data(self, text)
    renderer = get_pirate_key_renderer()
    dialog_class = getattr(renderer, "PirateKeyDialog", None) if renderer is not None else None
    if dialog_class is None:
        QMessageBox.information(self, "Klíč náhodné šifry", _random_easy_format_key(generated))
        return

    dialog = dialog_class(_random_easy_key_data(generated), self)
    dialog.exec()


SifratorSkinWidget.encrypt_selected_cipher = _random_easy_encrypt_selected_cipher
SifratorSkinWidget.decrypt_selected_cipher = _random_easy_decrypt_selected_cipher
SifratorSkinWidget.show_cipher_key = _random_easy_show_cipher_key


class CipherOverviewDialog(PirateModuleDialog):
    """Přehled použití, obtížnosti, věku, poznámek a náhodných šifer."""

    def __init__(self, owner_window):
        super().__init__(
            owner_window,
            "overview_BG.png",
            ((0.167, 0.255, 0.70), (0.870, 0.252, 0.70), (0.919, 0.797, 1.18)),
        )
        self.owner_window = owner_window
        self.central = getattr(owner_window, "central", None)
        self.notes = _cipher_notes_load()
        self.used = _cipher_used_details()
        self._updating_note = False

        self.setWindowTitle("Přehled šifer")
        self.resize(1050, 720)
        self.setMinimumSize(840, 560)
        self.setStyleSheet("""
            QDialog { background-color: #0a1626; color: #f3ddaa; }
            QLabel { color: #ead8b3; background: transparent; }
            QListWidget, QTextEdit {
                color: #f0e2c0;
                background: #071018;
                border: 1px solid #8a6938;
                border-radius: 6px;
                padding: 6px;
            }
            QListWidget::item { padding: 6px; }
            QListWidget::item:selected { background: #144a63; color: #fff2cc; }
            QPushButton {
                color: #f6e7bf;
                background-color: #10263e;
                border: 1px solid #c49344;
                border-radius: 8px;
                padding: 8px 12px;
                font-weight: bold;
            }
            QPushButton:hover { background-color: #144a63; }
        """)

        root = QVBoxLayout(self)
        root.setContentsMargins(12, 12, 12, 12)
        root.setSpacing(10)

        title = QLabel("Přehled použitých šifer", self)
        title.setStyleSheet("color: #f3d79a; font-size: 20px; font-weight: bold;")
        root.addWidget(title)

        body = QHBoxLayout()
        body.setSpacing(12)
        root.addLayout(body, 1)

        left = QVBoxLayout()
        left.addWidget(QLabel("Šifry:", self))
        self.cipher_list = QListWidget(self)
        self.cipher_list.currentRowChanged.connect(self.refresh_selected_cipher)
        left.addWidget(self.cipher_list, 1)
        body.addLayout(left, 2)

        right = QVBoxLayout()
        self.detail_label = QLabel(self)
        self.detail_label.setWordWrap(True)
        self.detail_label.setStyleSheet("color: #f3d79a; font-size: 15px; font-weight: bold;")
        right.addWidget(self.detail_label)

        self.note_edit = QTextEdit(self)
        self.note_edit.setPlaceholderText("Vlastní poznámka vedoucího k této šifře.")
        self.note_edit.textChanged.connect(self.save_current_note)
        right.addWidget(QLabel("Vlastní poznámka:", self))
        right.addWidget(self.note_edit, 1)

        button_row = QHBoxLayout()
        self.select_button = QPushButton("Vybrat šifru", self)
        self.random_unused_button = QPushButton("Náhodná nepoužitá", self)
        self.random_easy_button = QPushButton("Vymyslet lehkou šifru z textu", self)
        button_row.addWidget(self.select_button)
        button_row.addWidget(self.random_unused_button)
        button_row.addWidget(self.random_easy_button)
        right.addLayout(button_row)

        self.generated_edit = QTextEdit(self)
        self.generated_edit.setReadOnly(True)
        self.generated_edit.setPlaceholderText("Tady se zobrazí vymyšlená jednoduchá šifra, její klíč a nápověda.")
        right.addWidget(QLabel("Náhodně vymyšlená šifra:", self))
        right.addWidget(self.generated_edit, 2)

        body.addLayout(right, 3)

        bottom = QHBoxLayout()
        bottom.addStretch(1)
        self.refresh_button = QPushButton("Obnovit", self)
        self.close_button = QPushButton("Zavřít", self)
        bottom.addWidget(self.refresh_button)
        bottom.addWidget(self.close_button)
        root.addLayout(bottom)

        self.select_button.clicked.connect(self.select_current_cipher)
        self.random_unused_button.clicked.connect(self.choose_random_unused_cipher)
        self.random_easy_button.clicked.connect(self.generate_easy_cipher)
        self.refresh_button.clicked.connect(self.refresh_overview)
        self.close_button.clicked.connect(self.close)

        self.apply_pirate_glass()
        self.refresh_overview()

    def cipher_names(self) -> list[str]:
        return _cipher_known_names()

    def current_cipher_name(self) -> str:
        item = self.cipher_list.currentItem()
        return str(item.data(Qt.UserRole) or "") if item is not None else ""

    def refresh_overview(self):
        self.used = _cipher_used_details()
        current = self.current_cipher_name()
        self.cipher_list.clear()

        names = self.cipher_names()
        for cipher_name in names:
            meta = _cipher_metadata(cipher_name)
            used = cipher_name in self.used
            mark = "✓ použité" if used else "○ nepoužité"
            item = QListWidgetItem(f"{mark} | obtížnost {meta['difficulty']}/3 | věk {meta['age']} | {cipher_name}")
            item.setData(Qt.UserRole, cipher_name)
            if used:
                item.setToolTip(f"Použité v: {', '.join(self.used.get(cipher_name, []))}")
            self.cipher_list.addItem(item)

        target_row = 0
        if current:
            for row in range(self.cipher_list.count()):
                if self.cipher_list.item(row).data(Qt.UserRole) == current:
                    target_row = row
                    break
        if self.cipher_list.count():
            self.cipher_list.setCurrentRow(target_row)

        if self.central is not None:
            _cipher_overview_refresh_usage_marks(self.central)

    def refresh_selected_cipher(self):
        cipher_name = self.current_cipher_name()
        meta = _cipher_metadata(cipher_name)
        used = cipher_name in self.used
        source = ", ".join(self.used.get(cipher_name, [])) if used else "zatím nepoužitá"
        self.detail_label.setText(
            f"{cipher_name}\n"
            f"Stav: {'použitá' if used else 'nepoužitá'} ({source})\n"
            f"Obtížnost: {meta['difficulty']}/3, doporučený věk: {meta['age']}"
        )

        note = ""
        raw = self.notes.get(cipher_name)
        if isinstance(raw, dict):
            note = str(raw.get("note") or "")
        self._updating_note = True
        self.note_edit.setPlainText(note)
        self._updating_note = False

    def save_current_note(self):
        if self._updating_note:
            return
        cipher_name = self.current_cipher_name()
        if not cipher_name:
            return
        self.notes.setdefault(cipher_name, {})
        if not isinstance(self.notes[cipher_name], dict):
            self.notes[cipher_name] = {}
        self.notes[cipher_name]["note"] = self.note_edit.toPlainText().strip()
        _cipher_notes_save(self.notes)
        if self.central is not None:
            _cipher_overview_refresh_usage_marks(self.central)

    def select_current_cipher(self):
        cipher_name = self.current_cipher_name()
        if self.central is not None and cipher_name:
            self.central.select_cipher(cipher_name)
            self.central.auto_encrypt_action()

    def choose_random_unused_cipher(self):
        names = self.cipher_names()
        unused = [name for name in names if name not in self.used]
        if not unused:
            unused = names
            QMessageBox.information(self, "Náhodná šifra", "Všechny šifry už jsou použité, vyberu tedy náhodně ze všech.")
        cipher_name = random.choice(unused)
        for row in range(self.cipher_list.count()):
            if self.cipher_list.item(row).data(Qt.UserRole) == cipher_name:
                self.cipher_list.setCurrentRow(row)
                break
        self.select_current_cipher()

    def generate_easy_cipher(self):
        if self.central is None:
            return
        input_widget = getattr(self.central, "input_text", None)
        text = input_widget.toPlainText().strip() if input_widget is not None else ""
        if not text:
            QMessageBox.information(self, "Vymyslet šifru", "Nejdřív napiš text do hlavního vstupního pole aplikace.")
            return

        generated = _random_easy_current_data(self.central, text, force_new=True)
        output = _random_easy_format_output(generated)
        self.generated_edit.setPlainText(output)
        try:
            self.central.select_cipher(_RANDOM_EASY_CIPHER_NAME)
            self.central.auto_encrypt_action()
        except Exception:
            pass


def _batch_image_dir() -> str:
    path = os.path.join(tempfile.gettempdir(), APP_NAME, "batch_images")
    os.makedirs(path, exist_ok=True)
    return path


def _batch_escape_html(value: str) -> str:
    import html

    return html.escape(str(value or "")).replace("\n", "<br>")


def _batch_context_label(context: dict) -> str:
    return _history_context_label(context)


def _batch_cipher_context(dialog) -> dict:
    cipher_name = dialog.current_cipher_name()
    try:
        cipher_name = dialog.current_cipher_name_for_context()
    except Exception:
        pass
    if cipher_name == "Caesarova šifra":
        direction = "dozadu" if "DOZADU" in dialog.caesar_direction_combo.currentText().upper() else "dopředu"
        return {
            "caesar_shift": int(dialog.caesar_shift_input.value()),
            "caesar_direction": direction,
        }
    return {}


def _batch_cipher_context_for_name(dialog, cipher_name: str) -> dict:
    previous = getattr(dialog, "_batch_context_cipher_name", None)
    dialog._batch_context_cipher_name = cipher_name
    try:
        return _batch_cipher_context(dialog)
    finally:
        if previous is None:
            try:
                delattr(dialog, "_batch_context_cipher_name")
            except Exception:
                pass
        else:
            dialog._batch_context_cipher_name = previous


def _batch_available_cipher_names(dialog) -> list[str]:
    try:
        return [dialog.cipher_combo.itemText(index) for index in range(dialog.cipher_combo.count())]
    except Exception:
        return list_cipher_names()


def _batch_resolve_cipher_name(raw_name: str, fallback: str, available_names: list[str]) -> str:
    raw = str(raw_name or "").strip()
    if not raw:
        return fallback

    exact = {name.casefold(): name for name in available_names}
    if raw.casefold() in exact:
        return exact[raw.casefold()]

    starts = [name for name in available_names if name.casefold().startswith(raw.casefold())]
    if len(starts) == 1:
        return starts[0]

    contains = [name for name in available_names if raw.casefold() in name.casefold()]
    if len(contains) == 1:
        return contains[0]

    return ""


def _batch_encrypt_message(cipher_name: str, message: str, context: dict) -> str:
    logic = get_cipher_logic(cipher_name)
    if logic is None or not hasattr(logic, "encrypt"):
        return f"Chybí logika šifry: {cipher_name}"

    if cipher_name == "Caesarova šifra":
        shift = int(context.get("caesar_shift", 3))
        if context.get("caesar_direction") == "dozadu":
            shift = -shift
        return logic.encrypt(message, shift)

    return logic.encrypt(message)


def _batch_render_cipher_image(cipher_name: str, output_text: str, input_text: str, index: int) -> str:
    widget_class = get_cipher_widget_class(cipher_name)
    if widget_class is None:
        return ""

    widget = None
    try:
        widget = widget_class()
        widget.setAttribute(Qt.WA_DontShowOnScreen, True)
        widget.setAttribute(Qt.WA_TranslucentBackground, True)
        if hasattr(widget, "set_scale"):
            widget.set_scale(1.0)
        if hasattr(widget, "set_cipher_text"):
            widget.set_cipher_text(output_text)
        elif hasattr(widget, "setText"):
            widget.setText(output_text)

        width = 900
        height = 260
        if hasattr(widget, "calculate_required_height"):
            try:
                height = max(height, int(widget.calculate_required_height(width)))
            except Exception:
                pass
        widget.resize(width, height)
        if hasattr(widget, "update_content_size"):
            widget.update_content_size()
        app = QApplication.instance()
        if app is not None:
            app.processEvents()

        pixmap = widget.grab()
        if pixmap.isNull():
            return ""

        payload = json.dumps(
            {
                "cipher": cipher_name,
                "index": index,
                "input": input_text,
                "output": output_text,
                "time": time.strftime("%Y-%m-%d %H:%M:%S"),
            },
            ensure_ascii=False,
            sort_keys=True,
            default=str,
        )
        digest = hashlib.sha1(payload.encode("utf-8", errors="replace")).hexdigest()[:16]
        file_name = f"davka_{time.strftime('%Y%m%d_%H%M%S')}_{index:02d}_{digest}.png"
        path = os.path.join(_batch_image_dir(), file_name)
        if pixmap.save(path, "PNG"):
            return path
    except Exception as error:
        _sifrator_debug_log(f"Hromadné šifrování: obrázkový náhled selhal: {type(error).__name__}: {error}")
    finally:
        try:
            if widget is not None:
                widget.close()
                widget.deleteLater()
        except Exception:
            pass

    return ""


def _batch_results_plain_text(results: list[dict]) -> str:
    if not results:
        return ""

    blocks = []
    for item in results:
        title = f"{item.get('index', '?')}. stanoviště | {item.get('cipher', '')}"
        context = _batch_context_label(item.get("context", {}))
        if context:
            title += f" ({context})"
        block = (
            f"{title}\n"
            f"{'-' * min(72, max(16, len(title)))}\n"
            f"Zpráva:\n{item.get('input', '')}\n\n"
            f"Zašifrováno:\n{item.get('output', '')}"
        )
        image_path = str(item.get("image_path", "") or "")
        if image_path:
            block += f"\n\nObrázek:\n{image_path}"
        blocks.append(block)

    return ("\n\n" + "=" * 72 + "\n\n").join(blocks)


def _batch_station_title_from_template(item: dict, settings: dict) -> str:
    template = str(_print_settings_value(settings, "station_title_template", "{index}. stanoviště | {cipher}"))
    if not template.strip():
        return ""

    context = _batch_context_label(item.get("context", {}))
    values = {
        "index": str(item.get("index", "?")),
        "n": str(item.get("index", "?")),
        "station": str(item.get("index", "?")),
        "stanoviste": str(item.get("index", "?")),
        "cipher": str(item.get("cipher", "")),
        "sifra": str(item.get("cipher", "")),
        "context": context,
        "kontext": context,
    }

    title = template
    for key, value in values.items():
        title = title.replace("{" + key + "}", value)

    return title.strip()


def _batch_results_html(results: list[dict], print_mode: bool = False, settings: dict | None = None) -> str:
    if not results:
        empty = "Výsledky zatím nejsou připravené."
        return f"<html><body style='color:#f0e2c0; background:transparent;'>{empty}</body></html>"

    settings = settings or {}
    heading_size = int(_print_settings_value(settings, "heading_size", 14 if print_mode else 14))
    text_size = int(_print_settings_value(settings, "text_size", 12 if print_mode else 12))
    cipher_scale = int(_print_settings_value(settings, "cipher_scale", 90 if print_mode else 100))
    show_frames = bool(_print_settings_value(settings, "show_frames", True))
    show_station_title = bool(_print_settings_value(settings, "show_station_title", True))
    show_input = bool(_print_settings_value(settings, "show_input", True))
    show_output = bool(_print_settings_value(settings, "show_output", True))
    show_images = bool(_print_settings_value(settings, "show_images", True))
    input_label = str(_print_settings_value(settings, "input_label", "Zpráva:"))
    output_label = str(_print_settings_value(settings, "output_label", "Zašifrováno:"))

    body_color = "#111111" if print_mode else "#f0e2c0"
    background = "#ffffff" if print_mode else "transparent"
    title_color = "#10223a" if print_mode else "#f3d79a"
    line_color = "#777777" if print_mode else "#8a6938"
    card_style = (
        "page-break-inside: avoid; border:1px solid #999; border-radius:6px; "
        "padding:12px; margin:0 0 16px 0;"
        if print_mode and show_frames else
        "page-break-inside: avoid; border:none; padding:0; margin:0 0 16px 0;"
        if print_mode else
        "border:1px solid #8a6938; border-radius:6px; padding:10px; margin:0 0 14px 0;"
    )

    blocks = []
    for item in results:
        title = _batch_station_title_from_template(item, settings)

        image_html = ""
        image_path = str(item.get("image_path", "") or "")
        if show_images and image_path and os.path.exists(image_path):
            src = _batch_escape_html(QUrl.fromLocalFile(image_path).toString())
            image = QImage(image_path)
            width = 720
            if not image.isNull():
                width = min(width, max(1, image.width()))
            if print_mode:
                width = int(max(60, width * max(0, min(100, cipher_scale)) / 100.0))
            image_html = (
                f"<div style='margin-top:10px;'><img src='{src}' width='{width}' "
                "style='max-width:100%; height:auto;'></div>"
            )

        parts = [f"<div style='{card_style}'>"]
        if show_station_title and title:
            parts.append(
                f"<div style='font-size:{heading_size}pt; font-weight:bold; color:{title_color};'>"
                f"{_batch_escape_html(title)}</div>"
                f"<div style='height:1px; background:{line_color}; margin:6px 0 10px 0;'></div>"
            )
        if show_input:
            if input_label.strip():
                parts.append(f"<div style='font-weight:bold;'>{_batch_escape_html(input_label)}</div>")
            parts.append(
                f"<div style='white-space:pre-wrap;'>{_batch_escape_html(item.get('input', ''))}</div>"
            )
        if show_input and show_output:
            parts.append("<div style='height:8px;'></div>")
        if show_output:
            if output_label.strip():
                parts.append(f"<div style='font-weight:bold;'>{_batch_escape_html(output_label)}</div>")
            parts.append(
                f"<div style='white-space:pre-wrap;'>{_batch_escape_html(item.get('output', ''))}</div>"
            )
        parts.append(image_html)
        if not show_station_title and not show_input and not show_output and not image_html:
            parts.append("<div style='color:#777;'>Toto stanoviště nemá vybraný žádný obsah k tisku.</div>")
        parts.append("</div>")
        blocks.append("".join(parts))

    return (
        "<html><head><meta charset='utf-8'></head>"
        f"<body style='font-family:Georgia, serif; font-size:{text_size}pt; color:{body_color}; background:{background};'>"
        + "".join(blocks)
        + "</body></html>"
    )


def _batch_default_print_settings() -> dict:
    return {
        "show_frames": True,
        "show_station_title": True,
        "show_input": True,
        "show_output": True,
        "show_images": True,
        "station_title_template": "{index}. stanoviště | {cipher}",
        "input_label": "Zpráva:",
        "output_label": "Zašifrováno:",
        "heading_size": 14,
        "text_size": 12,
        "cipher_scale": 90,
    }


def _batch_build_print_document(results: list[dict], paper_name: str, orientation_name: str, settings: dict | None = None):
    from PySide6.QtGui import QTextDocument

    document = QTextDocument()
    document.setDefaultFont(QFont("Georgia", 11))
    document.setPageSize(_print_page_size_points(paper_name, orientation_name))
    document.setHtml(_batch_results_html(results, print_mode=True, settings=settings or {}))
    return document


def _batch_write_pdf(path: str, results: list[dict], paper_name: str, orientation_name: str, settings: dict | None = None) -> None:
    writer = _pdf_writer_for_path(path, paper_name, orientation_name)
    document = _batch_build_print_document(results, paper_name, orientation_name, settings or {})
    document.print_(writer)


def _batch_export_pdf(parent, results: list[dict], paper_name: str = "A4", orientation_name: str = "Na výšku", settings: dict | None = None) -> bool:
    if not results:
        QMessageBox.information(parent, "Export PDF", "Výsledky zatím nejsou připravené.")
        return False

    default_name = f"sifrator_hromadne_{time.strftime('%Y%m%d_%H%M%S')}.pdf"
    path, _ = QFileDialog.getSaveFileName(
        parent,
        "Exportovat hromadné šifrování jako PDF",
        default_name,
        "PDF soubor (*.pdf);;Všechny soubory (*.*)",
    )
    if not path:
        return False
    if not path.lower().endswith(".pdf"):
        path += ".pdf"

    try:
        _batch_write_pdf(path, results, paper_name, orientation_name, settings or {})
        QMessageBox.information(parent, "Export PDF", f"PDF bylo uloženo do:\n{path}")
        return True
    except Exception as error:
        QMessageBox.warning(parent, "Export PDF", f"PDF se nepodařilo uložit:\n{error}")
        return False


def _batch_save_result_to_history(item: dict) -> None:
    try:
        entry = {
            "created_at": time.strftime("%Y-%m-%d %H:%M:%S"),
            "cipher": str(item.get("cipher", "")),
            "context": item.get("context", {}),
            "input": _history_limit_text(str(item.get("input", ""))),
            "output": _history_limit_text(str(item.get("output", ""))),
        }

        image_path = str(item.get("image_path", "") or "")
        if image_path and os.path.exists(image_path):
            image = QImage(image_path)
            if not image.isNull():
                digest = hashlib.sha1(image_path.encode("utf-8", errors="replace")).hexdigest()[:16]
                image_name = f"historie_davka_{time.strftime('%Y%m%d_%H%M%S')}_{digest}.png"
                target = os.path.join(_history_image_dir(), image_name)
                if image.save(target, "PNG"):
                    entry["image"] = image_name

        _history_add_entry(entry)
    except Exception as error:
        _sifrator_debug_log(f"Hromadné šifrování: zápis do historie selhal: {type(error).__name__}: {error}")


class BatchEncryptDialog(PirateModuleDialog):
    """Hromadné šifrování více zpráv najednou."""

    def __init__(self, owner_window):
        super().__init__(
            owner_window,
            "batch_BG.png",
            ((0.063, 0.227, 0.78), (0.946, 0.834, 1.18)),
        )
        self.owner_window = owner_window
        self.results = []
        self.station_cipher_combos = []
        self.setWindowTitle("Hromadné šifrování")
        self.resize(980, 680)

        layout = QVBoxLayout(self)
        layout.setContentsMargins(12, 12, 12, 12)
        layout.setSpacing(8)

        top_row = QHBoxLayout()
        self.mode_label = QLabel("Režim", self)
        self.mode_label.setStyleSheet("color: #ead8b3;")
        top_row.addWidget(self.mode_label)

        self.mode_combo = QComboBox(self)
        self.mode_combo.addItems(["Jedna šifra pro všechna stanoviště", "Jiná šifra pro každé stanoviště"])
        top_row.addWidget(self.mode_combo)

        self.cipher_label = QLabel("Šifra", self)
        self.cipher_label.setStyleSheet("color: #ead8b3;")
        top_row.addWidget(self.cipher_label)

        self.cipher_combo = QComboBox(self)
        cipher_names = []
        try:
            cipher_names = [item.name for item in self.owner_window.central.ciphers]
        except Exception:
            cipher_names = list_cipher_names()
        self.cipher_combo.addItems(cipher_names)
        current_cipher = getattr(getattr(self.owner_window, "central", None), "selected_cipher", None)
        if current_cipher:
            index = self.cipher_combo.findText(current_cipher)
            if index >= 0:
                self.cipher_combo.setCurrentIndex(index)
        top_row.addWidget(self.cipher_combo, 1)

        self.caesar_direction_combo = QComboBox(self)
        self.caesar_direction_combo.addItems(["DOPŘEDU", "DOZADU"])
        top_row.addWidget(self.caesar_direction_combo)

        self.caesar_shift_input = QSpinBox(self)
        self.caesar_shift_input.setRange(0, 999)
        self.caesar_shift_input.setValue(3)
        self.caesar_shift_input.setPrefix("POSUN: ")
        top_row.addWidget(self.caesar_shift_input)

        layout.addLayout(top_row)

        editors_row = QHBoxLayout()
        left_col = QVBoxLayout()
        right_col = QVBoxLayout()

        self.messages_label = QLabel("Zprávy", self)
        self.messages_label.setStyleSheet("color: #ead8b3;")
        left_col.addWidget(self.messages_label)

        self.messages_edit = QTextEdit(self)
        self.messages_edit.setAcceptRichText(False)
        self.messages_edit.setPlaceholderText(
            "Co řádek, to stanoviště:\n"
            "Najdi mapu u ohniště\n"
            "Klíč je pod třetím kamenem\n"
            "Další stopa čeká u potoka"
        )
        self.messages_edit.setStyleSheet("""
            QTextEdit {
                color: #f0e2c0;
                background: #071018;
                border: 1px solid #8a6938;
                border-radius: 6px;
                padding: 8px;
                font-family: Consolas, Menlo, Monaco, monospace;
                font-size: 12px;
            }
        """)
        left_col.addWidget(self.messages_edit, 1)

        self.station_cipher_label = QLabel("Šifry pro stanoviště", self)
        self.station_cipher_label.setStyleSheet("color: #ead8b3;")
        left_col.addWidget(self.station_cipher_label)

        self.station_cipher_scroll = QScrollArea(self)
        self.station_cipher_scroll.setWidgetResizable(True)
        self.station_cipher_scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
        self.station_cipher_scroll.setVerticalScrollBarPolicy(Qt.ScrollBarAsNeeded)
        self.station_cipher_scroll.setFrameShape(QScrollArea.NoFrame)
        self.station_cipher_scroll.setMinimumHeight(96)
        self.station_cipher_scroll.setMaximumHeight(190)
        self.station_cipher_scroll.setStyleSheet("""
            QScrollArea {
                background: #071018;
                border: 1px solid #8a6938;
                border-radius: 6px;
            }
            QScrollBar:vertical {
                background: rgba(20, 17, 12, 160);
                width: 10px;
                margin: 2px;
                border-radius: 5px;
            }
            QScrollBar::handle:vertical {
                background: #b89b68;
                border-radius: 5px;
                min-height: 28px;
            }
            QScrollBar::add-line:vertical,
            QScrollBar::sub-line:vertical {
                height: 0px;
            }
        """)

        self.station_cipher_content = QWidget()
        self.station_cipher_content.setStyleSheet("background: #071018;")
        self.station_cipher_layout = QVBoxLayout(self.station_cipher_content)
        self.station_cipher_layout.setContentsMargins(8, 8, 8, 8)
        self.station_cipher_layout.setSpacing(6)
        self.station_cipher_scroll.setWidget(self.station_cipher_content)
        left_col.addWidget(self.station_cipher_scroll)

        self.results_label = QLabel("Výsledky", self)
        self.results_label.setStyleSheet("color: #ead8b3;")
        right_col.addWidget(self.results_label)

        self.results_edit = QTextEdit(self)
        self.results_edit.setReadOnly(True)
        self.results_edit.setAcceptRichText(True)
        self.results_edit.setLineWrapMode(QTextEdit.WidgetWidth)
        self.results_edit.setWordWrapMode(QTextOption.WrapAnywhere)
        self.results_edit.setStyleSheet("""
            QTextEdit {
                color: #f0e2c0;
                background: #071018;
                border: 1px solid #8a6938;
                border-radius: 6px;
                padding: 8px;
                font-family: Consolas, Menlo, Monaco, monospace;
                font-size: 12px;
            }
        """)
        right_col.addWidget(self.results_edit, 1)

        editors_row.addLayout(left_col, 1)
        editors_row.addLayout(right_col, 1)
        layout.addLayout(editors_row, 1)

        button_row = QHBoxLayout()
        self.count_label = QLabel("0 zpráv", self)
        self.count_label.setStyleSheet("color: #a8a295;")
        button_row.addWidget(self.count_label, 1)

        self.encrypt_button = QPushButton("Zašifrovat vše", self)
        self.copy_button = QPushButton("Kopírovat", self)
        self.export_button = QPushButton("Export TXT", self)
        self.export_pdf_button = QPushButton("Export PDF", self)
        self.print_button = QPushButton("Náhled tisku", self)
        self.close_button = QPushButton("Zavřít", self)
        button_row.addWidget(self.encrypt_button)
        button_row.addWidget(self.copy_button)
        button_row.addWidget(self.export_button)
        button_row.addWidget(self.export_pdf_button)
        button_row.addWidget(self.print_button)
        button_row.addWidget(self.close_button)
        layout.addLayout(button_row)

        self.mode_combo.currentIndexChanged.connect(lambda *_args: self.update_controls_visibility())
        self.cipher_combo.currentTextChanged.connect(lambda *_args: self.update_controls_visibility())
        self.messages_edit.textChanged.connect(self.messages_changed)
        self.encrypt_button.clicked.connect(self.encrypt_all)
        self.copy_button.clicked.connect(self.copy_results)
        self.export_button.clicked.connect(self.export_results)
        self.export_pdf_button.clicked.connect(self.export_pdf_results)
        self.print_button.clicked.connect(self.print_results)
        self.close_button.clicked.connect(self.close)

        self.update_controls_visibility()
        self.update_message_count()
        self.refresh_results()
        self.apply_pirate_glass()
        self.station_cipher_content.setStyleSheet(
            self.station_cipher_content.styleSheet() + "QWidget { background: transparent; }"
        )

    def current_cipher_name(self) -> str:
        return self.cipher_combo.currentText().strip()

    def current_cipher_name_for_context(self) -> str:
        return str(getattr(self, "_batch_context_cipher_name", None) or self.current_cipher_name())

    def is_per_station_mode(self) -> bool:
        return self.mode_combo.currentIndex() == 1

    def update_controls_visibility(self):
        per_station = self.is_per_station_mode()
        self.station_cipher_label.setVisible(per_station)
        self.station_cipher_scroll.setVisible(per_station)
        self.cipher_label.setText("Výchozí šifra" if per_station else "Šifra")

        is_caesar = self.current_cipher_name() == "Caesarova šifra" or per_station
        self.caesar_direction_combo.setVisible(is_caesar)
        self.caesar_shift_input.setVisible(is_caesar)
        self.caesar_direction_combo.setToolTip("Použije se pro Caesarovu šifru v dávce.")
        self.caesar_shift_input.setToolTip("Použije se pro Caesarovu šifru v dávce.")
        try:
            self.sync_station_cipher_rows()
            self.update_message_count()
        except Exception:
            pass

    def update_caesar_visibility(self):
        self.update_controls_visibility()

    def message_lines(self) -> list[str]:
        return [line.strip() for line in self.messages_edit.toPlainText().splitlines() if line.strip()]

    def messages_changed(self):
        self.sync_station_cipher_rows()
        self.update_message_count()

    def clear_station_cipher_rows(self):
        while self.station_cipher_layout.count():
            item = self.station_cipher_layout.takeAt(0)
            widget = item.widget()
            if widget is not None:
                widget.deleteLater()
        self.station_cipher_combos = []

    def create_station_cipher_row(self, station_index: int, selected_cipher: str, available_names: list[str]):
        row = QWidget(self.station_cipher_content)
        row.setStyleSheet("background: transparent;")
        row_layout = QHBoxLayout(row)
        row_layout.setContentsMargins(0, 0, 0, 0)
        row_layout.setSpacing(8)

        number_label = QLabel(f"{station_index} -", row)
        number_label.setFixedWidth(42)
        number_label.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
        number_label.setStyleSheet("color: #ead8b3; background: transparent;")
        row_layout.addWidget(number_label)

        combo = QComboBox(row)
        combo.addItems(available_names)
        combo.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)
        combo.setStyleSheet("""
            QComboBox {
                color: #f0e2c0;
                background: #1e1e1e;
                border: 1px solid #8a6938;
                border-radius: 5px;
                padding: 5px 8px;
            }
            QComboBox QAbstractItemView {
                color: #f0e2c0;
                background: #071018;
                border: 1px solid #8a6938;
                selection-background-color: rgba(200, 154, 76, 155);
            }
        """)
        if selected_cipher in available_names:
            combo.setCurrentText(selected_cipher)
        combo.currentTextChanged.connect(lambda *_args: self.update_message_count())
        row_layout.addWidget(combo, 1)

        self.station_cipher_layout.addWidget(row)
        self.station_cipher_combos.append(combo)

    def sync_station_cipher_rows(self):
        count = len(self.message_lines())
        default_cipher = self.current_cipher_name()
        previous_default = str(getattr(self, "_station_cipher_last_default", "") or "")

        if getattr(self, "_station_cipher_row_count", None) == count and self.station_cipher_combos:
            if default_cipher != previous_default:
                for combo in self.station_cipher_combos:
                    if combo.currentText() == previous_default:
                        combo.setCurrentText(default_cipher)
            self._station_cipher_last_default = default_cipher
            return

        previous_values = self.station_cipher_lines()
        self.clear_station_cipher_rows()

        available_names = _batch_available_cipher_names(self)
        if count <= 0:
            info = QLabel("Zadej zprávy vlevo. Tady se pak ukáže 1 - až poslední stanoviště.", self.station_cipher_content)
            info.setWordWrap(True)
            info.setStyleSheet("color: #a8a295; background: transparent;")
            self.station_cipher_layout.addWidget(info)
            self._station_cipher_row_count = count
            self._station_cipher_last_default = default_cipher
            return

        for index in range(1, count + 1):
            selected = previous_values[index - 1] if index - 1 < len(previous_values) else default_cipher
            if selected not in available_names:
                selected = default_cipher
            self.create_station_cipher_row(index, selected, available_names)

        self.station_cipher_layout.addStretch(1)
        self._station_cipher_row_count = count
        self._station_cipher_last_default = default_cipher

    def station_cipher_lines(self) -> list[str]:
        return [combo.currentText().strip() for combo in self.station_cipher_combos]

    def cipher_for_station(self, station_index: int) -> str:
        fallback = self.current_cipher_name()
        if not self.is_per_station_mode():
            return fallback

        lines = self.station_cipher_lines()
        raw = lines[station_index - 1] if station_index - 1 < len(lines) else ""
        return _batch_resolve_cipher_name(raw, fallback, _batch_available_cipher_names(self))

    def invalid_station_cipher_lines(self, station_count: int) -> list[tuple[int, str]]:
        if not self.is_per_station_mode():
            return []

        invalid = []
        lines = self.station_cipher_lines()
        available = _batch_available_cipher_names(self)
        fallback = self.current_cipher_name()
        for index in range(1, station_count + 1):
            raw = lines[index - 1].strip() if index - 1 < len(lines) else ""
            if raw and not _batch_resolve_cipher_name(raw, fallback, available):
                invalid.append((index, raw))
        return invalid

    def update_message_count(self):
        count = len(self.message_lines())
        if self.is_per_station_mode():
            filled = len(self.station_cipher_lines())
            self.count_label.setText(f"{count} stanovišť, šifer zvoleno {filled}")
        else:
            self.count_label.setText(f"{count} stanovišť")

    def refresh_results(self):
        self.results_label.setText(f"Výsledky ({len(self.results)})")
        self.results_edit.setHtml(_batch_results_html(self.results, print_mode=False))
        self.results_edit.moveCursor(QTextCursor.Start)

    def encrypt_all(self):
        messages = self.message_lines()
        if not messages:
            QMessageBox.information(self, "Hromadné šifrování", "Nejdřív zadej zprávy.")
            return

        invalid = self.invalid_station_cipher_lines(len(messages))
        if invalid:
            details = "\n".join(f"{index}. stanoviště: {name}" for index, name in invalid[:12])
            if len(invalid) > 12:
                details += "\n..."
            QMessageBox.warning(
                self,
                "Neznámá šifra",
                "Některé názvy šifer v seznamu pro stanoviště nejdou jednoznačně poznat:\n\n"
                f"{details}\n\n"
                "Použij přesný název šifry nebo jednoznačný začátek názvu.",
            )
            return

        results = []
        for index, message in enumerate(messages, start=1):
            cipher_name = self.cipher_for_station(index)
            context = _batch_cipher_context_for_name(self, cipher_name)
            output = _batch_encrypt_message(cipher_name, message, context)
            item = {
                "index": index,
                "cipher": cipher_name,
                "context": dict(context),
                "input": message,
                "output": output,
                "image_path": _batch_render_cipher_image(cipher_name, output, message, index),
            }
            results.append(item)
            _batch_save_result_to_history(item)

        self.results = results
        self.refresh_results()

        try:
            if self.owner_window is not None and hasattr(self.owner_window, "central"):
                self.owner_window.central.update_status()
            mode = "po stanovištích" if self.is_per_station_mode() else "jedna šifra"
            _sifrator_debug_log(f"Hromadné šifrování hotovo: režim={mode}, zpráv={len(results)}")
        except Exception:
            pass

    def copy_results(self):
        text = _batch_results_plain_text(self.results)
        if not text.strip():
            QMessageBox.information(self, "Kopírování", "Výsledky zatím nejsou připravené.")
            return
        clipboard = QApplication.clipboard()
        if clipboard is not None:
            clipboard.setText(text)

    def export_results(self):
        text = _batch_results_plain_text(self.results)
        if not text.strip():
            QMessageBox.information(self, "Export", "Výsledky zatím nejsou připravené.")
            return

        default_name = f"sifrator_hromadne_{time.strftime('%Y%m%d_%H%M%S')}.txt"
        path, _ = QFileDialog.getSaveFileName(
            self,
            "Exportovat hromadné šifrování jako TXT",
            default_name,
            "Textový soubor (*.txt);;Všechny soubory (*.*)",
        )
        if not path:
            return
        if not path.lower().endswith(".txt"):
            path += ".txt"

        try:
            with open(path, "w", encoding="utf-8") as file:
                file.write(text.strip() + "\n")
            QMessageBox.information(self, "Export", f"Výsledky byly uloženy do:\n{path}")
        except Exception as error:
            QMessageBox.warning(self, "Export", f"Výsledky se nepodařilo uložit:\n{error}")

    def export_pdf_results(self):
        _batch_export_pdf(self, self.results, "A4", "Na výšku", _batch_default_print_settings())

    def print_results(self):
        if not self.results:
            QMessageBox.information(self, "Tisk", "Výsledky zatím nejsou připravené.")
            return

        try:
            from PySide6.QtPrintSupport import QPrinter, QPrintDialog
            from PySide6.QtWidgets import QFrame, QFormLayout
        except Exception as error:
            QMessageBox.warning(self, "Tisk není dostupný", f"Nepodařilo se načíst podporu tisku:\n{error}")
            return

        dialog = QDialog(self)
        dialog.setWindowTitle("Náhled tisku hromadného šifrování – Šifrátor Mraveniště")
        dialog.setModal(True)

        try:
            screen = self.window().screen() if self.window() is not None else QApplication.primaryScreen()
            available = screen.availableGeometry() if screen is not None else None
        except Exception:
            available = None

        if available is not None:
            dialog_w = min(1280, max(760, int(available.width() * 0.94)))
            dialog_h = min(850, max(560, int(available.height() * 0.90)))
        else:
            dialog_w, dialog_h = 1280, 850

        compact = dialog_w < 1080 or dialog_h < 720
        dialog.resize(dialog_w, dialog_h)
        dialog.setMinimumSize(min(760, dialog_w), min(540, dialog_h))
        dialog.setStyleSheet("""
            QDialog {
                background-color: #0a1626;
                color: #f3ddaa;
            }
            QLabel {
                color: #f3ddaa;
                background: transparent;
            }
            QFrame#sidePanel, QFrame#pageFrame {
                background-color: rgba(10, 20, 34, 230);
                border: 1px solid #9b6b2f;
                border-radius: 12px;
            }
            QFrame#sectionFrame {
                background-color: rgba(11, 31, 52, 200);
                border: 1px solid rgba(179, 130, 55, 160);
                border-radius: 10px;
            }
            QCheckBox {
                spacing: 10px;
                padding: 3px 0;
                font-size: 13px;
            }
            QCheckBox::indicator {
                width: 16px;
                height: 16px;
            }
            QCheckBox::indicator:unchecked {
                border: 1px solid #d4ae67;
                background: #10263e;
            }
            QCheckBox::indicator:checked {
                border: 1px solid #d4ae67;
                background: #0f8aa8;
            }
            QComboBox, QSpinBox, QLineEdit {
                background-color: #10263e;
                color: #f6e7bf;
                border: 1px solid #c49344;
                border-radius: 8px;
                padding: 6px 8px;
                min-height: 20px;
            }
            QComboBox QAbstractItemView {
                background-color: #10263e;
                color: #f6e7bf;
                border: 1px solid #c49344;
                selection-background-color: #155d75;
            }
            QScrollArea#paperPreviewArea {
                background: #07111f;
                border: 1px solid #c49344;
                border-radius: 8px;
            }
            QScrollArea {
                background: transparent;
                border: none;
            }
            QPushButton {
                background-color: #10263e;
                color: #f6e7bf;
                border: 1px solid #c49344;
                border-radius: 10px;
                padding: 10px 18px;
                font-weight: bold;
                min-height: 22px;
            }
            QPushButton:hover {
                background-color: #144a63;
            }
            QPushButton#cancelButton {
                background-color: #4a1520;
                border-color: #b56a5c;
            }
            QPushButton#cancelButton:hover {
                background-color: #6a1b2c;
            }
        """)

        main_layout = QVBoxLayout(dialog)
        main_layout.setContentsMargins(16, 16, 16, 16)
        main_layout.setSpacing(12)

        title = QLabel("Hromadné šifrování – náhled tisku", dialog)
        title.setStyleSheet("font-weight: 700; font-size: 18px; color: #f8e8c2;")
        main_layout.addWidget(title)

        content_layout = QVBoxLayout() if compact else QHBoxLayout()
        content_layout.setSpacing(10 if compact else 14)
        main_layout.addLayout(content_layout, 1)

        side_scroll = QScrollArea(dialog)
        side_scroll.setWidgetResizable(True)
        side_scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
        if compact:
            side_scroll.setMinimumHeight(min(220, max(170, int(dialog_h * 0.28))))
            side_scroll.setMaximumHeight(min(330, max(220, int(dialog_h * 0.38))))
        else:
            side_scroll.setMinimumWidth(min(350, max(300, int(dialog_w * 0.27))))
            side_scroll.setMaximumWidth(min(410, max(330, int(dialog_w * 0.32))))

        side_panel = QFrame()
        side_panel.setObjectName("sidePanel")
        side_layout = QVBoxLayout(side_panel)
        side_layout.setContentsMargins(14, 14, 14, 14)
        side_layout.setSpacing(12)
        side_scroll.setWidget(side_panel)

        def make_block(title_text):
            frame = QFrame(side_panel)
            frame.setObjectName("sectionFrame")
            block_layout = QVBoxLayout(frame)
            block_layout.setContentsMargins(12, 12, 12, 12)
            block_layout.setSpacing(8)
            label = QLabel(title_text, frame)
            label.setStyleSheet("font-weight: 700; font-size: 15px;")
            block_layout.addWidget(label)
            return frame, block_layout

        block_page, block_page_layout = make_block("Nastavení stránky")
        page_form = QFormLayout()
        page_form.setHorizontalSpacing(10)
        page_form.setVerticalSpacing(10)
        paper_combo = QComboBox(block_page)
        paper_combo.addItems(["A4", "A5", "A3", "Letter", "Legal"])
        paper_combo.setCurrentText("A4")
        orientation_combo = QComboBox(block_page)
        orientation_combo.addItems(["Na výšku", "Na šířku"])
        orientation_combo.setCurrentText("Na výšku")
        page_form.addRow("Papír:", paper_combo)
        page_form.addRow("Orientace:", orientation_combo)
        block_page_layout.addLayout(page_form)
        side_layout.addWidget(block_page)

        block_content, block_content_layout = make_block("Co tisknout")
        show_station_title_check = QCheckBox("Nadpis stanoviště", block_content)
        show_input_check = QCheckBox("Původní zpráva", block_content)
        show_output_check = QCheckBox("Zašifrovaný text", block_content)
        show_images_check = QCheckBox("Obrázkový náhled", block_content)
        for check in (show_station_title_check, show_input_check, show_output_check, show_images_check):
            check.setChecked(True)
            block_content_layout.addWidget(check)
        side_layout.addWidget(block_content)

        block_texts, block_texts_layout = make_block("Texty")
        texts_form = QFormLayout()
        texts_form.setHorizontalSpacing(8)
        texts_form.setVerticalSpacing(8)
        station_title_edit = QLineEdit("{index}. stanoviště | {cipher}", block_texts)
        input_label_edit = QLineEdit("Zpráva:", block_texts)
        output_label_edit = QLineEdit("Zašifrováno:", block_texts)
        station_title_edit.setToolTip("Můžeš použít {index}, {cipher}, {context}. Prázdné pole nadpis skryje.")
        input_label_edit.setToolTip("Prázdné pole skryje popisek původní zprávy.")
        output_label_edit.setToolTip("Prázdné pole skryje popisek zašifrovaného textu.")
        texts_form.addRow("Nadpis:", station_title_edit)
        texts_form.addRow("Zpráva:", input_label_edit)
        texts_form.addRow("Výsledek:", output_label_edit)
        block_texts_layout.addLayout(texts_form)
        side_layout.addWidget(block_texts)

        block_layout_opts, block_layout_opts_layout = make_block("Vzhled listu")
        show_frames_check = QCheckBox("Zobrazit rámečky stanovišť", block_layout_opts)
        show_frames_check.setChecked(True)
        block_layout_opts_layout.addWidget(show_frames_check)

        size_form = QFormLayout()
        heading_size_spin = QSpinBox(block_layout_opts)
        heading_size_spin.setRange(8, 34)
        heading_size_spin.setValue(14)
        heading_size_spin.setSuffix(" bodů")
        text_size_spin = QSpinBox(block_layout_opts)
        text_size_spin.setRange(8, 28)
        text_size_spin.setValue(12)
        text_size_spin.setSuffix(" bodů")
        cipher_scale_spin = QSpinBox(block_layout_opts)
        cipher_scale_spin.setRange(20, 100)
        cipher_scale_spin.setValue(90)
        cipher_scale_spin.setSuffix(" %")
        size_form.addRow("Nadpis:", heading_size_spin)
        size_form.addRow("Text:", text_size_spin)
        size_form.addRow("Obrázky:", cipher_scale_spin)
        block_layout_opts_layout.addLayout(size_form)
        side_layout.addWidget(block_layout_opts)

        block_zoom, block_zoom_layout = make_block("Náhled")
        zoom_row = QHBoxLayout()
        zoom_out_button = QPushButton("−", block_zoom)
        zoom_in_button = QPushButton("+", block_zoom)
        zoom_fit_button = QPushButton("Přizpůsobit", block_zoom)
        zoom_label = QLabel("Automaticky", block_zoom)
        zoom_label.setStyleSheet("color: #d8c392;")
        zoom_out_button.setFixedWidth(46)
        zoom_in_button.setFixedWidth(46)
        zoom_row.addWidget(zoom_out_button)
        zoom_row.addWidget(zoom_in_button)
        zoom_row.addWidget(zoom_fit_button)
        block_zoom_layout.addLayout(zoom_row)
        block_zoom_layout.addWidget(zoom_label)
        side_layout.addWidget(block_zoom)

        side_hint = QLabel(
            "Náhled vpravo používá stejný papír jako finální tisk. Tlačítka + a − mění jen zobrazení náhledu.",
            side_panel,
        )
        side_hint.setWordWrap(True)
        side_hint.setStyleSheet("color: #cdb57e; font-size: 12px;")
        side_layout.addWidget(side_hint)
        side_layout.addStretch(1)

        content_layout.addWidget(side_scroll)

        page_frame = QFrame(dialog)
        page_frame.setObjectName("pageFrame")
        page_layout = QVBoxLayout(page_frame)
        page_layout.setContentsMargins(10 if compact else 18, 10 if compact else 18, 10 if compact else 18, 10 if compact else 18)
        page_layout.setSpacing(8 if compact else 10)
        preview_header = QLabel("Náhled stránky", page_frame)
        preview_header.setStyleSheet("font-size: 16px; font-weight: 700;")
        page_layout.addWidget(preview_header)
        preview_detail = QLabel(_print_page_setup_label("A4", "Na výšku"), page_frame)
        preview_detail.setStyleSheet("color: #d8c392; font-size: 12px;")
        page_layout.addWidget(preview_detail)
        preview_scroll = QScrollArea(page_frame)
        preview_scroll.setObjectName("paperPreviewArea")
        preview_scroll.setWidgetResizable(False)
        preview_scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarAsNeeded)
        preview_scroll.setVerticalScrollBarPolicy(Qt.ScrollBarAsNeeded)
        preview_scroll.setAlignment(Qt.AlignHCenter | Qt.AlignTop)
        preview = _PrintPaperPreviewWidget(preview_scroll)
        preview_scroll.setWidget(preview)
        page_layout.addWidget(preview_scroll, 1)
        content_layout.addWidget(page_frame, 1)

        buttons_layout = QHBoxLayout()
        buttons_layout.addStretch(1)
        export_pdf_button = QPushButton("Export PDF", dialog)
        print_button = QPushButton("Tisknout", dialog)
        cancel_button = QPushButton("Zrušit", dialog)
        cancel_button.setObjectName("cancelButton")
        export_pdf_button.setMinimumWidth(150)
        print_button.setMinimumWidth(150)
        cancel_button.setMinimumWidth(120)
        buttons_layout.addWidget(export_pdf_button)
        buttons_layout.addWidget(print_button)
        buttons_layout.addWidget(cancel_button)
        main_layout.addLayout(buttons_layout)

        def current_page_setup():
            return paper_combo.currentText(), orientation_combo.currentText()

        def current_settings():
            return {
                "show_frames": show_frames_check.isChecked(),
                "show_station_title": show_station_title_check.isChecked(),
                "show_input": show_input_check.isChecked(),
                "show_output": show_output_check.isChecked(),
                "show_images": show_images_check.isChecked(),
                "station_title_template": station_title_edit.text(),
                "input_label": input_label_edit.text(),
                "output_label": output_label_edit.text(),
                "heading_size": heading_size_spin.value(),
                "text_size": text_size_spin.value(),
                "cipher_scale": cipher_scale_spin.value(),
            }

        def update_zoom_label():
            zoom_label.setText(f"Měřítko náhledu: {preview.zoom_percent()} %")

        def update_preview():
            paper_name, orientation_name = current_page_setup()
            preview_detail.setText(_print_page_setup_label(paper_name, orientation_name))
            heading_size_spin.setEnabled(show_station_title_check.isChecked())
            cipher_scale_spin.setEnabled(show_images_check.isChecked())
            station_title_edit.setEnabled(show_station_title_check.isChecked())
            input_label_edit.setEnabled(show_input_check.isChecked())
            output_label_edit.setEnabled(show_output_check.isChecked())
            document = _batch_build_print_document(self.results, paper_name, orientation_name, current_settings())
            preview.set_preview_document(document, paper_name, orientation_name)
            update_zoom_label()

        preview_timer = QTimer(dialog)
        preview_timer.setSingleShot(True)
        preview_timer.timeout.connect(update_preview)

        def update_preview_now(*_args):
            preview_timer.stop()
            update_preview()

        def schedule_update_preview(*_args):
            preview_timer.start(70)

        paper_combo.currentTextChanged.connect(update_preview_now)
        orientation_combo.currentTextChanged.connect(update_preview_now)
        show_station_title_check.toggled.connect(update_preview_now)
        show_input_check.toggled.connect(update_preview_now)
        show_output_check.toggled.connect(update_preview_now)
        show_images_check.toggled.connect(update_preview_now)
        show_frames_check.toggled.connect(update_preview_now)
        station_title_edit.textChanged.connect(schedule_update_preview)
        input_label_edit.textChanged.connect(schedule_update_preview)
        output_label_edit.textChanged.connect(schedule_update_preview)
        heading_size_spin.valueChanged.connect(update_preview_now)
        text_size_spin.valueChanged.connect(update_preview_now)
        cipher_scale_spin.valueChanged.connect(update_preview_now)

        zoom_out_button.clicked.connect(lambda: (preview.set_zoom_percent(preview.zoom_percent() - 10), update_zoom_label()))
        zoom_in_button.clicked.connect(lambda: (preview.set_zoom_percent(preview.zoom_percent() + 10), update_zoom_label()))
        zoom_fit_button.clicked.connect(lambda: (preview.reset_zoom(), update_preview()))

        cancel_button.clicked.connect(dialog.reject)

        def do_export_pdf():
            paper_name, orientation_name = current_page_setup()
            _batch_export_pdf(dialog, self.results, paper_name, orientation_name, current_settings())

        def do_print():
            paper_name, orientation_name = current_page_setup()
            document = _batch_build_print_document(self.results, paper_name, orientation_name, current_settings())
            printer = QPrinter(QPrinter.HighResolution)
            _print_apply_page_setup_to_printer(printer, paper_name, orientation_name)
            print_dialog = QPrintDialog(printer, dialog)
            print_dialog.setWindowTitle("Tisk hromadného šifrování")
            if print_dialog.exec() != QDialog.Accepted:
                return
            document.print_(printer)
            dialog.accept()

        export_pdf_button.clicked.connect(do_export_pdf)
        print_button.clicked.connect(do_print)
        preview.set_message("Připravuji náhled...", "A4", "Na výšku")
        update_preview_now()
        dialog.exec()


def _footer_prepare_links(self):
    """Připraví samostatné klikací texty ve spodním řádku."""
    if (
        hasattr(self, "batch_status_link")
        and hasattr(self, "cipher_overview_status_link")
        and hasattr(self, "planner_status_link")
        and hasattr(self, "history_status_link")
        and hasattr(self, "log_status_link")
        and hasattr(self, "update_status_link")
    ):
        return

    base_style = "background: transparent; color: #d9c697;"
    link_style = "background: transparent; color: #f3d79a;"

    self.batch_status_link = QLabel(self)
    self.batch_status_link.setStyleSheet(link_style)
    self.batch_status_link.setCursor(Qt.PointingHandCursor)
    self.batch_status_link.setToolTip("Otevřít hromadné šifrování.")
    self.batch_status_link.mousePressEvent = lambda event: _footer_batch_clicked(self, event)

    self.cipher_overview_status_link = QLabel(self)
    self.cipher_overview_status_link.setStyleSheet(link_style)
    self.cipher_overview_status_link.setCursor(Qt.PointingHandCursor)
    self.cipher_overview_status_link.setToolTip("Otevřít přehled použití, obtížnosti, věku a poznámek šifer.")
    self.cipher_overview_status_link.mousePressEvent = lambda event: _footer_cipher_overview_clicked(self, event)

    self.planner_status_link = QLabel(self)
    self.planner_status_link.setStyleSheet(link_style)
    self.planner_status_link.setCursor(Qt.PointingHandCursor)
    self.planner_status_link.setToolTip("Otevřít plánovač šifer na dny tábora.")
    self.planner_status_link.mousePressEvent = lambda event: _footer_planner_clicked(self, event)

    self.history_status_link = QLabel(self)
    self.history_status_link.setStyleSheet(link_style)
    self.history_status_link.setCursor(Qt.PointingHandCursor)
    self.history_status_link.setToolTip("Otevřít historii zašifrovaných zpráv.")
    self.history_status_link.mousePressEvent = lambda event: _footer_history_clicked(self, event)

    self.log_status_link = QLabel(self)
    self.log_status_link.setStyleSheet(link_style)
    self.log_status_link.setCursor(Qt.PointingHandCursor)
    self.log_status_link.setToolTip("Kliknutím zapneš logování a otevřeš live logy.")
    self.log_status_link.mousePressEvent = lambda event: _footer_log_clicked(self, event)

    self.src_status_label = QLabel(self)
    self.src_status_label.setStyleSheet(base_style)

    self.update_status_link = QLabel(self)
    self.update_status_link.setStyleSheet(link_style)
    self.update_status_link.setCursor(Qt.PointingHandCursor)
    self.update_status_link.setToolTip("Ručně zkontrolovat aktualizace.")
    self.update_status_link.mousePressEvent = lambda event: _footer_update_clicked(self, event)

    for label in (
        self.batch_status_link,
        self.cipher_overview_status_link,
        self.planner_status_link,
        self.history_status_link,
        self.log_status_link,
        self.src_status_label,
        self.update_status_link,
    ):
        label.show()
        label.raise_()


def _footer_window(self):
    try:
        window = self.window()
        if isinstance(window, SifratorWindow):
            return window
    except Exception:
        pass
    return None


def _footer_log_clicked(self, event=None):
    window = _footer_window(self)
    if window is not None and hasattr(window, "show_live_log_window"):
        window.show_live_log_window()


def _footer_batch_clicked(self, event=None):
    window = _footer_window(self)
    if window is not None and hasattr(window, "show_batch_encrypt_window"):
        window.show_batch_encrypt_window()


def _footer_cipher_overview_clicked(self, event=None):
    window = _footer_window(self)
    if window is not None and hasattr(window, "show_cipher_overview_window"):
        window.show_cipher_overview_window()


def _footer_planner_clicked(self, event=None):
    window = _footer_window(self)
    if window is not None and hasattr(window, "show_camp_planner_window"):
        window.show_camp_planner_window()


def _footer_history_clicked(self, event=None):
    window = _footer_window(self)
    if window is not None and hasattr(window, "show_history_window"):
        window.show_history_window()


def _footer_update_clicked(self, event=None):
    window = _footer_window(self)
    if window is not None and hasattr(window, "manual_check_updates"):
        window.manual_check_updates()


def _footer_create_widgets(self):
    _FOOTER_ORIGINAL_CREATE_WIDGETS(self)
    _footer_prepare_links(self)
    self.update_status()


def _footer_layout_status(self):
    _footer_prepare_links(self)

    y_rect = self.sr(52, 881, 950, 22)
    x = y_rect.x()
    y = y_rect.y()
    h = y_rect.height()
    max_right = max(x + self.sr(0, 0, 1560, 0).width(), self.width() - self.fs(20))

    font = QFont("Georgia", self.fs(10))
    footer_labels = (
        self.status,
        self.batch_status_link,
        self.cipher_overview_status_link,
        self.planner_status_link,
        self.history_status_link,
        self.log_status_link,
        self.src_status_label,
        self.update_status_link,
    )

    for label in footer_labels:
        label.setFont(font)
        label.setFixedHeight(h)

    fm = self.status.fontMetrics()
    available_total = max(300, max_right - x)
    batch_text = self.batch_status_link.text()
    cipher_overview_text = self.cipher_overview_status_link.text()
    planner_text = self.planner_status_link.text()
    history_text = self.history_status_link.text()
    log_text = self.log_status_link.text()
    src_text = self.src_status_label.text()
    update_text = self.update_status_link.text()
    fixed_w = (
        fm.horizontalAdvance(batch_text)
        + fm.horizontalAdvance(cipher_overview_text)
        + fm.horizontalAdvance(planner_text)
        + fm.horizontalAdvance(history_text)
        + fm.horizontalAdvance(log_text)
        + fm.horizontalAdvance(src_text)
        + fm.horizontalAdvance(update_text)
        + self.fs(86)
    )
    max_selected_w = max(self.fs(130), available_total - fixed_w)

    selected_text = self.status.property("full_status_text") or self.status.text()
    selected_text = fm.elidedText(str(selected_text), Qt.ElideRight, max_selected_w)
    self.status.setText(selected_text)

    for label in footer_labels:
        label.adjustSize()

    gap = self.fs(10)
    for label in footer_labels:
        width = min(label.width() + self.fs(4), max(40, max_right - x))
        label.setGeometry(x, y, width, h)
        label.show()
        label.raise_()
        x += width + gap


def _footer_update_status(self):
    _footer_prepare_links(self)
    selected_text = self.selected_cipher if self.selected_cipher else "Žádná"
    _cipher_overview_refresh_usage_marks(self)

    window = _footer_window(self)
    logging_enabled = bool(window and hasattr(window, "is_live_logging_enabled") and window.is_live_logging_enabled())
    logging_text = "Zapnuto" if logging_enabled else "Vypnuto"
    used_count = len(_cipher_used_details())
    known_count = len(_cipher_known_names())

    self.status.setProperty("full_status_text", f"VYBRANÁ ŠIFRA:  {selected_text}   |")
    self.status.setText(f"VYBRANÁ ŠIFRA:  {selected_text}   |")
    self.batch_status_link.setText("HROMADNĚ   |")
    self.cipher_overview_status_link.setText(f"ŠIFRY:  {used_count}/{known_count}   |")
    self.planner_status_link.setText(f"PLÁN:  {_planner_used_count()}/{len(_planner_known_cipher_names())}   |")
    self.history_status_link.setText(f"HISTORIE:  {_history_count()}   |")
    self.log_status_link.setText(f"LOGOVÁNÍ:  {logging_text}   |")
    self.src_status_label.setText("SRC SLOŽKA:  Nalezena   |")
    self.update_status_link.setText("AKTUALIZACE")
    self.batch_status_link.setToolTip("Otevřít hromadné šifrování.")
    self.cipher_overview_status_link.setToolTip("Otevřít přehled použití, obtížnosti, věku a poznámek šifer.")
    self.planner_status_link.setToolTip("Otevřít plánovač šifer na dny tábora.")
    self.history_status_link.setToolTip("Otevřít historii zašifrovaných zpráv.")
    self.log_status_link.setToolTip(
        "Kliknutím otevřeš live logy. Zavřením okna se logování vypne."
        if logging_enabled else
        "Kliknutím zapneš logování a otevřeš live logy."
    )
    _footer_layout_status(self)


def _footer_update_layout_positions(self):
    _FOOTER_ORIGINAL_UPDATE_LAYOUT_POSITIONS(self)
    _footer_layout_status(self)


SifratorSkinWidget.create_widgets = _footer_create_widgets
SifratorSkinWidget.update_status = _footer_update_status
SifratorSkinWidget.update_layout_positions = _footer_update_layout_positions


try:
    _HISTORY_ORIGINAL_SET_RESULT_OUTPUT = SifratorSkinWidget.set_result_output

    def _history_set_result_output(self, result: str):
        self._history_last_result_output = str(result or "")
        return _HISTORY_ORIGINAL_SET_RESULT_OUTPUT(self, result)

    SifratorSkinWidget.set_result_output = _history_set_result_output
except Exception:
    pass


try:
    _HISTORY_ORIGINAL_AUTO_ENCRYPT_ACTION = SifratorSkinWidget.auto_encrypt_action

    def _history_auto_encrypt_action(self):
        result = _HISTORY_ORIGINAL_AUTO_ENCRYPT_ACTION(self)
        _history_schedule_capture(self, 900)
        return result

    SifratorSkinWidget.auto_encrypt_action = _history_auto_encrypt_action
except Exception:
    pass


_AUTOMATIC_ORIGINAL_CHECK_UPDATES = SifratorWindow.check_updates_after_start


def _window_is_live_logging_enabled(self) -> bool:
    return bool(getattr(self, "_live_logging_enabled", False))


def _window_set_live_logging_enabled(self, enabled: bool):
    self._live_logging_enabled = bool(enabled)
    try:
        if hasattr(self, "central"):
            self.central.update_status()
    except Exception:
        pass


def _window_write_live_log(self, message: str):
    _sifrator_debug_log(message)
    try:
        dialog = getattr(self, "_live_log_dialog", None)
        if dialog is not None:
            dialog.refresh_log()
    except Exception:
        pass


def _window_show_live_log_window(self):
    dialog = getattr(self, "_live_log_dialog", None)
    if dialog is not None and dialog.isVisible():
        dialog.raise_()
        dialog.activateWindow()
        return

    self.set_live_logging_enabled(True)
    self._live_log_dialog = LiveLogDialog(self)
    self.write_live_log("Logování bylo zapnuto přes spodní stavový řádek.")
    self._live_log_dialog.show()
    self._live_log_dialog.raise_()
    self._live_log_dialog.activateWindow()


def _window_show_history_window(self):
    dialog = getattr(self, "_history_dialog", None)
    if dialog is not None:
        dialog.refresh_history()
    else:
        dialog = HistoryDialog(self)
        self._history_dialog = dialog
    self._show_embedded_dialog_page(dialog, "Historie zpráv")


def _window_show_cipher_overview_window(self):
    dialog = getattr(self, "_cipher_overview_dialog", None)
    if dialog is not None:
        dialog.refresh_overview()
    else:
        dialog = CipherOverviewDialog(self)
        self._cipher_overview_dialog = dialog
    self._show_embedded_dialog_page(dialog, "Přehled šifer")


def _window_show_batch_encrypt_window(self):
    dialog = getattr(self, "_batch_encrypt_dialog", None)
    if dialog is None:
        dialog = BatchEncryptDialog(self)
        self._batch_encrypt_dialog = dialog
    self._show_embedded_dialog_page(dialog, "Hromadné šifrování")


def _window_show_camp_planner_window(self):
    dialog = getattr(self, "_camp_planner_dialog", None)
    if dialog is not None:
        dialog.refresh_overview(save=False)
    else:
        dialog = CampPlannerDialog(self)
        self._camp_planner_dialog = dialog
    self._show_embedded_dialog_page(dialog, "Plánovač tábora")


def _window_show_update_offer(self, update_data: dict, manual: bool = False):
    remote_version = update_data.get("version", "")
    notes = update_data.get("notes", "")
    platform_key = update_data.get("platform_key", "")
    file_name = update_data.get("file_name", "")

    message = (
        f"Je dostupná nová verze {remote_version}.\n\n"
        f"Aktuální verze: {APP_VERSION}\n"
        f"Platforma: {platform_key}\n"
        f"Balíček: {file_name}\n\n"
        f"{notes}\n\n"
        "Chceš aplikaci aktualizovat?"
    )

    answer = QMessageBox.question(
        self,
        "Dostupná aktualizace",
        message,
        QMessageBox.Yes | QMessageBox.No,
        QMessageBox.Yes,
    )

    if answer == QMessageBox.Yes:
        try:
            self.write_live_log(f"Uživatel potvrdil aktualizaci na verzi {remote_version}.")
            update_manager.download_and_install_update(update_data)
        except Exception as error:
            self.write_live_log(f"Chyba aktualizace: {type(error).__name__}: {error}")
            QMessageBox.critical(
                self,
                "Chyba aktualizace",
                f"Aktualizaci se nepodařilo dokončit:\n\n{error}",
            )
    else:
        self.write_live_log("Uživatel aktualizaci odmítl.")


def _window_check_updates_common(self, manual: bool = False):
    platform_key = update_manager.get_platform_key() if hasattr(update_manager, "get_platform_key") else "unknown"
    self.write_live_log(
        ("Ruční" if manual else "Automatická") +
        f" kontrola aktualizací: current={APP_VERSION}, platform={platform_key}"
    )

    update_data = update_manager.check_for_update(APP_VERSION)
    if update_data:
        self.write_live_log(
            f"Aktualizace dostupná: version={update_data.get('version')}, "
            f"platform={update_data.get('platform_key')}, file={update_data.get('file_name')}"
        )
        _window_show_update_offer(self, update_data, manual=manual)
        return

    self.write_live_log("Aktualizace není dostupná nebo se nepodařilo načíst update.json.")

    if manual:
        log_path = _sifrator_debug_log_path()
        QMessageBox.information(
            self,
            "Aktualizace",
            "Novější verze nebyla nalezena.\n\n"
            f"Aktuální verze: {APP_VERSION}\n"
            f"Platforma: {platform_key}\n\n"
            f"Diagnostický log:\n{log_path}",
        )


def _window_check_updates_after_start(self):
    _window_check_updates_common(self, manual=False)


def _window_manual_check_updates(self):
    _window_check_updates_common(self, manual=True)


SifratorWindow.is_live_logging_enabled = _window_is_live_logging_enabled
SifratorWindow.set_live_logging_enabled = _window_set_live_logging_enabled
SifratorWindow.write_live_log = _window_write_live_log
SifratorWindow.show_live_log_window = _window_show_live_log_window
SifratorWindow.show_history_window = _window_show_history_window
SifratorWindow.show_cipher_overview_window = _window_show_cipher_overview_window
SifratorWindow.show_batch_encrypt_window = _window_show_batch_encrypt_window
SifratorWindow.show_camp_planner_window = _window_show_camp_planner_window
SifratorWindow.manual_check_updates = _window_manual_check_updates
SifratorWindow.check_updates_after_start = _window_check_updates_after_start


def _user_data_backup_default_dir() -> str:
    desktop = os.path.join(os.path.expanduser("~"), "Desktop")
    if os.path.isdir(desktop):
        return desktop
    home = os.path.expanduser("~")
    if os.path.isdir(home):
        return home
    return get_user_data_dir()


def _user_data_format_size(byte_count: int) -> str:
    size = float(max(0, int(byte_count or 0)))
    for unit in ("B", "KB", "MB", "GB"):
        if size < 1024 or unit == "GB":
            if unit == "B":
                return f"{int(size)} {unit}"
            return f"{size:.1f} {unit}"
        size /= 1024
    return f"{size:.1f} GB"


class UserDataBackupDialog(PirateModuleDialog):
    """Pirátský panel pro export a import uživatelských dat."""

    def __init__(self, owner_window):
        super().__init__(
            owner_window,
            "menu_BG.png",
            (
                (0.058, 0.404, 0.48),
                (0.949, 0.838, 1.25),
            ),
        )
        self.owner_window = owner_window
        self.setWindowTitle("Záloha uživatelských dat")
        self.resize(920, 560)
        self.setMinimumSize(760, 500)
        self.setStyleSheet(self.styleSheet() + """
            QLabel#backupTitle {
                color: #f4dea4;
                font-family: Georgia;
                font-size: 26px;
                font-weight: bold;
                letter-spacing: 1px;
            }
            QLabel#backupSubtitle {
                color: #d8c59a;
                font-family: Georgia;
                font-size: 13px;
                font-style: italic;
            }
            QLabel#backupSmall {
                color: #cdb98d;
                font-family: Georgia;
                font-size: 11px;
            }
            QFrame#backupSummary {
                background-color: rgba(3, 18, 27, 82);
                border: 1px solid rgba(211, 162, 78, 145);
                border-radius: 11px;
            }
            QLabel#backupSummaryName {
                color: #d4aa64;
                font-family: Georgia;
                font-size: 10px;
                font-weight: bold;
            }
            QLabel#backupSummaryValue {
                color: #f1e4c4;
                font-family: Georgia;
                font-size: 11px;
            }
            QFrame#backupPanel {
                background-color: rgba(3, 18, 27, 92);
                border: 1px solid rgba(211, 162, 78, 175);
                border-radius: 10px;
            }
            QLabel#backupBadge {
                min-width: 58px;
                min-height: 58px;
                max-width: 58px;
                max-height: 58px;
                color: #fff0bd;
                background-color: rgba(14, 73, 83, 185);
                border: 1px solid rgba(238, 191, 94, 225);
                border-radius: 29px;
                font-family: Georgia;
                font-size: 14px;
                font-weight: bold;
            }
            QLabel#backupPanelTitle {
                color: #f4dea4;
                font-family: Georgia;
                font-size: 17px;
                font-weight: bold;
            }
            QLabel#backupPanelText {
                color: #ead8b3;
                font-family: Georgia;
                font-size: 12px;
            }
            QLabel#backupPanelItem {
                color: #d9c69c;
                font-family: Georgia;
                font-size: 11px;
            }
            QFrame#backupNote {
                background-color: rgba(5, 24, 33, 96);
                border: 1px solid rgba(205, 159, 78, 130);
                border-radius: 9px;
            }
            QPushButton#backupPrimary {
                min-height: 38px;
                color: #fff0bd;
                background-color: rgba(14, 73, 83, 205);
                border: 1px solid rgba(238, 191, 94, 230);
                border-radius: 9px;
                padding: 9px 18px;
                font-family: Georgia;
                font-weight: bold;
            }
            QPushButton#backupPrimary:hover {
                background-color: rgba(20, 96, 105, 235);
                border: 2px solid #f3d79a;
            }
            QPushButton#backupDanger {
                min-height: 38px;
                color: #f0c2ae;
                background-color: rgba(95, 30, 36, 172);
                border: 1px solid #a95a55;
                border-radius: 9px;
                padding: 9px 18px;
                font-family: Georgia;
                font-weight: bold;
            }
            QPushButton#backupDanger:hover {
                color: #ffe3d8;
                background-color: rgba(128, 40, 46, 220);
                border: 2px solid #f0a39b;
            }
            QPushButton#backupSecondary {
                min-height: 32px;
                color: #f1e4c4;
                background-color: rgba(7, 42, 52, 168);
                border: 1px solid rgba(205, 159, 78, 185);
                border-radius: 8px;
                padding: 7px 14px;
                font-family: Georgia;
                font-weight: bold;
            }
            QPushButton#backupSecondary:hover {
                color: #fff0bd;
                background-color: rgba(12, 60, 70, 214);
                border: 2px solid #f3d79a;
            }
        """)

        root = QVBoxLayout(self)
        root.setContentsMargins(34, 26, 34, 26)
        root.setSpacing(12)

        title = QLabel("ZÁLOHA UŽIVATELSKÝCH DAT", self)
        title.setObjectName("backupTitle")
        title.setAlignment(Qt.AlignCenter)
        root.addWidget(title)

        subtitle = QLabel("Plán, historie, přílohy, oddíly, ubytování a sportovní den v jednom ZIPu.", self)
        subtitle.setObjectName("backupSubtitle")
        subtitle.setAlignment(Qt.AlignCenter)
        root.addWidget(subtitle)

        summary = QFrame(self)
        summary.setObjectName("backupSummary")
        summary_layout = QHBoxLayout(summary)
        summary_layout.setContentsMargins(16, 9, 16, 9)
        summary_layout.setSpacing(18)
        self.path_value_label = self._summary_value(summary, "ÚLOŽIŠTĚ")
        self.file_count_value_label = self._summary_value(summary, "SOUBORŮ")
        self.size_value_label = self._summary_value(summary, "VELIKOST")
        summary_layout.addWidget(self.path_value_label, 1)
        summary_layout.addWidget(self.file_count_value_label)
        summary_layout.addWidget(self.size_value_label)
        root.addWidget(summary)

        actions = QHBoxLayout()
        actions.setSpacing(14)
        actions.addWidget(self._action_panel(
            "EXPORT ZIP",
            "Uloží aktuální uživatelská data do jednoho přenositelného souboru.",
            ("Plán tábora", "Oddíly a ubytování", "Sportovní den", "Historie zpráv"),
            "ZIP",
            "VYTVOŘIT ZÁLOHU",
            self._export_clicked,
            "backupPrimary",
        ))
        actions.addWidget(self._action_panel(
            "IMPORT ZIP",
            "Obnoví data ze zálohy a před přepsáním uloží bezpečnostní kopii.",
            ("Kontrola zálohy", "Bezpečnostní kopie", "Obnova dat", "Nový start aplikace"),
            "IN",
            "OBNOVIT ZE ZÁLOHY",
            self._import_clicked,
            "backupDanger",
        ))
        root.addLayout(actions)

        note = QFrame(self)
        note.setObjectName("backupNote")
        note_layout = QHBoxLayout(note)
        note_layout.setContentsMargins(14, 9, 14, 9)
        note_text = QLabel(
            "Import před přepsáním dat automaticky uloží bezpečnostní kopii do složky aplikace.",
            note,
        )
        note_text.setObjectName("backupSmall")
        note_text.setWordWrap(True)
        note_text.setAlignment(Qt.AlignCenter)
        note_layout.addWidget(note_text)
        root.addWidget(note)
        root.addStretch(1)

        bottom = QHBoxLayout()
        open_folder = QPushButton("OTEVŘÍT SLOŽKU DAT", self)
        open_folder.setObjectName("backupSecondary")
        open_folder.clicked.connect(self._open_data_folder)
        bottom.addWidget(open_folder)
        bottom.addStretch(1)
        close_button = QPushButton("ZAVŘÍT", self)
        close_button.setObjectName("backupSecondary")
        close_button.clicked.connect(self.close)
        bottom.addWidget(close_button)
        root.addLayout(bottom)

        self._refresh_summary()
        self.apply_pirate_glass()

    def _summary_value(self, parent, label: str):
        box = QWidget(parent)
        layout = QVBoxLayout(box)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(2)
        name = QLabel(label, box)
        name.setObjectName("backupSummaryName")
        name.setAlignment(Qt.AlignCenter)
        value = QLabel("", box)
        value.setObjectName("backupSummaryValue")
        value.setAlignment(Qt.AlignCenter)
        value.setWordWrap(True)
        layout.addWidget(name)
        layout.addWidget(value)
        box.value_label = value
        return box

    def _action_panel(
        self,
        title: str,
        text: str,
        items,
        badge_text: str,
        button_text: str,
        callback,
        button_object_name: str,
    ):
        panel = QFrame(self)
        panel.setObjectName("backupPanel")
        panel.setMinimumHeight(250)
        layout = QVBoxLayout(panel)
        layout.setContentsMargins(18, 14, 18, 16)
        layout.setSpacing(10)
        badge = QLabel(badge_text, panel)
        badge.setObjectName("backupBadge")
        badge.setAlignment(Qt.AlignCenter)
        layout.addWidget(badge, 0, Qt.AlignCenter)
        title_label = QLabel(title, panel)
        title_label.setObjectName("backupPanelTitle")
        title_label.setAlignment(Qt.AlignCenter)
        layout.addWidget(title_label)
        text_label = QLabel(text, panel)
        text_label.setObjectName("backupPanelText")
        text_label.setWordWrap(True)
        text_label.setAlignment(Qt.AlignCenter)
        layout.addWidget(text_label)
        item_box = QWidget(panel)
        item_layout = QGridLayout(item_box)
        item_layout.setContentsMargins(6, 2, 6, 2)
        item_layout.setHorizontalSpacing(10)
        item_layout.setVerticalSpacing(5)
        for index, item in enumerate(tuple(items or ())):
            item_label = QLabel(f"• {item}", item_box)
            item_label.setObjectName("backupPanelItem")
            item_label.setAlignment(Qt.AlignLeft | Qt.AlignVCenter)
            item_layout.addWidget(item_label, index // 2, index % 2)
        layout.addWidget(item_box)
        layout.addStretch(1)
        button = QPushButton(button_text, panel)
        button.setObjectName(button_object_name)
        button.clicked.connect(callback)
        layout.addWidget(button)
        return panel

    def _refresh_summary(self):
        data_dir = get_user_data_dir()
        file_count = 0
        total_bytes = 0
        try:
            for current_root, _dirs, files in os.walk(data_dir):
                for file_name in files:
                    path = os.path.join(current_root, file_name)
                    file_count += 1
                    try:
                        total_bytes += os.path.getsize(path)
                    except OSError:
                        pass
        except Exception:
            pass
        self.path_value_label.value_label.setText(data_dir)
        self.file_count_value_label.value_label.setText(str(file_count))
        self.size_value_label.value_label.setText(_user_data_format_size(total_bytes))

    def _export_clicked(self):
        if self.owner_window is not None:
            self.owner_window.export_user_data_backup()
        self._refresh_summary()

    def _import_clicked(self):
        if self.owner_window is not None:
            self.owner_window.import_user_data_backup()
        self._refresh_summary()

    def _open_data_folder(self):
        QDesktopServices.openUrl(QUrl.fromLocalFile(get_user_data_dir()))


def _window_build_user_data_menu(self):
    self._user_data_menu_ready = True


def _window_export_user_data_backup(self):
    default_path = os.path.join(_user_data_backup_default_dir(), default_backup_filename(APP_VERSION))
    path, _filter = QFileDialog.getSaveFileName(
        self,
        "Exportovat uživatelská data",
        default_path,
        "Záloha Šifrátoru (*.zip);;Všechny soubory (*)",
    )
    if not path:
        return
    try:
        stats = export_user_data_zip(path, app_version=APP_VERSION)
        self.write_live_log(
            f"Export uživatelských dat dokončen: {stats['file_count']} souborů -> {stats['path']}"
        )
        QMessageBox.information(
            self,
            "Export uživatelských dat",
            (
                "Záloha byla vytvořena.\n\n"
                f"Soubor: {stats['path']}\n"
                f"Zahrnuto souborů: {stats['file_count']}\n"
                f"Velikost dat: {_user_data_format_size(stats['bytes'])}"
            ),
        )
    except UserDataBackupError as error:
        self.write_live_log(f"Export uživatelských dat selhal: {error}")
        QMessageBox.warning(self, "Export uživatelských dat", str(error))
    except Exception as error:
        self.write_live_log(f"Export uživatelských dat selhal: {type(error).__name__}: {error}")
        QMessageBox.warning(self, "Export uživatelských dat", f"Zálohu se nepodařilo vytvořit:\n{error}")


def _window_import_user_data_backup(self):
    path, _filter = QFileDialog.getOpenFileName(
        self,
        "Importovat uživatelská data",
        _user_data_backup_default_dir(),
        "Záloha Šifrátoru (*.zip);;Všechny soubory (*)",
    )
    if not path:
        return

    try:
        info = inspect_user_data_backup(path)
    except UserDataBackupError as error:
        QMessageBox.warning(self, "Import uživatelských dat", str(error))
        return

    created = info.get("created_at") or "neznámé datum"
    version = info.get("app_version") or "neznámá verze"
    answer = QMessageBox.question(
        self,
        "Importovat uživatelská data?",
        (
            "Opravdu chcete obnovit uživatelská data z vybraného ZIPu?\n\n"
            f"Záloha: {path}\n"
            f"Vytvořeno: {created}\n"
            f"Verze aplikace v záloze: {version}\n"
            f"Souborů v záloze: {info['file_count']}\n"
            f"Velikost dat: {_user_data_format_size(info['bytes'])}\n\n"
            "Aktuální data se před importem automaticky uloží do bezpečnostní ZIP zálohy. "
            "Po importu aplikaci restartujte, aby se všechny otevřené části načetly z obnovených souborů."
        ),
        QMessageBox.Yes | QMessageBox.No,
        QMessageBox.No,
    )
    if answer != QMessageBox.Yes:
        return

    try:
        stats = import_user_data_zip(path, app_version=APP_VERSION)
        self.write_live_log(
            f"Import uživatelských dat dokončen: {stats['imported_files']} souborů z {path}"
        )
        QMessageBox.information(
            self,
            "Import uživatelských dat",
            (
                "Import byl dokončen.\n\n"
                f"Obnoveno souborů: {stats['imported_files']}\n"
                f"Automatická záloha původních dat:\n{stats['backup_path']}\n\n"
                "Teď aplikaci zavřete a znovu spusťte, aby se všude načetla obnovená data."
            ),
        )
        try:
            self.central.update_status()
        except Exception:
            pass
    except UserDataBackupError as error:
        self.write_live_log(f"Import uživatelských dat selhal: {error}")
        QMessageBox.warning(self, "Import uživatelských dat", str(error))
    except Exception as error:
        self.write_live_log(f"Import uživatelských dat selhal: {type(error).__name__}: {error}")
        QMessageBox.warning(self, "Import uživatelských dat", f"Import se nepodařil:\n{error}")


def _window_show_user_data_backup_window(self):
    dialog = getattr(self, "_user_data_backup_dialog", None)
    if dialog is None:
        dialog = UserDataBackupDialog(self)
        self._user_data_backup_dialog = dialog
    else:
        dialog._refresh_summary()
    self._show_embedded_dialog_page(dialog, "Záloha dat")


SifratorWindow.build_user_data_menu = _window_build_user_data_menu
SifratorWindow.export_user_data_backup = _window_export_user_data_backup
SifratorWindow.import_user_data_backup = _window_import_user_data_backup
SifratorWindow.show_user_data_backup_window = _window_show_user_data_backup_window

try:
    _USER_DATA_MENU_ORIGINAL_WINDOW_INIT = SifratorWindow.__init__

    def _user_data_menu_window_init(self, *args, **kwargs):
        result = _USER_DATA_MENU_ORIGINAL_WINDOW_INIT(self, *args, **kwargs)
        self.build_user_data_menu()
        return result

    SifratorWindow.__init__ = _user_data_menu_window_init
except Exception:
    pass


# Důležité uživatelské akce zapisujeme do stejného logu jako aktualizace a cache.
def _sifrator_widget_window(widget):
    try:
        window = widget.window()
        if isinstance(window, SifratorWindow):
            return window
    except Exception:
        pass
    return None


def _sifrator_log_from_widget(widget, message: str) -> None:
    window = _sifrator_widget_window(widget)
    if window is not None and hasattr(window, "write_live_log"):
        window.write_live_log(message)
    else:
        _sifrator_debug_log(message)


def _sifrator_text_len(widget, attr_name: str) -> int:
    try:
        edit = getattr(widget, attr_name, None)
        return len(edit.toPlainText().strip()) if edit is not None else 0
    except Exception:
        return 0


try:
    _LOG_ORIGINAL_WINDOW_INIT = SifratorWindow.__init__

    def _log_window_init(self, *args, **kwargs):
        result = _LOG_ORIGINAL_WINDOW_INIT(self, *args, **kwargs)
        try:
            self.write_live_log(
                f"Start aplikace: version={APP_VERSION}, app_dir={get_app_dir()}, "
                f"icons_dir={get_icons_dir()}, cache_log={_sifrator_debug_log_path()}"
            )
        except Exception:
            pass
        return result

    SifratorWindow.__init__ = _log_window_init
except Exception:
    pass


try:
    _LOG_ORIGINAL_SELECT_CIPHER = SifratorSkinWidget.select_cipher

    def _log_select_cipher(self, name):
        previous = getattr(self, "selected_cipher", None) or "žádná"
        try:
            _sifrator_log_from_widget(self, f"Výběr šifry: {previous} -> {name}")
        except Exception:
            pass
        try:
            return _LOG_ORIGINAL_SELECT_CIPHER(self, name)
        except Exception as error:
            _sifrator_log_from_widget(self, f"Chyba při výběru šifry {name}: {type(error).__name__}: {error}")
            raise

    SifratorSkinWidget.select_cipher = _log_select_cipher
except Exception:
    pass


try:
    _LOG_ORIGINAL_ENCRYPT_ACTION = SifratorSkinWidget.encrypt_action

    def _log_encrypt_action(self):
        cipher_name = getattr(self, "selected_cipher", None) or "žádná"
        input_len = _sifrator_text_len(self, "input_text")
        _sifrator_log_from_widget(self, f"Ruční šifrování: šifra={cipher_name}, délka_vstupu={input_len}")
        try:
            result = _LOG_ORIGINAL_ENCRYPT_ACTION(self)
        except Exception as error:
            _sifrator_log_from_widget(self, f"Chyba ručního šifrování: {type(error).__name__}: {error}")
            raise
        output_len = _sifrator_text_len(self, "output_text")
        _sifrator_log_from_widget(self, f"Ruční šifrování hotovo: šifra={cipher_name}, délka_výstupu={output_len}")
        return result

    SifratorSkinWidget.encrypt_action = _log_encrypt_action
except Exception:
    pass


try:
    _RANDOM_EASY_ORIGINAL_ENCRYPT_ACTION = SifratorSkinWidget.encrypt_action

    def _random_easy_encrypt_action(self):
        if getattr(self, "selected_cipher", None) == _RANDOM_EASY_CIPHER_NAME:
            text = self.get_input_text() if hasattr(self, "get_input_text") else ""
            if str(text or "").strip():
                self._random_easy_force_next_encrypt = True
        return _RANDOM_EASY_ORIGINAL_ENCRYPT_ACTION(self)

    SifratorSkinWidget.encrypt_action = _random_easy_encrypt_action
except Exception:
    pass


try:
    _LOG_ORIGINAL_DECRYPT_ACTION = SifratorSkinWidget.decrypt_action

    def _log_decrypt_action(self):
        cipher_name = getattr(self, "selected_cipher", None) or "žádná"
        input_len = _sifrator_text_len(self, "input_text")
        _sifrator_log_from_widget(self, f"Ruční dešifrování: šifra={cipher_name}, délka_vstupu={input_len}")
        try:
            result = _LOG_ORIGINAL_DECRYPT_ACTION(self)
        except Exception as error:
            _sifrator_log_from_widget(self, f"Chyba ručního dešifrování: {type(error).__name__}: {error}")
            raise
        output_len = _sifrator_text_len(self, "output_text")
        _sifrator_log_from_widget(self, f"Ruční dešifrování hotovo: šifra={cipher_name}, délka_výstupu={output_len}")
        return result

    SifratorSkinWidget.decrypt_action = _log_decrypt_action
except Exception:
    pass


try:
    _LOG_ORIGINAL_SHOW_CIPHER_KEY = SifratorSkinWidget.show_cipher_key

    def _log_show_cipher_key(self):
        cipher_name = getattr(self, "selected_cipher", None) or "žádná"
        _sifrator_log_from_widget(self, f"Otevření klíče: šifra={cipher_name}")
        try:
            return _LOG_ORIGINAL_SHOW_CIPHER_KEY(self)
        except Exception as error:
            _sifrator_log_from_widget(self, f"Chyba při otevření klíče {cipher_name}: {type(error).__name__}: {error}")
            raise

    SifratorSkinWidget.show_cipher_key = _log_show_cipher_key
except Exception:
    pass


try:
    _LOG_ORIGINAL_PRINT_CURRENT_RESULT = SifratorSkinWidget.print_current_result

    def _log_print_current_result(self):
        cipher_name = getattr(self, "selected_cipher", None) or "žádná"
        _sifrator_log_from_widget(
            self,
            f"Otevření náhledu tisku: šifra={cipher_name}, "
            f"délka_vstupu={_sifrator_text_len(self, 'input_text')}, "
            f"délka_výstupu={_sifrator_text_len(self, 'output_text')}",
        )
        try:
            return _LOG_ORIGINAL_PRINT_CURRENT_RESULT(self)
        except Exception as error:
            _sifrator_log_from_widget(self, f"Chyba náhledu tisku: {type(error).__name__}: {error}")
            raise

    SifratorSkinWidget.print_current_result = _log_print_current_result
except Exception:
    pass


# ============================================================
# TRVALÁ CACHE KLÍČŮ PODLE VERZE APLIKACE
# ============================================================
# Renderer dostane aktuální APP_VERSION. Pokud se po aktualizaci verze změní,
# pirate_key_renderer.py automaticky smaže starou diskovou cache klíčů.

try:
    _PERSISTENT_CACHE_ORIGINAL_GET_RENDERER = get_pirate_key_renderer

    def get_pirate_key_renderer():
        renderer = _PERSISTENT_CACHE_ORIGINAL_GET_RENDERER()
        try:
            if renderer is not None and hasattr(renderer, "set_persistent_cache_version"):
                renderer.set_persistent_cache_version(APP_VERSION)
        except Exception as error:
            try:
                _sifrator_debug_log(f"Cache klíčů: nastavení verze selhalo: {type(error).__name__}: {error}")
            except Exception:
                pass
        return renderer
except Exception:
    pass


def _clear_persistent_key_cache_from_main() -> bool:
    """Ručně smaže trvalou cache klíčů, pokud ji renderer podporuje."""
    try:
        renderer = get_pirate_key_renderer()
        if renderer is not None and hasattr(renderer, "clear_persistent_key_cache"):
            ok = bool(renderer.clear_persistent_key_cache())
            try:
                _sifrator_debug_log("Cache klíčů byla vyčištěna ručně." if ok else "Cache klíčů se nepodařilo vyčistit.")
            except Exception:
                pass
            return ok
    except Exception as error:
        try:
            _sifrator_debug_log(f"Ruční čištění cache klíčů selhalo: {type(error).__name__}: {error}")
        except Exception:
            pass
    return False


# Do okna LOGOVÁNÍ přidáme tlačítko pro ruční vyčištění cache klíčů.
try:
    _CACHE_DIALOG_ORIGINAL_INIT = LiveLogDialog.__init__

    def _cache_dialog_init(self, owner_window):
        _CACHE_DIALOG_ORIGINAL_INIT(self, owner_window)
        try:
            self.clear_cache_button = QPushButton("Vyčistit cache klíčů", self)
            self.clear_cache_button.setToolTip("Smaže uložené PNG klíče. Při dalším použití se vytvoří znovu.")
            self.clear_cache_button.clicked.connect(self.clear_key_cache)

            # Tlačítko vložíme vedle ostatních tlačítek ve spodním řádku dialogu.
            layout = self.layout()
            if layout is not None:
                row = layout.itemAt(layout.count() - 1)
                if row is not None and row.layout() is not None:
                    row.layout().insertWidget(max(0, row.layout().count() - 1), self.clear_cache_button)
        except Exception:
            pass

    def _cache_dialog_clear_key_cache(self):
        ok = _clear_persistent_key_cache_from_main()
        if ok:
            QMessageBox.information(self, "Cache klíčů", "Cache klíčů byla vyčištěna.")
        else:
            QMessageBox.warning(self, "Cache klíčů", "Cache klíčů se nepodařilo vyčistit.")
        try:
            self.refresh_log()
        except Exception:
            pass

    LiveLogDialog.__init__ = _cache_dialog_init
    LiveLogDialog.clear_key_cache = _cache_dialog_clear_key_cache
except Exception:
    pass

def run_smoke_test() -> int:
    """Rychlý test pro GitHub Actions bez otevření hlavního okna aplikace."""
    try:
        if hasattr(sys.stdout, "reconfigure"):
            sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        if hasattr(sys.stderr, "reconfigure"):
            sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

    print("Smoke test: start")
    print(f"APP_NAME: {APP_NAME}")
    print(f"APP_VERSION: {APP_VERSION}")
    print(f"APP_DIR: {get_app_dir()}")
    print(f"ICONS_DIR: {get_icons_dir()}")

    errors = []

    def check(condition: bool, ok_message: str, error_message: str) -> None:
        if condition:
            print(f"OK: {ok_message}")
        else:
            errors.append(error_message)

    check(os.path.isdir(get_icons_dir()), "icons", "Chybi slozka icons")
    check(
        os.path.exists(os.path.join(get_icons_dir(), "groups_BG.png")),
        "groups_BG.png",
        "Chybi piratske pozadi modulu Oddily",
    )
    for relative_asset in (
        os.path.join("diplomas", "sports_d.png"),
        os.path.join("documents", "daily_a.png"),
        os.path.join("documents", "daily_b.png"),
        os.path.join("documents", "cleaning_award_a.png"),
        os.path.join("documents", "cleaning_award_b.png"),
        os.path.join("documents", "meal_a.png"),
        os.path.join("documents", "meal_b.png"),
    ):
        check(
            os.path.exists(os.path.join(get_icons_dir(), relative_asset)),
            relative_asset,
            f"Chybi tiskovy podklad {relative_asset}",
        )
    check(
        os.path.isdir(os.path.join(get_app_dir(), "logika_sifer"))
        or os.path.isdir(os.path.join(get_script_dir(), "logika_sifer"))
        or os.path.isdir(os.path.join(get_app_dir(), "logika_sifer"))
        or os.path.isdir(os.path.join(get_script_dir(), "logika_sifer")),
        "logika_sifer",
        "Chybi slozka logika_sifer",
    )
    check(os.path.exists(get_pirate_key_renderer_file()), "pirate_key_renderer.py", "Chybi pirate_key_renderer.py")

    try:
        renderer = get_pirate_key_renderer()
        check(renderer is not None, "pirate_key_renderer import", "Nepodarilo se nacist pirate_key_renderer.py")
    except Exception as error:
        errors.append(f"Chyba pirate_key_renderer.py: {error}")

    try:
        import openpyxl  # noqa: F401
        try:
            import xlrd  # noqa: F401
        except ImportError:
            from vendor import xlrd  # noqa: F401
        check(True, "Excel import/export", "Chybi knihovny pro Excel")
    except Exception as error:
        errors.append(f"Chybi knihovny pro Excel import/export: {error}")

    logic_tests = list_cipher_names()

    for name in logic_tests:
        try:
            module = get_cipher_logic(name)
            if module is None:
                errors.append(f"Nepodarilo se nacist logiku: {name}")
            else:
                print(f"OK: {name}")
        except Exception as error:
            errors.append(f"Chyba logiky {name}: {error}")

    if errors:
        print("Smoke test: CHYBA")
        for error in errors:
            print(" -", str(error))
        return 1

    print("Smoke test: OK")
    return 0


def _prebuild_key_cache_output_dir() -> str | None:
    flag = "--prebuild-key-cache"
    if flag not in sys.argv:
        return None
    index = sys.argv.index(flag)
    if index + 1 < len(sys.argv) and not sys.argv[index + 1].startswith("--"):
        return sys.argv[index + 1]
    return ""


def _default_key_cache_context(cipher_name: str) -> dict:
    if cipher_name == "Caesarova šifra":
        return {"caesar_shift": 3, "caesar_direction": "dopředu"}
    return {}


def _install_key_cache_fonts() -> str:
    """Načte systémový font pro offscreen generování cache klíčů."""
    try:
        from PySide6.QtGui import QFontDatabase

        candidates = []
        windir = os.environ.get("WINDIR") or r"C:\Windows"
        candidates.extend([
            os.path.join(windir, "Fonts", "georgia.ttf"),
            os.path.join(windir, "Fonts", "georgiab.ttf"),
            os.path.join(windir, "Fonts", "times.ttf"),
            os.path.join(windir, "Fonts", "timesbd.ttf"),
            os.path.join(windir, "Fonts", "arial.ttf"),
        ])
        candidates.extend([
            "/System/Library/Fonts/Supplemental/Georgia.ttf",
            "/System/Library/Fonts/Supplemental/Times New Roman.ttf",
            "/System/Library/Fonts/Supplemental/Arial.ttf",
            "/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf",
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        ])

        loaded_families = []
        for path in candidates:
            if not path or not os.path.exists(path):
                continue
            font_id = QFontDatabase.addApplicationFont(path)
            if font_id >= 0:
                try:
                    loaded_families.extend(QFontDatabase.applicationFontFamilies(font_id))
                except Exception:
                    pass

        available = set(QFontDatabase.families()) | set(loaded_families)
        app = QApplication.instance()
        for family in ("Georgia", "Times New Roman", "Times", "DejaVu Serif", "Arial", "DejaVu Sans"):
            if family in available:
                if app is not None:
                    app.setFont(QFont(family, 10))
                return family
    except Exception:
        pass
    return ""


def _write_cli_progress(phase: str, done: int, total: int, detail: str, status_label: str) -> None:
    import shutil

    total = max(1, int(total or 1))
    done = max(0, min(total, int(done or 0)))
    width = 30
    filled = int(width * done / total)
    bar = "#" * filled + "-" * (width - filled)
    percent = int(round(done * 100 / total))

    columns = max(60, min(160, shutil.get_terminal_size((100, 20)).columns))
    base = f"{phase} [{bar}] {percent:3d}% {done}/{total} {status_label}"
    if detail:
        room = max(0, columns - len(base) - 4)
        text = str(detail)
        if room and len(text) > room:
            text = text[:max(0, room - 3)] + "..."
        line = f"{base}: {text}" if room else base
    else:
        line = base

    try:
        line = line[:max(1, columns - 1)]
        previous_len = int(getattr(_write_cli_progress, "_previous_len", 0))
        clear_tail = " " * max(0, previous_len - len(line))
        sys.stdout.write("\r" + line + clear_tail)
        if done >= total:
            sys.stdout.write("\n")
            _write_cli_progress._previous_len = 0
        else:
            _write_cli_progress._previous_len = len(line)
        sys.stdout.flush()
    except Exception:
        pass


def _print_key_cache_clear_progress(done: int, total: int, file_name: str, status: str) -> None:
    if int(total or 0) <= 0:
        _write_cli_progress("Mazání   ", 1, 1, "", "nic ke smazání")
        print("Mazání hotovo: nic ke smazání.")
        return
    if int(done or 0) <= 1:
        _print_key_cache_clear_progress.deleted = 0
        _print_key_cache_clear_progress.failed = 0
    if status == "deleted":
        _print_key_cache_clear_progress.deleted = int(getattr(_print_key_cache_clear_progress, "deleted", 0)) + 1
    elif status == "failed":
        _print_key_cache_clear_progress.failed = int(getattr(_print_key_cache_clear_progress, "failed", 0)) + 1

    status_label = {
        "deleted": "smazáno",
        "failed": "chyba",
    }.get(str(status), str(status or "pracuji"))
    _write_cli_progress("Mazání   ", done, total, file_name, status_label)
    if int(done or 0) >= int(total or 0):
        deleted = int(getattr(_print_key_cache_clear_progress, "deleted", 0))
        failed = int(getattr(_print_key_cache_clear_progress, "failed", 0))
        print(f"Mazání hotovo: smazáno {deleted}, chyb {failed}.")


def _print_key_cache_progress(done: int, total: int, cipher_name: str, print_mode: bool, status: str) -> None:
    mode = "tisk" if print_mode else "náhled"
    status_label = {
        "created": "vytvořeno",
        "skipped": "existuje",
        "failed": "chyba",
    }.get(str(status), str(status or "pracuji"))
    _write_cli_progress("Generuji ", done, total, f"{cipher_name} ({mode})", status_label)


def run_prebuild_key_cache() -> int:
    """Předgeneruje cache PNG klíčů do složky cache/key_cache pro publikování."""
    try:
        if hasattr(sys.stdout, "reconfigure"):
            sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        if hasattr(sys.stderr, "reconfigure"):
            sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

    try:
        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
        os.environ.setdefault("QT_LOGGING_RULES", "qt.qpa.fonts=false")
        app = QApplication.instance() or QApplication(sys.argv[:1])
        _ = app
        cache_font = _install_key_cache_fonts()
    except Exception as error:
        print(f"Cache klíčů: nepodařilo se spustit Qt: {error}")
        return 1

    renderer = get_pirate_key_renderer()
    if renderer is None:
        print("Cache klíčů: chybí pirate_key_renderer.py")
        return 1
    if not hasattr(renderer, "prebuild_key_cache_to_dir"):
        print("Cache klíčů: renderer neumí předgenerovat cache")
        return 1

    output_arg = _prebuild_key_cache_output_dir()
    if output_arg:
        root = os.path.abspath(output_arg)
    elif hasattr(renderer, "get_bundled_key_cache_dir"):
        root = renderer.get_bundled_key_cache_dir()
    else:
        root = os.path.join(get_app_dir(), "cache", "key_cache")

    items = []
    skipped_logic = 0
    for cipher_name in list_cipher_names():
        module = get_cipher_logic(cipher_name)
        if module is None:
            skipped_logic += 1
            print(f"Cache klíčů: přeskočeno, nejde načíst logiku: {cipher_name}")
            continue
        items.append((cipher_name, module, _default_key_cache_context(cipher_name)))

    print("Cache klíčů: start")
    print(f"Cíl: {root}")
    print(f"Šifer: {len(items)}")
    if cache_font:
        print(f"Font: {cache_font}")
    print("Staré soubory cache v cílové složce se smažou a vytvoří znovu.")
    stats = renderer.prebuild_key_cache_to_dir(
        items,
        root=root,
        width=1400,
        include_print=True,
        include_ui=True,
        progress_callback=_print_key_cache_progress,
        clear_progress_callback=_print_key_cache_clear_progress,
        clear_existing=True,
    )

    created = int(stats.get("created", 0))
    skipped = int(stats.get("skipped", 0))
    failed = int(stats.get("failed", 0))
    total = int(stats.get("total", 0))
    deleted = int(stats.get("deleted", 0))
    delete_failed = int(stats.get("delete_failed", 0))

    print(f"Hotovo: celkem {total}, smazáno {deleted}, chyb mazání {delete_failed}, vytvořeno {created}, už existovalo {skipped}, chyb {failed}, bez logiky {skipped_logic}")
    return 0 if failed == 0 and delete_failed == 0 and skipped_logic == 0 and (created + skipped) > 0 else 1


if "--smoke-test" in sys.argv:
    sys.exit(run_smoke_test())

if "--prebuild-key-cache" in sys.argv:
    sys.exit(run_prebuild_key_cache())

if __name__ == "__main__":
    # Musí být před QApplication, jinak si Windows může držet ikonu python.exe.
    set_windows_app_id()

    app = QApplication(sys.argv)
    app.setApplicationName("Šifrátor Mraveniště")
    app.setApplicationDisplayName("Šifrátor Mraveniště")

    window = SifratorWindow()
    window.show()

    sys.exit(app.exec())
